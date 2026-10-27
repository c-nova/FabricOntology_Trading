"""Structural checks and a real local LibreOffice render of the customer deck.

Renders decks/TradingOntology_Customer_Deck_JA.pptx to presentation/qa/customer/ (git-ignored): one PNG per
slide and contact sheets. Needs macOS LibreOffice and a Japanese font (see fonts.conf).
"""
import json
import os
from pathlib import Path
import subprocess
import zipfile

import fitz
from PIL import Image, ImageDraw
from lxml import etree
from pptx import Presentation

HERE = Path(__file__).resolve().parent
QA = HERE / "qa" / "customer"
DECK = HERE.parent / "decks" / "TradingOntology_Customer_Deck_JA.pptx"
SOFFICE = os.environ.get("SOFFICE", "/Applications/LibreOffice.app/Contents/MacOS/soffice")
QA.mkdir(parents=True, exist_ok=True)
for old in QA.glob("slide-*.png"):
    old.unlink()
prs = Presentation(DECK)
n_slides = len(prs.slides)
oob, notes_len = [], []
for n, slide in enumerate(prs.slides, 1):
    notes_len.append(len(slide.notes_slide.notes_text_frame.text))
    for shape in slide.shapes:
        if (min(shape.left, shape.top) < -100 or shape.left + shape.width > prs.slide_width + 100
                or shape.top + shape.height > prs.slide_height + 100):
            oob.append([n, shape.name])
assert not oob, oob
with zipfile.ZipFile(DECK) as z:
    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            etree.fromstring(z.read(name))
    externals = []
    for name in z.namelist():
        if name.startswith("ppt/slides/_rels/") and name.endswith(".rels"):
            for r in etree.fromstring(z.read(name)):
                if r.get("TargetMode") == "External":
                    externals.append(r.get("Target"))
assert not externals, externals

env = dict(os.environ, TMPDIR=str(QA / "tmp"), FONTCONFIG_FILE=str(HERE / "fonts.conf"))
(QA / "tmp").mkdir(exist_ok=True)
subprocess.run([SOFFICE, f"-env:UserInstallation={(QA / 'lo-profile').as_uri()}", "--headless",
                "--convert-to", "pdf", "--outdir", str(QA), str(DECK)],
               env=env, check=True, capture_output=True, text=True, timeout=240)
pdf = QA / (DECK.stem + ".pdf")
doc = fitz.open(pdf)
assert len(doc) == n_slides, (len(doc), n_slides)
no_cjk = []
for n, page in enumerate(doc, 1):
    page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(QA / f"slide-{n:02}.png")
    if not any("\u3040" <= c <= "\u9fff" for c in page.get_text()):
        no_cjk.append(n)
for start in range(0, n_slides, 6):
    sheet = Image.new("RGB", (1980, 790), "white")
    for j, n in enumerate(range(start + 1, min(start + 7, n_slides + 1))):
        im = Image.open(QA / f"slide-{n:02}.png").convert("RGB")
        im.thumbnail((640, 360))
        canvas = Image.new("RGB", (660, 395), "#d9e1e6")
        canvas.paste(im, (10, 25))
        ImageDraw.Draw(canvas).text((10, 5), f"{n:02}", fill="black")
        sheet.paste(canvas, ((j % 3) * 660, (j // 3) * 395))
    sheet.save(QA / f"contact-{start + 1:02}.png")
print(json.dumps({"slides": n_slides, "pdf_pages": len(doc), "min_notes_chars": min(notes_len),
                  "slides_without_cjk_text": no_cjk, "external_links": len(externals)}, ensure_ascii=False))
