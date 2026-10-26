"""Structural checks and real local LibreOffice rendering of the presenter deck."""
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation
from lxml import etree

HERE = Path(__file__).resolve().parent
QA = HERE / "qa"
DECK = HERE.parent / "decks" / "TradingOntology_Demo_Presenter_Guide_JA.pptx"
SOFFICE = os.environ.get("SOFFICE", "/Applications/LibreOffice.app/Contents/MacOS/soffice")
QA.mkdir(exist_ok=True)
prs = Presentation(DECK)
assert len(prs.slides) == 31
slide_count = len(prs.slides)
notes = []
out_of_bounds = []
for n, slide in enumerate(prs.slides, 1):
    note = slide.notes_slide.notes_text_frame.text
    assert len(note) >= 300, (n, "Missing detailed notes")
    notes.append(len(note))
    for shape in slide.shapes:
        if min(shape.left, shape.top) < 0 or (
                shape.left + shape.width > prs.slide_width + 100 or
                shape.top + shape.height > prs.slide_height + 100):
            out_of_bounds.append([n, shape.name])
assert not out_of_bounds, out_of_bounds
with zipfile.ZipFile(DECK) as z:
    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            etree.fromstring(z.read(name))
    links = []
    for name in z.namelist():
        if name.startswith("ppt/slides/_rels/") and name.endswith(".rels"):
            for r in etree.fromstring(z.read(name)):
                if r.get("TargetMode") == "External":
                    links.append(r.get("Target"))
    assert all(x.startswith("https://app.fabric.microsoft.com/") for x in links)
    assert len(links) >= 6

env = dict(os.environ, TMPDIR=str(QA / "local-runtime"),
           FONTCONFIG_FILE=str(HERE / "fonts.conf"))
(QA / "local-runtime").mkdir(exist_ok=True)
subprocess.run([
    SOFFICE,
    f"-env:UserInstallation={(QA / 'lo-profile-final').as_uri()}",
    "--headless", "--convert-to", "pdf", "--outdir", str(QA), str(DECK)
], env=env, check=True, capture_output=True, text=True, timeout=120)
pdf = QA / (DECK.stem + ".pdf")
doc = fitz.open(pdf)
assert len(doc) == slide_count
rendered_text = []
for n, page in enumerate(doc, 1):
    page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(QA / f"slide-{n:02}.png")
    rendered_text.append(page.get_text())
for start in range(0, slide_count, 9):
    sheet = Image.new("RGB", (1980, 1185), "white")
    for j, n in enumerate(range(start + 1, min(start + 10, slide_count + 1))):
        im = Image.open(QA / f"slide-{n:02}.png").convert("RGB")
        im.thumbnail((640, 360))
        canvas = Image.new("RGB", (660, 395), "#d9e1e6")
        canvas.paste(im, (10, 25))
        ImageDraw.Draw(canvas).text((10, 5), f"{n:02}", fill="black")
        sheet.paste(canvas, ((j % 3) * 660, (j // 3) * 395))
    sheet.save(QA / f"contact-sheet-{start+1:02}.png")

# Extraction is independent of the LibreOffice glyph rendering.
extracted = []
for n, slide in enumerate(prs.slides, 1):
    extracted.append(f"=== SLIDE {n} ===\n" + "\n".join(
        s.text for s in slide.shapes if s.has_text_frame))
(QA / "extracted_text.txt").write_text("\n\n".join(extracted))
(QA / "rendered_pdf_text.txt").write_text("\n\n".join(rendered_text))
report = {
    "deck": str(DECK), "slide_count": slide_count,
    "notes_characters_by_slide": notes,
    "all_slides_have_detailed_notes": True,
    "xml_well_formed": True, "out_of_bounds_shapes": out_of_bounds,
    "fabric_clickable_hyperlinks": len(links),
    "rendering": "LibreOffice 26.8.1.1 PDF and PyMuPDF PNG",
    "render_font_substitute": "Yu Gothic -> installed Noto Sans CJK JP via local fonts.conf",
    "chart_fix": "Signed bars use editable shapes; native chart negative rendering was incorrect in LibreOffice. Underlying source amounts verified.",
    "notes": "Native line chart remains editable. Actual DAG is a local browser capture of saved SVG, not a Fabric UI screenshot.",
    "markitdown": "Installed namespace lacks API/CLI; text extracted directly with python-pptx and independently from PDF.",
    "fresh_visual_review": "See session review of revised slide PNGs; not inferred from the original deck review.",
    "user_final_check": "Windows PowerPoint font substitution, linebreaks, notes, links",
}
(QA / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps({"slides": slide_count, "notes_min_chars": min(notes),
                  "rendered_pages": len(doc), "out_of_bounds": out_of_bounds}))
