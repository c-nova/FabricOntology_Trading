# -*- coding: utf-8 -*-
"""Build the Japanese presenter guide deck from the verified Trading demo artifacts.

Run from any directory with python3 (python-pptx and Pillow are required).
Writes decks/TradingOntology_Demo_Presenter_Guide_JA.pptx, presentation/slide_sources.json
and prompts/presenter_prompts_ja.txt. No Fabric API is called. python-pptx is used because
the pptxgenjs probe failed in the original environment. The DAG image is a browser capture
of the saved local SVG (reports/step3_food_loss_dag.svg), not a Fabric UI screenshot.
"""
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
OUT = PROJECT / "decks" / "TradingOntology_Demo_Presenter_Guide_JA.pptx"
OUT.parent.mkdir(exist_ok=True)
QA = HERE / "qa"
QA.mkdir(exist_ok=True)
manifest = json.loads((PROJECT / "config/deployment_manifest.json").read_text())
audit = json.loads((PROJECT / "reports/step3_live_audit.json").read_text())
step3 = json.loads((PROJECT / "config/step3_deployment.json").read_text())
events = json.loads((PROJECT / "reports/qa_expected_event_totals.json").read_text())
dataset = json.loads((PROJECT / "data/demo_dataset.json").read_text())["LH_Group_SAP"]

NAVY = "17334A"
TEAL = "087F83"
CORAL = "BD4354"
INK = "21374A"
MUTED = "566A78"
BG = "F4F7F7"
WHITE = "FFFFFF"
PALE = "E6F2EF"
GRAY = "E5EBEE"
PURPLE = "7352AD"
FONT = "Yu Gothic"
SW, SH = 13.333333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
prs.core_properties.title = "Trading Fabric Ontology デモ説明者ガイド"
prs.core_properties.subject = "営業向け基礎解説、DA入力全文と対応結果、食品損失DAG、説明台本"
prs.core_properties.author = "Microsoft Scout"
prs.core_properties.keywords = "Japanese, Presenter Enablement, Synthetic, Fabric, Ontology"
slide_metadata = []

DA = manifest["data_agent"]["url"]
ONTOLOGY = manifest["ontology"]["url"]
NOTEBOOK = step3["url"]
WORKSPACE = manifest["workspace"]["url"]

ONT = PROJECT / "ontology_tmdl"
node_counts = {r["entity_type"]: r["instance_count"] for r in json.loads(
    (PROJECT / "reports/graph_node_counts.json").read_text())["result"]["data"]}
edge_counts = {r["relationship_type"]: r["edge_count"] for r in json.loads(
    (PROJECT / "reports/graph_edge_counts.json").read_text())["result"]["data"]}
cloud = json.loads((PROJECT / "reports/cloud_validation.json").read_text())


def parse_ontology():
    entities = {}
    for f in sorted((ONT / "entities").glob("*.tmdl")):
        t = f.read_text()
        name = re.search(r"^entity (\w+)", t, re.M).group(1)
        backing = re.search(r"backingTable: (\w+)", t).group(1)
        table_def = (ONT / "tables" / f"{backing}.tmdl").read_text()
        entities[name] = {
            "key": re.search(r"keyProperty: (\w+)", t).group(1),
            "properties": len(re.findall(r"^\tproperty ", t, re.M)),
            "table": re.search(r"entityName: (\w+)", table_def).group(1),
            "mode": re.search(r"mode: (\w+)", table_def).group(1),
        }
    relationships = []
    for block in re.split(r"\n(?=entityRelationship )", (ONT / "entityRelationships.tmdl").read_text()):
        m = re.search(r"label: (\w+)\n\s+fromEntity: (\w+)\n\s+toEntity: (\w+)", block)
        if m:
            relationships.append({"name": m.group(1), "from": m.group(2), "to": m.group(3),
                                  "link_table": "type: table" in block})
    return entities, relationships


ontology_entities, ontology_relationships = parse_ontology()
assert len(ontology_entities) == 17 and len(ontology_relationships) == 39
assert sum(node_counts.values()) == 33960 and sum(edge_counts.values()) == 65684
assert sum(r["link_table"] for r in ontology_relationships) == 3

Q1 = """2025-10-01～2026-09-30を対象に、全額連結EBITDAのbaselineとactualを月次で集計し、折れ線グラフで比較してください。横軸は「2025-10」のような年月の文字列で、2025-10から2026-09までの12か月だけを載せ、期間合計はグラフに含めないでください。縦軸は円を1億で割った億円にしてください。actualとbaselineは別系列にし、同じ行を重複集計しないでください。グラフとは別に、期間全体のactualの全額連結EBITDAと持分管理寄与EBITDAを億円で示し、51%子会社の扱いの違いを説明してください。"""
Q2 = """2025-10-01～2026-09-30の全10イベントについて、全額連結EBITDAへの影響を符号付きの億円で集計し、棒グラフで表示してください。各effect_idは一度だけ合計し、CROSS_NEGとCROSS_POSを含めてください。グラフの項目名は次の短い日本語ラベルを使ってください。
MTL_NEG=製錬停止、MTL_POS=再生合金、ENE_NEG=LNG遅延、ENE_POS=配分最適化、FOD_NEG=食品損失、FOD_POS=調達改善、FIN_NEG=与信費増、FIN_POS=ヘッジ手数料、CROSS_NEG=横断混乱、CROSS_POS=横断回復。
元のevent_idと影響額の対応表も示してください。"""
Q3 = """食品損失イベントFOD_NEGの全額連結EBITDA影響額と持分管理寄与額、重複のないeffect_id件数を示してください。根拠ID、代表する活動ID、契約ID、効果ID、日次指標ID、全社KPI IDを確認できる形で短く説明してください。代表効果はeffect_idの昇順で先頭の1件にし、日次指標IDと全社KPI IDは、その代表効果と同じ日付のactual・全額連結EBITDAのものにしてください。架空の設計上の影響配賦であることを明記してください。図やMermaidは不要です。関係の図は別のNotebookで確認します。"""
Q4 = """ENE_POSについて、根拠、対象活動、契約、効果、EBITDA、全社KPIのつながりをID付きで説明してください。全額JPYと持分管理JPYを区別し、配分最適化の設計上の機構と未確認事項を示してください。架空データであり、実世界の因果証明ではないと明記してください。"""
Q5 = """CROSS_NEGとCROSS_POSを、配賦先会社のsectorで4セクターに分類して比較してください。primary_sector='all'を除外せず、effect_idごとに一度だけ合計し、全額と持分管理のEBITDA影響を億円で示してください。時期の異なるイベントなので回復率とは呼ばないでください。"""


Q_ACT = """2025-10-01～2026-09-30のネガティブイベント（CROSS_NEGを含む5件）について、全額連結EBITDAへの影響が大きかった契約の上位5件と、2件以上のネガティブイベントの影響を受けた契約を、重複を除いて1つの表にまとめ、契約ID、子会社、商品、影響したイベント、影響額（億円）を示してください。primary_sector='all'のCROSS_NEGを除外せず、各effect_idは一度だけ合計し、影響を受けたネガティブイベントの数も示してください。続いて、表の各契約への対応を高、中、低の優先度で提案し、優先度は影響額と、複数のイベントの影響を受けたかで判定して、判定基準も示してください。データから言えることと、仮説としてのアイデアを分け、取引先名はデータにないため作らず、契約IDと商品で示してください。架空データであり、実世界の因果証明ではないと明記してください。"""
DA_ACTION_TEXT = """対応案
高：SUB_FIN_US_01_C08、SUB_FIN_US_01_C07、SUB_FIN_JP_01_C08、SUB_FIN_JP_01_C07

データから言えること：4契約はすべてtrade_financeで、CROSS_NEGとFIN_NEGの両方の影響を受け、-0.43～-0.45億円のマイナスです。
仮説としてのアイデア：契約別の信用枠、担保・ヘッジ条件、例外承認ルールを優先点検し、横断混乱と与信費上昇が同時に発生した場合の早期エスカレーション条件を設定します。
中：SUB_MTL_JP_01_C07

データから言えること：copper_alloyの契約で、CROSS_NEG単独により-0.37億円の影響です。
仮説としてのアイデア：代替調達、在庫バッファ、納期・価格条件の見直しなど、横断的な混乱に耐える契約運用を検討します。
低：SUB_FIN_JP_01_C09、SUB_FIN_US_01_C09

データから言えること：両契約はtrade_financeで、2イベントの影響を受けていますが、影響額は各-0.24億円で上位4契約より小さい水準です。
仮説としてのアイデア：高優先契約で有効性を確認した与信・ヘッジの監視ルールを横展開し、影響拡大時に優先度を引き上げる基準を設けます。"""


# Chapter structure. The divider pages and the agenda page are generated from this table, so the
# page numbers always agree with the final slide order. Page titles must match the slide titles.
T_APP_DEMO = "任意の実演：業務アプリで条件を変える"
T_APP_FDEMO = "任意の実演：予測にイベントを仮定して比べる"
T_APP_QA = "業務アプリの想定質問と答え方"
T_APP_FQA = "予測と将来の仮定の想定質問と答え方"

CHAPTERS = [
    {"n": 1, "name": "全体像と準備",
     "agenda": "10分の進行、企業、環境、前提、確認事項",
     "summary": "デモの進行、想定企業、環境と、\n実演前に確認する事項を示します。",
     "when": "実演の前",
     "note": "【進め方】最初に、10分の進行を確認します。続いて、デモが想定する企業と場面、Fabricの環境の構成を押さえます。さらに、8社を同じ尺度で比べる前提（8社、JPY、対象年度）と、「実演前の確認事項」（環境、製品の状態、日本語の扱い、リンク）を点検します。\n"
             "【所要】読むだけなら8分ほどです。確認事項は、実演の前日までに済ませます。\n"
             "【つなぎ】前提がそろったら、経営指標になじみのない相手に説明する場合は第2章、すぐ実演する場合は第3章へ進みます。",
     "pages": [
         ("10分デモの進行と画面切替", "10分の進行表"),
         ("想定企業：架空の総合商社グループ", "想定企業と場面"),
         ("デモ環境：データから画面までの構成", "デモ環境"),
         ("架空の8社を同じ尺度で比較", "デモの前提（8社、通貨、年度）"),
         ("実演前の確認事項", "環境と制約の確認"),
     ]},
    {"n": 2, "name": "営業向け基礎",
     "agenda": "EBITDA、baselineとactual、51%所有、通貨と現金",
     "summary": "経営指標になじみのない相手に、\nEBITDAなどの意味を\n数値例で説明します。",
     "when": "初めて説明する前",
     "when_long": "経営指標になじみのない相手に説明する前。事前学習用で、実演の10分には含めません",
     "note": "【進め方】EBITDA、baselineとactual、51%所有、通貨と現金の4ページです。各ページの数値例は説明用で、8社デモの集計結果ではありません。\n"
             "【所要】読むのに10分ほどです。実演の10分には含めず、事前学習や、相手に合わせた補足に使います。\n"
             "【つなぎ】相手が経営指標に詳しければ、この章は飛ばして第3章へ進みます。",
     "pages": [
         ("売上とEBITDAの違い", "1 EBITDA（金額は説明用の例）"),
         ("食品損失によるEBITDAの変化", "2 baselineとactual"),
         ("51%所有の金額を区別", "3 全額連結と持分管理"),
         ("通貨と現金の読み違い", "4 通貨とキャッシュフロー"),
     ]},
    {"n": 3, "name": "実演",
     "agenda": "4つの入力、結果と図、金額の区別、アプリ2種",
     "summary": "4つの入力をDAに順に入れ、\n結果の読み方、食品損失の経路、\n金額の区別を示します。\n最後に、任意で業務アプリを見せます。",
     "when": "ライブの10分（④は任意）",
     "when_long": "ライブの10分（Step04は時間があるときだけ）\n業務アプリは、10分のあとの追加",
     "note": "【進め方】4つの入力を、DAに順に入れます。必須はStep01からStep03で、Step04は時間があるときだけ実行します。各入力のページに全文があり、次のページに参考結果と読み方があります。\n"
             "【所要】ライブは10分です。ページは、説明者が実演の前に読む順に並んでいます。業務アプリの実演（最後の2ページ。過去の差額と、将来の予測）は、10分に含めず、求められたときに使う追加の時間です（過去が5分、将来が3分）。\n"
             "【注意】参考結果のページは、保存データから作った照合用です。ライブの回答とは呼びません。\n"
             "【つなぎ】実演のあとに質問が出たら、第4章へ進みます。",
     "pages": [
         ("4つの入力の使い分け", "進め方"),
         ("入力① 年間と月次の稼ぎ", "Step01 全社の稼ぎ（DAに入力）"),
         ("Step01は基準値と実績の比較", "Step01 全社の稼ぎ（参考結果）"),
         ("入力② 増減の理由を比較", "Step02 イベントの比較（DAに入力）"),
         ("Step02は10イベントの増減", "Step02 イベントの比較（参考結果）"),
         ("DAの金額を照合する正解表", "Step02 イベントの比較（照合用）"),
         ("入力③ 食品損失の根拠", "Step03 食品損失の根拠（DAに入力）"),
         ("食品損失の回答を読む", "Step03 食品損失の根拠（回答の読み方）"),
         ("Step03で食品損失の経路を表示", "Step03 食品損失の根拠（Notebookで表示）"),
         ("Notebookは保存済みの結果から始める", "Step03 食品損失の根拠（Notebookの操作）"),
         ("矢印は3種類の意味を持つ", "Step03 食品損失の根拠（図の読み方）"),
         ("代表記録でわかること、取れる行動", "Step03 食品損失の根拠（記録の読み方）"),
         ("4つの金額は範囲が違う", "Step03 食品損失の根拠（金額の読み方）"),
         ("入力④ 影響点の特定と対応案", "Step04 影響点と対応案（DAに入力）"),
         ("影響が大きい契約と繰り返し影響を受けた契約", "Step04 影響点と対応案（参考結果）"),
         ("優先度別の対応案（DA実行例）", "Step04 影響点と対応案（DAの実行例）"),
         ("デモで見せた3つの価値", "締め"),
         (T_APP_DEMO, "業務アプリ（任意の実演）"),
         (T_APP_FDEMO, "業務アプリ（任意の実演：将来）"),
     ]},
    {"n": 4, "name": "想定質問と答え方",
     "agenda": "短い回答、SQL、全社波及、カジュアル、アプリ",
     "summary": "質疑でよく出る質問への短い答え、\nOntologyの価値、全社波及の扱い、\nカジュアルな聞き方、業務アプリと\n予測の質問です。",
     "when": "質疑のとき",
     "note": "【進め方】質疑でよく出る質問に、短く答えるための章です。SQLとの違いは、Ontologyの価値を問われたときに、1ページを使って説明します。全社波及の2ページは、「全社に及ぶ出来事も扱えるか」と聞かれたときに使います。カジュアルな質問の2ページは、質問の自由度を聞かれたときに使います。業務アプリの質問の1ページと、予測と将来の仮定の質問の1ページは、第3章の最後の2ページで画面を見せたあとに使います。\n"
             "【所要】必要なページだけを選んで使います。\n"
             "【つなぎ】お客様のデータで進めたい場合は第5章、技術の詳細を聞かれたら第6章へ進みます。",
     "pages": [
         ("想定質問への短い回答", "短い回答集"),
         ("SQLとOntologyの違い", "価値の説明"),
         ("全社に及ぶイベントの影響範囲", "全社波及（影響範囲）"),
         ("全社波及で扱える範囲と未対応の範囲", "全社波及（扱える範囲）"),
         ("カジュアルな質問に答えさせる", "カジュアルな質問（結果と設定）"),
         ("カジュアルに聞くときの使い方と限界", "カジュアルな質問（使い方と限界）"),
         (T_APP_QA, "業務アプリの質問"),
         (T_APP_FQA, "予測と将来の仮定の質問"),
     ]},
    {"n": 5, "name": "お客様データへの当てはめ",
     "agenda": "イベントの集め方、台帳の列、小さく始める",
     "summary": "お客様のデータで同じことを\nするときの、イベントの集め方、\n台帳の列、進め方です。",
     "when": "お客様と次の一歩を話すとき",
     "note": "【進め方】お客様から「そんなデータは社内にない」と言われたときの進め方です。イベントの拾い先、台帳に持たせる列、最小の構成からの始め方を説明します。\n"
             "【所要】10分から15分の個別の打ち合わせを想定します。\n"
             "【つなぎ】どの実体が既存のデータで、どの実体が整備の対象かは、第6章の「Ontologyの全体マップ」にあります。",
     "pages": [
         ("社内にないイベントを集める方法", "1 イベントを集める"),
         ("イベントの列の例", "2 台帳の列（イベント）"),
         ("影響額と根拠の列の例", "3 台帳の列（影響額と根拠）"),
         ("小さく始める：最小のOntologyから広げる", "4 小さく始める"),
     ]},
    {"n": 6, "name": "構成の解説",
     "agenda": "Lakehouse、データの流れ、Ontology、設計の理由",
     "summary": "どのデータを、どう分け、\nどう結んで作ったかを\n説明します。",
     "when": "技術の質問が出たとき",
     "note": "【進め方】データをどう分け、どう結んで作ったかを説明する章です。技術に詳しい相手や、実装の質問が出たときに使います。営業向けの説明では、各ページの見出しだけで足ります。\n"
             "【所要】ページごとに60秒から90秒です。\n"
             "【つなぎ】新しい機能について聞かれたら、第7章へ進みます。",
     "pages": [
         ("9つのLakehouseの役割分担", "1 Lakehouse"),
         ("ゼロコピーと実体化の使い分け", "2 データの流れ"),
         ("4つの用語を食品の例で区別", "3 用語の整理"),
         ("Ontologyの全体マップ", "4 全体マップと整備状況"),
         ("実体と元の表の対応", "5 実体と表の対応"),
         ("39種類の関係と読み方", "6 関係"),
         ("この構成にした3つの理由", "7 設計の理由"),
         ("関係はGraph、金額はDeltaで照合", "8 GraphとDeltaの役割分担"),
     ]},
    {"n": 7, "name": "Ontologyの新機能",
     "agenda": "名前空間、ルール、共有プロパティ、継承",
     "summary": "新しいOntologyの機能と、\n入れる順番です。\nいずれもPreviewです。",
     "when": "今後の拡張を聞かれたとき",
     "note": "【進め方】新しいOntologyの機能と、このデモに入れる順番を説明します。いずれもPreviewで、このデモのOntologyには、まだ入れていません。\n"
             "【所要】全体像の1ページと、聞かれた機能の1ページで足ります。\n"
             "【注意】現在の製品の挙動は、実演の前に画面で確認します。",
     "pages": [
         ("新しいOntologyで増えた機能", "1 全体像"),
         ("名前空間：同じ名前を領域ごとに区別する", "2 名前空間"),
         ("ビジネスルール：業務の決まりを文章で定義する", "3 ビジネスルール"),
         ("共有プロパティ：同じ定義を1回だけ書く", "4 共有プロパティ"),
         ("継承：共通の親から実体を派生させる", "5 継承"),
         ("新機能を入れる順番と注意点", "6 導入の順番"),
     ]},
]
FRONT_TITLES = ["Trading Ontology デモ説明者ガイド", "資料の使い方と章立て"]
KICKER = {}
FINAL_TITLES = list(FRONT_TITLES)
for _ch in CHAPTERS:
    FINAL_TITLES.append(f"第{_ch['n']}章 {_ch['name']}")
    for _title, _suffix in _ch["pages"]:
        KICKER[_title] = f"第{_ch['n']}章 {_ch['name']}  |  {_suffix}"
        FINAL_TITLES.append(_title)
assert len(FINAL_TITLES) == len(set(FINAL_TITLES))
PAGE_OF = {t: i + 1 for i, t in enumerate(FINAL_TITLES)}


def page_range(ch):
    return f"p.{PAGE_OF[ch['pages'][0][0]]}〜{PAGE_OF[ch['pages'][-1][0]]}"


# Titles and page references for the group-wide ripple pages (page numbers come from the final order)
T_RIPPLE = "全社に及ぶイベントの影響範囲"
T_RIPPLE_SCOPE = "全社波及で扱える範囲と未対応の範囲"
RIPPLE_PAGES = f"p.{PAGE_OF[T_RIPPLE]}〜{PAGE_OF[T_RIPPLE_SCOPE]}"
CASUAL_PAGES = f"p.{PAGE_OF['カジュアルな質問に答えさせる']}〜{PAGE_OF['カジュアルに聞くときの使い方と限界']}"


def est_lines(value, width, size, ratio=.9):
    """Rough line count of Japanese text in a box of the given width (inches) and font size (pt)."""
    per_line = max(1, int(width / (size / 72) * ratio))
    lines = 0
    for para in value.split("\n"):
        w = sum(.5 if ord(c) < 128 else 1 for c in para)
        lines += max(1, -(-int(w + .999) // per_line))
    return lines


def stack_items(slide, x, y, w, items, accent, size=11.5, gap=.13, label_h=.22):
    """Stack (label, body) pairs: a bold label line, then wrapped body text. Returns the end y."""
    for label, body in items:
        text(slide, label, x, y, w, label_h, size, accent, True)
        y += label_h + .02
        lines = est_lines(body, w, size)
        text(slide, body, x, y, w, lines * size * 1.3 / 72, size, INK)
        y += lines * size * 1.3 / 72 + gap
    return y


def grid_table(slide, x, y, widths, header, rows, row_h, size=11.5, head_size=12,
               colors=None, bolds=None, fills=None, pad=.12, cell_colors=None):
    """Draw a table as rounded row cards. Returns the y below the last row."""
    cx = x
    for w, label in zip(widths, header):
        text(slide, label, cx + pad, y, w - 2 * pad, .26, head_size, TEAL, True)
        cx += w
    ry = y + .33
    for i, row in enumerate(rows):
        h = row_h[i] if isinstance(row_h, (list, tuple)) else row_h
        box(slide, x, ry, sum(widths), h - .07, fills[i] if fills else WHITE, rounded=True)
        cx = x
        for j, (w, val) in enumerate(zip(widths, row)):
            color = (cell_colors or {}).get((i, j), colors[j] if colors else INK)
            text(slide, val, cx + pad, ry + .07, w - 2 * pad, h - .16, size,
                 color, bolds[j] if bolds else False)
            cx += w
        ry += h
    return ry

def box(slide, x, y, w, h, fill=WHITE, line=None, rounded=False, flat=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    shp.line.fill.background() if line is None else None
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
    if rounded:
        shp.adjustments[0] = 0.08
    if flat:
        shp.shadow.inherit = False
    return shp


def text(slide, value, x, y, w, h, size=19, color=INK, bold=False,
         align=None, fill=None):
    if fill:
        box(slide, x, y, w, h, fill)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        p.line_spacing = Pt(round(size * 1.25, 1))
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        if align is not None:
            p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        rPr = r._r.get_or_add_rPr()
        ea = OxmlElement("a:ea")
        ea.set("typeface", FONT)
        rPr.append(ea)
        rPr.set("lang", "ja-JP")
    return shape


def arrow(slide, x1, y1, x2, y2, color=TEAL, dashed=False):
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(2.2)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    c.line._get_or_add_ln().append(tail)
    return c


def pill(slide, label, x, y, w, url=None, color=TEAL):
    shp = box(slide, x, y, w, .44, color, rounded=True)
    if url:
        shp.click_action.hyperlink.address = url
    t = text(slide, label, x + .13, y + .1, w - .2, .32, 13, WHITE, True)
    if url:
        t.click_action.hyperlink.address = url


def card(slide, x, y, w, h, title, body, accent=TEAL, number=None):
    box(slide, x, y, w, h, WHITE, rounded=True)
    if number:
        text(slide, number, x + .22, y + .18, .65, .52, 30, accent, True)
        tx = x + .98
        tw = w - 1.2
    else:
        tx, tw = x + .25, w - .5
    text(slide, title, tx, y + .22, tw, .48, 22, accent, True)
    text(slide, body, x + .25, y + .95, w - .5, h - 1.03, 18)


def new_slide(title, lead, section, notes, sources, takeaway=None, foot=None):
    section = KICKER.get(title, section)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    n = len(prs.slides)
    text(slide, section, .6, .32, 11.4, .25, 11, TEAL, True)
    text(slide, title, .6, .8, 12.05, .66, 32, NAVY, True)
    text(slide, lead, .62, 1.64, 12.02, .88, 18, MUTED)
    if takeaway:
        box(slide, .6, 6.54, 12.12, .43, PALE, rounded=True)
        text(slide, takeaway, .78, 6.62, 11.78, .29, 13, NAVY, True)
    text(slide, foot or "架空の業務データ／公開ECB為替を使用  |  検証済み構成：2026-10-27",
         .62, 7.12, 11.3, .22, 9, MUTED)
    text(slide, f"{n:02}", 12.17, 7.06, .55, .3, 13, TEAL, True)
    source_text = "\n".join(sources)
    slide.notes_slide.notes_text_frame.text = (
        notes.strip() + "\n\n【出典と照合先】\n" + source_text)
    slide_metadata.append({"slide": n, "title": title, "sources": sources})
    return slide


def dashed_frame(slide, x, y, w, h, color=PURPLE):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = .05
    shp.fill.background()
    shp.shadow.inherit = False
    shp.line.color.rgb = RGBColor.from_string(color)
    shp.line.width = Pt(1.75)
    shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shp


def nodes(slide, labels, y=3.2, xs=None, w=2.7, h=1.5, edge_colors=None):
    xs = xs or [.65, 3.8, 6.95, 10.1]
    for i, (title, body, color) in enumerate(labels):
        x = xs[i]
        box(slide, x, y, w, h, WHITE, rounded=True)
        text(slide, title, x + .18, y + .2, w - .36, .42, 20, color, True)
        text(slide, body, x + .18, y + .76, w - .36, .64, 17)
        if i:
            arrow(slide, xs[i - 1] + w + .03, y + h / 2,
                  x - .06, y + h / 2,
                  edge_colors[i - 1] if edge_colors else color)


def prompt_slide(title, question, purpose, expected, sources, size=16,
                 result_note=None, notes_extra=""):
    slide = new_slide(
        title,
        purpose,
        "DAへの入力  |  枠内の全文をコピー",
        "【操作】枠内の文章をすべてコピーし、公開済みFabric DAへ入力します。\n"
        "【入力全文】\n" + question +
        "\n【回答の見方】" + expected +
        "\n【保存結果との区別】" + (result_note or (
            "次ページは保存データから作った照合用の参考結果です。"
            "今回この質問を再実行した回答ではありません。文章やグラフの見た目は変わり得ます。")) + "\n"
        "【失敗時】対象期間、指標、通貨、件数を確かめます。エラーや欠落を隠して、"
        "参考結果をライブの回答として見せないでください。DAの結果を待つ間は、"
        "質問で何を確かめたいかを説明します。\n"
        "【説明者向け】英文の列名や長いIDを読み上げる必要はありません。"
        "入力する条件を示した後、結果の金額と業務上の意味に注目してもらいます。"
        + notes_extra,
        sources, "次ページ：この入力に対応する参考結果と、営業向けの読み方")
    box(slide, .65, 2.76, 12.02, 3.05, WHITE, rounded=True)
    text(slide, "入力するテキスト（全文）", .9, 2.95, 11.5, .38, 18, TEAL, True)
    text(slide, question, .9, 3.5, 11.5, 2.25, size)
    box(slide, .65, 5.98, 12.02, .43, PALE, rounded=True)
    text(slide, expected, .9, 6.06, 11.5, .3, 15, NAVY)
    pill(slide, "Fabric DAを開く", 10.1, .3, 2.55, DA)
    return slide


# 01
s = new_slide(
    "Trading Ontology デモ説明者ガイド",
    "損益グラフだけでは、食品の損失がどの記録や取引に結び付くかを説明できません。\n金額の比較と、根拠をたどる操作を分けて案内します。",
    "説明者ガイド  |  営業とSE向け",
    """【目安】40秒。これは説明者の準備資料です。全ページをそのまま10分で投影する構成ではありません。基礎解説は事前学習用で、ライブ10分枠とは別です。
【話す】「架空の総合商社を使い、まず損益の変化を見ます。その後、食品損失を一つ選んで、根拠記録や契約と結び付けて確認します。今日見せるのは、数字と説明を同じデータに照らす方法です。」
【指す】3つのカードを順に指し、「概況、イベント、根拠」と読みます。Ontology、GQL、DAGという名前はここでは暗記させません。
【操作】事前に「実演前の確認事項」のリンクで公開済みFabric Data AgentとStep03 Notebookを別タブで開きます。
【避ける】実在企業の実データ、実世界の因果分析、自動の原因究明、正式な連結会計の完成版とは言いません。
【つなぎ】「この順序なら、画面の移動も一度で済みます。」次のページで、この資料の章立てを確認します。""",
    ["config/deployment_manifest.json", "config/step3_deployment.json"],
    "本番前：DAの公開済み画面とNotebookの保存済みHTMLを別タブで確認")
card(s, .6, 2.85, 3.83, 2.96, "全社の概況", "月次の基準値と実績\n年間の全額／持分管理", number="1")
card(s, 4.75, 2.85, 3.83, 2.96, "10イベントの比較", "符号付きの増減\n食品損失を選択", accent=CORAL, number="2")
card(s, 8.9, 2.85, 3.83, 2.96, "食品損失の根拠", "実体の関係をDAGで表示\n記録、活動、契約、KPI", number="3")

# 01b agenda: page numbers come from the final order
s = new_slide(
    "資料の使い方と章立て",
    "この資料は、デモを説明する人の準備用です。全ページを10分で見せる構成ではありません。\n"
    "場面によって使う章が違うので、最初に章立てで、自分が使う章を選びます。",
    "はじめに  |  資料の使い方",
    """【目安】45秒。最初に、この資料の使い方を伝えます。
【話す】「この資料は、デモを説明する人の準備用です。全ページを10分で見せるのではありません。実演で使うのは第3章で、前後の章は、準備と、質問への答え方です。」
【指す】7つの章を上から順に指します。「使う場面」の列を見て、自分がいま使う章を選びます。
【読む順】初めて読むときは、第1章から第3章までを順に読みます。経営指標になじみのない相手に説明する場合は、第2章を先に読みます。質疑の準備は第4章から、技術の質問への備えは第6章からです。
【ページ番号】PowerPointのスライド番号です。各章の最初のページ（章扉）に、その章のページ一覧があります。
【時間】実演の10分に使うのは第3章です。第1章は実演の前に、第2章は初めて説明する前に読みます。第4章から第7章は、質問が出たときに、必要なページだけを使います。""",
    ["この資料の構成（presentation/build_presenter_deck.py の CHAPTERS）"],
    "初めて読むときは第1章から第3章まで。質疑の準備は第4章から")
text(s, "章", .8, 2.56, .9, .26, 11, TEAL, True)
text(s, "内容", 5.1, 2.56, 3.0, .26, 11, TEAL, True)
text(s, "使う場面", 9.05, 2.56, 2.2, .26, 11, TEAL, True)
text(s, "ページ", 11.5, 2.56, 1.05, .26, 11, TEAL, True, PP_ALIGN.RIGHT)
for i, ch in enumerate(CHAPTERS):
    y = 2.86 + i * .52
    box(s, .62, y, 12.08, .46, PALE if ch["n"] == 3 else WHITE, rounded=True)
    text(s, f"第{ch['n']}章", .8, y + .08, .95, .3, 14, TEAL, True)
    text(s, ch["name"], 1.75, y + .08, 3.25, .3, 14, NAVY, True)
    text(s, ch["agenda"], 5.1, y + .1, 3.85, .3, 11.5, INK)
    text(s, ch["when"], 9.05, y + .1, 2.4, .3, 11, MUTED)
    text(s, page_range(ch), 11.5, y + .1, 1.05, .3, 12.5, TEAL, True, PP_ALIGN.RIGHT)


# 02
s = new_slide(
    "10分デモの進行と画面切替",
    "自然言語の応答と、関係図の表示は別の画面です。\nDAで金額を確かめてから、食品の実体経路をNotebookで示します。",
    "進行表  |  説明者が手元で参照",
    """【目安】準備時に60秒。下記の時間はライブデモ全体の割当です。
【0:00～1:00】架空データ、対象年度、全額と持分管理を一言で説明します。「想定企業：架空の総合商社グループ」と「架空の8社を同じ尺度で比較」のページを、それぞれ20秒ほどで使います。「デモ環境」のページは、聞かれたときに使います。
【1:00～3:00】公開済みFabric DAにQ1。月次のbaselineとactualの線を指し、年間actual 447.57億円を確認します。年度を本日からの直近1年に置き換えません。
【3:00～5:00】Q2。符号付き棒グラフを指し、横断2件を含めて10件あること、食品損失が約-1.14億円であることを確認します。色や横向きの指定で失敗した場合は、それらを外します。
【5:00～9:00】Q3で食品の要約とIDを確認し、Step03へ一度だけ切り替えます。先頭の保存済みHTML、カード、8ノードのDAG、根拠の原文を順に示します。
【9:00～10:00】関係の共有と金額照合の価値を説明します。余裕がなければ追加質問は実行しません。
【Step04】余裕があれば、DAでQ4（影響点の特定と対応案）を実行します。Q4はNotebookの前後どちらでも実行できます。画面の切替を減らしたい場合は、Q3の直後にDAでQ4を実行してからNotebookへ移ります。
【待ち時間】DAの処理中は「金額を先に確認し、関係は次の画面で確認します」と説明します。無応答時は既存の保存結果／この資料に切り替え、ライブ応答が成功したとは言いません。
【操作】質問全文は「入力①～④」の各スライドの枠内からコピーします。ノートにも全文があります。基礎解説は事前学習用で、10分の実演枠には含めません。
【つなぎ】「最初に、このデモが想定する企業と場面を説明します。」""",
    ["prompts/demo_questions_ja.txt", "reports/step3_ui_verification.json"],
    "必須はStep01～03。Step04（影響点と対応案）、ENE_POS、横断比較は時間が余った場合のみ")
rows = [
    ("0～1分", "前提", "架空の8社、対象年度、金額の定義"),
    ("1～3分", "Step01：DA", "月次のbaseline／actual"),
    ("3～5分", "Step02：DA", "全10イベントの符号付き影響"),
    ("5～9分", "Step03：Notebook", "食品の要約 → 保存済みDAG → 根拠"),
    ("9～10分", "締め", "共通の意味、追跡、正確な金額照合"),
]
for i, (tm, screen, purpose) in enumerate(rows):
    y = 2.73 + i * .67
    box(s, .62, y, 12.08, .55, WHITE if i != 3 else PALE, rounded=True)
    text(s, tm, .83, y + .12, 1.3, .3, 17, TEAL, True)
    text(s, screen, 2.25, y + .12, 3.35, .3, 17, INK, True)
    text(s, purpose, 5.72, y + .12, 6.64, .3, 17)

# Company and scene, then the demo environment. Both pages sit between the 10-minute flow and the
# same-scale page. Every figure is computed from the repo artifacts and asserted.
T_COMPANY = "想定企業：架空の総合商社グループ"
T_ENV = "デモ環境：データから画面までの構成"
subsidiaries = json.loads((PROJECT / "config/subsidiaries.json").read_text())
SECTOR_JP = {"metals": "金属", "energy": "エネルギー", "foods": "食品", "finance": "金融"}
COMMODITY_JP = {"copper_alloy": "銅合金", "LNG": "LNG", "milling_wheat": "製粉用小麦",
                "trade_finance": "貿易金融"}
commodity_of = {c["entity_id"]: c["commodity"] for c in dataset["semantic_contracts"]}
assert len(subsidiaries) == 8 and len(set(commodity_of.values())) == 4
OKU = Decimal(10) ** 8
actual_total = defaultdict(Decimal)
for r in dataset["semantic_financial_metrics_daily"]:
    if r["scenario_tag"] == "actual":
        actual_total[r["metric_code"]] += Decimal(r["amount_group"])
assert actual_total["EBITDA"] == Decimal(manifest["annual_actual_jpy"]["ebitda_full_consolidated"])
revenue_oku = round(actual_total["REVENUE"] / OKU)
ebitda_oku = round(actual_total["EBITDA"] / OKU)
assert (revenue_oku, ebitda_oku) == (3111, 448)
own_of = {s["entity_id"]: s["ownership_ratio"] for s in subsidiaries}
food_loss_entities = {r["entity_id"] for r in dataset["semantic_event_financial_effects"]
                      if r["event_id"] == "FOD_NEG"}
assert food_loss_entities == {"SUB_FOD_JP_01", "SUB_FOD_US_01"}
assert (own_of["SUB_FOD_JP_01"], own_of["SUB_FOD_US_01"]) == ("0.51", "1.00")
assert sorted(own_of.values()) == ["0.51"] * 4 + ["1.00"] * 4

s = new_slide(
    T_COMPANY,
    "月次の経営会議で、全社の稼ぎが基準から動いた理由を問われる場面を想定します。\n"
    "事業も通貨も所有比率も違う8社の数字を、同じ定義で説明できることが前提です。",
    "",
    """【目安】準備時に60秒。ライブでは、表と右のカードを指して20秒ほどで話します（10分の進行の「0～1分」）。
【話す】「今日のデモは、架空の総合商社グループです。金属、エネルギー、食品、金融の4つの事業があり、それぞれに日本の会社と米国の会社があります。全部で8社です。月次の経営会議で、全社の稼ぎが基準から動いた理由を問われた場面を想定しています。」
【背景】経営会議では、全社の稼ぎ（EBITDA）が基準値から動くと、「なぜか」「どの契約が効いたか」「根拠の記録は何か」が問われます。事業、通貨、所有比率が違う会社の数字は、定義をそろえないと比べられません。根拠の確認にも、詳しい人が表をつなぎ直す手間がかかることが多い、という場面です。このデモは、その場面で、金額、つながり、根拠を同じデータから示します。
【指す】表を上から順に読みます。事業ごとに、日本の会社、米国の会社の順です。次に、右上の「親会社」で、全社の管理データと、全額連結の扱いを伝えます。最後に、右下の「年間の規模」を指します。
【所有比率】100%所有は4社（金属の日本、エネルギーの日本、食品の米国、金融の日本）、51%所有も4社です。国では決めていません。全8社は支配子会社なので、全額連結には100%を取り込み、所有比率を掛けた持分管理寄与は別の指標として持ちます。違いは「51%所有の金額を区別」のページで説明します。
【規模】年間（2025-10-01～2026-09-30）の実績は、売上が約3,111億円、EBITDAが約448億円です。いずれも全額連結のJPYです。EBITDAはQ1の結果（447.57億円）と同じ値です。売上は保存データからの集計で、DAに聞いて確認した値ではありません。
【取り扱い】銅合金、LNG、製粉用小麦、貿易金融は、このデモの設計上の取り扱いです。実際の総合商社の事業構成や規模を表すものではありません。
【食品損失との関係】実演で追う食品損失（FOD_NEG）は、持分51%の日本の食品会社と、持分100%の米国の食品会社の2社に効きます（各42件、合計84件の影響明細）。そのため、全額（−1.1446億円）と持分管理（−0.8763億円）は、異なる値になります。
【避ける】社名は架空で、実在の会社を指しません。「実際の総合商社の規模」「業界の標準的な利益率」とは言いません。
【つなぎ】「次に、このデータが、Fabricのどこで動いているかを見ます。」""",
    ["config/subsidiaries.json", "data/demo_dataset.json:semantic_financial_metrics_daily（actual）",
     "data/demo_dataset.json:semantic_event_financial_effects（FOD_NEG）",
     "config/deployment_manifest.json:annual_actual_jpy"],
    "実演で追う食品損失は、持分51%の日本の会社と、持分100%の米国の会社に影響する")
company_rows, company_fills = [], []
for i, sub in enumerate(subsidiaries):
    share = f"{int(Decimal(sub['ownership_ratio']) * 100)}%"
    company_rows.append([
        SECTOR_JP[sub["sector_code"]], sub["entity_name"], COMMODITY_JP[commodity_of[sub["entity_id"]]],
        f"{'日本' if sub['country_code'] == 'JP' else '米国'}／{sub['currency_code']}", share])
    company_fills.append(WHITE if (i // 2) % 2 == 0 else PALE)
grid_table(s, .62, 2.58, [1.2, 2.9, 1.4, 1.25, .95],
           ["事業", "会社名", "取り扱い", "国／通貨", "所有比率"], company_rows, .43, size=12,
           colors=[TEAL, NAVY, INK, INK, INK], bolds=[True, True, False, False, True],
           fills=company_fills)
box(s, 8.52, 2.58, 4.18, 1.75, WHITE, rounded=True)
text(s, "親会社：全社の管理", 8.77, 2.72, 3.7, .3, 16, TEAL, True)
text(s, "SAP風の管理データを持つ。会社マスタ、為替、仕訳、内部取引、連結消去。\n"
        "全8社を支配子会社として全額連結する。\n所有比率を掛けた持分管理寄与は、別の指標。",
     8.77, 3.13, 3.7, 1.15, 12, INK)
box(s, 8.52, 4.43, 4.18, 1.85, WHITE, rounded=True)
text(s, "年間の規模（実績）", 8.77, 4.57, 3.7, .3, 16, TEAL, True)
text(s, f"売上　約{revenue_oku:,}億円", 8.77, 5.0, 3.7, .32, 17, NAVY, True)
text(s, f"EBITDA　約{ebitda_oku:,}億円", 8.77, 5.42, 3.7, .32, 17, NAVY, True)
text(s, "全額連結のJPY。2025-10-01から2026-09-30", 8.77, 5.85, 3.7, .4, 11, MUTED)

env_counts = {
    "contracts": len(dataset["semantic_contracts"]),
    "activities": len(dataset["semantic_business_activities"]),
    "events": len(dataset["semantic_risk_events"]),
    "evidence": len(dataset["semantic_event_evidence"]),
    "effects": len(dataset["semantic_event_financial_effects"]),
}
assert env_counts == {"contracts": 96, "activities": 8760, "events": 10, "evidence": 64, "effects": 1344}
assert (manifest["workspace"]["name"], manifest["workspace"]["capacity_sku"]) == ("Trading", "F16")
assert (manifest["data_agent"]["selected_tables"], manifest["data_agent"]["runtime"]) == (25, "Preview")
assert (manifest["storage"]["lakehouses"], manifest["storage"]["zero_copy_shortcuts"]) == (9, 8)
_lh_tables = defaultdict(int)
for _key in cloud["counts"]:
    _lh_tables[_key.split(".")[0]] += 1
lh_tables_total, lh_rows_total, lh_tables_parent = cloud["tables"], cloud["rows"], _lh_tables["LH_Group_SAP"]
assert (lh_tables_total, lh_rows_total, lh_tables_parent) == (95, 90424, 17)

s = new_slide(
    T_ENV,
    "実演中は、「このデータはどこにあるのか」「何が答えを出しているのか」と聞かれます。\n"
    "答えの出どころを、データ、意味と関係、画面の3層で言い分けられるようにします。",
    "",
    """【目安】準備時に60秒。ライブでは使わず、「どこで動いているのか」「何が答えを出しているのか」と聞かれたときに使います。
【話す】「このデモは、Fabricの1つのワークスペースで動いています。データは9つのLakehouseに置き、その上にOntologyとGraphを作ります。Data Agentが日本語の質問に答え、食品損失の経路はNotebookで図にします。」
【指す】下の段から上へ、データ、意味と関係、画面の順に指します。実演で開くのは、一番上の2つの画面です。
【データ】9つのLakehouse（親1、子会社8）に、物理表が95、行数が90,424あります。親は17表とショートカット8本で、SAP風の管理データ、全社の分析用の表、全社KPIを置いています。子会社は78表で、各社の原本です。詳細は第6章の「9つのLakehouseの役割分担」です。
【意味と関係】Ontologyは、実体の種類が17、関係の種類が39です。データは持たず、表に対応付けます。Graphは、Ontologyの定義から作った、実体と関係のグラフです。実体の1行が1つのノード（点）、関係の1行が1本の矢印（線）になり、ノードが33,960、矢印が65,684本あります。
【画面】Data Agentは公開済みで、ランタイムはPreviewです。Ontologyと親Lakehouseの表（25本）を参照し、日本語の質問に、金額と説明で答えます。Notebook（Step03）は、Graphで関係をたどり、金額をDeltaの表で照合して、食品損失の経路を図にします。実演の入口は、公開済みのFabric DAの画面です。Microsoft 365 Copilotには公開していません。
【規模】期間は365日（2025-10-01～2026-09-30）です。契約は96件（各社12件）、活動は8,760件（各社1日3件。金融の2社は手数料の計上）、イベントは10件、根拠は64件、効果は1,344件です。容量はF16、ワークスペースの名前はTradingです。
【避ける】「本番の基幹システムとつないでいる」とは言いません。データは架空で、為替だけが実データです。
【つなぎ】「次に、8社を同じ尺度で比べるための前提を確認します。」""",
    ["config/deployment_manifest.json", "reports/cloud_validation.json", "reports/graph_node_counts.json",
     "reports/graph_edge_counts.json", "docs/architecture.md", "config/lakehouse_inventory.json"],
    "容量（F2以上）、Previewの扱い、日本語の対応は、実演前の確認事項で点検する。構成の詳細は第6章")
env_layers = [
    ("画面", NAVY, [
        ("Data Agent", "公開済み、Preview",
         "Ontologyと親Lakehouseの表25本を参照し、\n日本語の質問に、金額と説明で答える"),
        ("Notebook（Step03）", "保存済みの図から開始",
         "Graphで関係をたどり、金額をDeltaの表で照合して、\n食品損失の経路を図にする")]),
    ("意味と関係", PURPLE, [
        ("Ontology", "17種類の実体、39種類の関係",
         "業務の言葉と、つながりを定義する。\nデータは持たず、表に対応付ける"),
        ("Graph", f"ノード{sum(node_counts.values()):,}、矢印{sum(edge_counts.values()):,}本",
         "Ontologyの定義から作った、実体と関係のグラフ。\n経路をたどる質問に使う")]),
    ("データ", TEAL, [
        ("親のLakehouse", "17表、ショートカット8本",
         "SAP風の管理データ、全社の分析用の表、全社KPI。\n業種拡張は、子会社の表を参照する"),
        ("子会社8社のLakehouse", f"{lh_tables_total - lh_tables_parent}表",
         "各社の原本。会社ごとに分けたまま残し、\n書き換えない")]),
]
for li, (layer, accent, boxes) in enumerate(env_layers):
    y = 2.58 + li * .96
    box(s, .62, y, 1.45, .88, accent, rounded=True)
    text(s, layer, .62, y + .26, 1.45, .35, 15, WHITE, True, PP_ALIGN.CENTER)
    for bi, (name, stat, desc) in enumerate(boxes):
        bx, bw = (2.2, 5.05) if bi == 0 else (7.4, 5.3)
        box(s, bx, y, bw, .88, WHITE, rounded=True)
        text(s, name, bx + .18, y + .07, bw - .36, .26, 13.5, NAVY, True)
        text(s, stat, bx + .18, y + .08, bw - .36, .26, 11.5, accent, True, PP_ALIGN.RIGHT)
        text(s, desc, bx + .18, y + .34, bw - .36, .46, 11, INK)
tiles = [
    ("365日", "2025-10-01から2026-09-30"),
    (f"契約{env_counts['contracts']}件", "各社12件"),
    (f"活動{env_counts['activities']:,}件", "受注や売上の明細\n各社1日3件"),
    (f"イベント{env_counts['events']}件", f"根拠{env_counts['evidence']}件、効果{env_counts['effects']:,}件\n効果は影響額の明細"),
    ("容量F16", "ワークスペース　Trading"),
]
for ti, (big, small) in enumerate(tiles):
    x = .62 + ti * 2.445
    box(s, x, 5.5, 2.3, .9, WHITE, rounded=True)
    text(s, big, x + .16, 5.58, 2.0, .3, 16, TEAL, True)
    text(s, small, x + .16, 5.93, 2.0, .42, 10.5, MUTED)

# 03
s = new_slide(
    "架空の8社を同じ尺度で比較",
    "国や事業が異なる会社のデータは、通貨や損益の定義がそのままでは揃いません。\nこのデモはJPYに換算し、全額と持分管理寄与を分けます。",
    "背景  |  最初に共有する業務設定",
    """【目安】準備時に45秒。ライブでは、通貨と所有比率だけを20秒ほどで話します（10分の進行の「0～1分」）。
【話す】「ここまでに紹介した8社を、同じ尺度で比べます。日本円とドルの会社は円にそろえ、所有比率が100%の会社と51%の会社は、全額と持分で分けて持ちます。各社の原本は別のLakehouseに残し、共通の分析用データを用意しています。」
【指す】左から、通貨、所有比率、対象年度の3つのカードを指します。会社の一覧は、「想定企業：架空の総合商社グループ」のページの表です。
【補足】100%所有は4社、51%所有も4社です。国だけで決めた割当ではありません。全8社は支配子会社という設定なので、51%でも全額ベースの管理デモ集計には100%を取り込みます。一方で持分管理寄与には所有比率を掛けます。法定の連結財務諸表や持分法損益とは区別します。
【実データ】業務、契約、根拠、イベント、価格は架空です。為替だけは公開されたECBの日次観測値を使用しています。USD/JPYはEUR建ての2通貨から計算し、休日は直前観測を最大5日補完します。
【為替の言い方】「為替だけは実データです。ECBが公開している日次の参考レートで、ドル建ての金額を円に換算しています。」
対象年度のUSD/JPYは約146.8円から約163.9円で、期初は147.11円、期末は157.00円です。
【避ける】「実際の市場や業績を反映したデモ」とは言いません。為替が実データでも、会社、取引、金額、イベントは架空です。
ECBの参考レートは公表用の参考値で、企業の決算で使う換算レートではありません。
【つなぎ】「次に、実演の前に確認しておく環境と、製品の状態を見ます。」""",
    ["docs/architecture.md", "config/subsidiaries.json", "data/ecb_fx_raw.csv"],
    "休日は直前の観測値を最大5日まで補完。ECBの参考レートは、決算で使う換算レートではない")
usd_jpy = [Decimal(r["fx_rate"]) for r in dataset["fx_rates"] if r["from_currency"] == "USD"]
assert len(usd_jpy) == 365 and (round(min(usd_jpy)), round(max(usd_jpy))) == (147, 164)
for i, (head, big, small) in enumerate([
    ("通貨", "JPY／USD → JPY",
     f"為替は実データ。\nECBが公開する日次の参考レートを使う。\n対象年度のUSD/JPYは、\n約{round(min(usd_jpy))}円から約{round(max(usd_jpy))}円。"),
    ("所有比率", "100%×4社／51%×4社",
     "国では決めていない。\n全額連結と、所有比率を掛けた\n持分管理寄与を分けて持つ。"),
    ("対象年度", "2025-10-01〜2026-09-30",
     "365日分。\n最新データ日は2026-09-30。"),
]):
    x = .62 + i * 4.09
    box(s, x, 3.0, 3.89, 2.85, WHITE, rounded=True)
    text(s, head, x + .25, 3.25, 3.4, .4, 22, TEAL, True)
    text(s, big, x + .25, 3.9, 3.45, .4, 19, NAVY, True)
    text(s, small, x + .25, 4.5, 3.4, 1.25, 13, INK)

# Sales foundations
foundation_notes = """
【使い方】経済や会計を専門にしない営業向けの事前学習ページです。基礎説明を含める場合は
10分枠とは別に時間を取ります。実演中にすべてのページを読む必要はありません。
【注意】ここに置く小さな数値例は概念を説明するための例であり、8社デモの集計結果ではありません。
顧客固有の会計定義を確認せずに、同じ指標名なら同じ計算方法だと説明しないでください。
【話す順】まず日常的な言葉で意味を伝え、次に数値例を指し、最後に今回のデモでの使い方を説明します。
用語の英語名や計算式の暗記は求めません。読者が結果の単位、対象、比較相手を説明できれば十分です。
【対象】今回の業務データは架空です。金額例や関係の説明を、実在企業の業績や実証済みの因果として
外部に説明しないでください。疑問が出た場合は、既存のデータ定義と保存結果に戻って確認します。
"""
s = new_slide(
    "売上とEBITDAの違い",
    "たくさん売れていても、仕入れや運営にお金がかかれば利益は残りません。\n"
    "事業の稼ぎを比べるため、このデモではEBITDA（イービットディーエー）を使います。",
    "営業向け基礎 1  |  金額は説明用の例",
    foundation_notes + """
【話す】「売上100万円だけでは、儲かっているかは分かりません。仕入れ60万円と、人件費など25万円を
差し引くと15万円が残ります。この簡単な例では15万円をEBITDAと考えます。」
【定義】EBITDAは利払い前、税引き前、減価償却前の利益です。一般には営業利益に減価償却費などを
足し戻す見方がありますが、実務では会社ごとの定義と調整項目を確認します。
【減価償却】設備の購入額を使用期間に分けて費用化する会計処理です。その年の現金支出とは一致しません。
【避ける】売上、最終利益、現金残高と同じだと説明しないでください。""",
    ["説明用の単純化した数値例", "docs/architecture.md"],
    "EBITDAは事業の稼ぎを見る指標。最終利益や手元の現金とは別")
nodes(s, [
    ("売上", "100万円\n販売の対価", TEAL),
    ("仕入れ", "−60万円\n仕入れ費用", CORAL),
    ("運営費", "−25万円\n人件費など", CORAL),
    ("EBITDA", "15万円\n事業の稼ぎ", TEAL),
], xs=[.65, 3.84, 7.03, 10.22], w=2.45, h=1.7)
text(s, "EBITDA：利息、税金、減価償却を差し引く前の利益。\n"
     "減価償却は、設備の購入額を使用期間に分けて費用にする会計処理。", .85, 5.35, 11.7, .85, 18, MUTED)

s = new_slide(
    "食品損失によるEBITDAの変化",
    "前ページと同じ例で、食品廃棄により費用が3万円増えた場合を考えます。\n"
    "同じ期間の「損失がない想定」と「損失を反映した結果」を比べます。",
    "営業向け基礎 2  |  前ページの15万円から続く説明用の例",
    foundation_notes + """
【前ページからのつなぎ】「先ほど、売上100万円から仕入れ60万円と運営費25万円を引くと、
EBITDAは15万円でした。では、その事業で食品廃棄が起き、費用が3万円増えたらどうでしょう。」
【話す】「売上と仕入れは変わらず、運営費が25万円から28万円になったとします。
稼ぎは15万円から12万円へ、3万円減ります。比較している金額は、どちらも同じ期間のEBITDAです。」
【用語の順】先に日本語で例を説明してから、「損失がない場合の15万円を基準値、baselineと呼びます。
損失を反映した12万円をこのデモ上の実績、actualと呼びます。両者の差はマイナス3万円、
これが差分、deltaです」と対応付けます。3つの英単語を別の経営指標として暗記させません。
【用語】baselineはこのデモでイベントを加える前の基準値であり、会社の正式な予算を意味しません。
actualはイベント反映後のデモ上の実績です。実在企業から収集した実績という意味ではありません。
【期間】前月と今月の比較ではありません。同じ期間について、イベントがない想定と反映後の値を比較します。
【符号】追加費用はプラス3万円ですが、稼ぎを減らすため、EBITDAへの影響はマイナス3万円です。
【二重計上】actualの12万円には損失がすでに反映されています。さらに3万円を引いて9万円にしません。
【実データとの区別】この食品廃棄の3万円は理解のための説明用で、FOD_NEGの実際のデモ配賦額ではありません。
【実演】年間や月次を比較する際もこの区別を保ちます。丸めと集計粒度で微差が出る場合は照合元に戻ります。""",
    ["前ページから継続する説明用の数値例",
     "data/demo_dataset.json:semantic_financial_metrics_daily", "prompts/da_instructions_ja.txt"],
    "baselineとactualはEBITDAの比較条件。actualの12万円から損失をもう一度引かない")
card(s, .65, 2.85, 3.78, 2.6, "損失なしの稼ぎ",
     "EBITDA 15万円\n基準値（baseline）\n前ページと同じ計算")
card(s, 4.77, 2.85, 3.78, 2.6, "食品廃棄の影響",
     "追加費用 ＋3万円\n稼ぎへの影響 −3万円\n差分（delta）", accent=CORAL)
card(s, 8.89, 2.85, 3.78, 2.6, "損失反映後の稼ぎ",
     "EBITDA 12万円\n実績（actual）\nこのデモ上の呼び方")
text(s, "基準15万円 ＋ 影響（−3万円）＝ 実績12万円",
     .85, 5.67, 11.7, .45, 25, NAVY, True)
text(s, "反映後の計算：売上100万円 − 仕入れ60万円 − 運営費28万円 ＝ EBITDA 12万円",
     .85, 6.17, 11.7, .28, 14, MUTED)

s = new_slide(
    "51%所有の金額を区別",
    "グループ全体の事業規模と、出資割合を掛けた管理上の貢献額は目的が異なります。\n"
    "このデモでは8社すべてを支配子会社として扱います。",
    "営業向け基礎 3  |  子会社EBITDAが1億円の場合",
    foundation_notes + """
【話す】「51%の株式を持つ子会社が1億円を稼いだ場合です。支配子会社として全額を見る指標は1億円。
所有比率を掛ける管理指標は5,100万円です。目的が違うので合計しません。」
【定義】連結は、企業グループを一つのまとまりとして見ることです。実務上の支配の判定は、
単に株式比率だけで決めるものではありません。このデモでは全社FULL_CONTROLという設定です。
【注意】持分管理寄与は簡易指標です。配当、親会社の最終利益、法定の持分法損益ではありません。
残る49%分との差をそのまま非支配株主利益と呼ぶことも避けます。
【全額連結の側も】全額連結EBITDAも、連結財務諸表の数値ではなくデモの管理値です（親会社の費用、税、未実現利益の消去は対象外）。EBITDAは会計基準が定める項目ではないため、「法定連結EBITDA」と言うより「連結財務諸表の数値とは別のデモ用の値」と説明するほうが正確です。
【デモの実数】各社の全額EBITDAに所有割合を掛けて合計すると335.48億円で、全額447.57億円との差は約112.1億円です。この差は、51%の4社のEBITDA（約228.8億円）の49%分に当たります。ただし非支配株主利益そのものではありません。""",
    ["config/subsidiaries.json", "reports/da_qa_published.json"],
    "2つの金額は同じ稼ぎの別の見方。足し合わせない")
card(s, .65, 2.85, 5.83, 2.8, "全額連結：事業全体を見る", "1億円 × 100% ＝ 1億円\n51%所有の子会社も全額を集計\n今回の管理デモ上の扱い")
card(s, 6.85, 2.85, 5.83, 2.8, "持分管理：出資割合で見る", "1億円 × 51% ＝ 5,100万円\n所有比率を反映した簡易的な貢献額\n配当や最終利益ではない", accent=PURPLE)
text(s, "「51%だから、全額連結にも51%を掛ける」という説明はしません。", .85, 5.9, 11.6, .45, 18, MUTED)

s = new_slide(
    "通貨と現金の読み違い",
    "日本円と米ドルをそのまま足したり、利益をそのまま現金と考えたりすると判断がずれます。\n"
    "金額を見るときは、通貨と指標の定義も確認します。",
    "営業向け基礎 4  |  レートは実データ、100万ドルは説明用の例",
    foundation_notes + """
【為替】「100万ドルの稼ぎは、1ドル156.25円なら約1億5,625万円です。事業のドル額が同じでも、
換算レートが変われば円表示は変わります。」100万ドルは説明用の仮の金額です。
156.25円は、デモが実際に使う2026-09-04のレート（ECBの円/ユーロとドル/ユーロから算出した156.24677336円）を丸めた値です。
【レートの幅】対象年度のUSD/JPYは約146.8円から約163.9円まで動きました。同じ100万ドルでも、
円換算は約1億4,700万円から約1億6,400万円の幅になります。
【確認先】DAの検証記録に同じ156.24677336円があります。2026-09-05は土曜のため、前営業日2026-09-04の値を1日補完した適用です。
休日は直前値を最大5日まで補完し、補間や将来レートの推測はしていません。
【避ける】「実際の市場や業績を反映したデモ」とは言いません。為替が実データでも、会社、取引、金額は架空です。
ECBの参考レートは公表用の参考値で、企業の決算で使う換算レートではありません。
【CF】キャッシュフローは、一定期間に現金がどのように増減するかという見方です。
商品を売って売上や利益が立っても、代金の入金が翌月ならその時点で現金はまだ増えていません。
【今回のCF】EBITDAの80%を起点にイベントのCF効果を加えた簡易シミュレーションです。
正式な営業CFやフリーCFとして説明しないでください。""",
    ["data/ecb_fx_raw.csv", "reports/da_qa_final_draft.json"],
    "主役はEBITDA。CFは補足の簡易値として説明")
card(s, .65, 2.85, 5.83, 2.8, "為替：共通の通貨に換算", "100万ドル × 156.25円/ドル\n＝ 約1億5,625万円\n2026-09-04のデモの実レート\n出典：ECB公開の日次参考レート")
card(s, 6.85, 2.85, 5.83, 2.8, "CF：現金の増減", "売上が立っても、入金は翌月の例\n利益と現金の動きは一致しない\n今回のCFは簡易シミュレーション", accent=PURPLE)
text(s, "億円の換算：1億円＝100,000,000円 ／ 0.01億円＝100万円", .85, 5.9, 11.6, .45, 18, MUTED)

# 04
s = new_slide(
    "ゼロコピーと実体化の使い分け",
    "8社の原本をどこまでコピーし、どこから参照にするかで、データの重複と更新の手間が変わります。\n"
    "この構成は、業種固有の8表だけを参照にし、全社分析用の表とグラフはコピーして作っています。",
    "",
    """【目安】60秒。「データはコピーしているのか」「全部ゼロコピーか」と聞かれたときに使います。
【話す】「左が8社の原本です。上の経路は、全社分析用の表で、8社分を1つにまとめてコピーして作っています。これを実体化と呼んでいます。下の経路は、業種ごとの固有の8表です。こちらはコピーせず、ショートカットで参照するだけです。これがゼロコピーです。両者を同じものとして説明しないでください。」
【用語】ゼロコピー：OneLakeのショートカットで、元の場所のデータを複製せずに、別のLakehouseから見る仕組みです。実体化：分析用に統合や集計をした表や、グラフを、データとして作って保持することです。
【指す】色の凡例、左の原本（灰色）、中央の上（紫、実体化）と下（青緑、ゼロコピー）、Ontology（定義だけ）、GraphModel（紫、実体化）の順です。親LakehouseはSAP風の管理データも持ち、子会社8個と合わせて9個です。物理表95個とショートカット8個を足して物理表103個とは数えません。
【Ontology】共通の実体名と関係を定義し、各表へ対応付けます。Ontology自体はデータを持たず、定義と対応付けだけです。その定義からGraphModelをマテリアライズしています。GraphModelは参照可能な実体のグラフで、NotebookはそこへGQLを実行します。
【DA】構成上のソースは、Ontologyと親Lakehouseの論理参照面です。DAの金額質問ではSQL集計が確認されています。すべての質問が自動的にGQLへ流れるという主張はしません。
【更新】投影、GraphModel、保存HTMLは、それぞれ更新タイミングが別です。原本の変更が即座にすべてへ反映するとは約束しません。
【Q：全部ゼロコピー？】A：「いいえ。ゼロコピーは業種拡張8表のショートカットだけです。分析用の投影とGraphModelは、コピーして作っています。」
【つなぎ】「ここで登場した4つの用語を食品の例に置き換えます。」""",
    ["docs/architecture.md:81-123", "config/deployment_manifest.json", "config/sector_shortcuts.json"],
    "投影、GraphModel、保存HTMLは、更新のタイミングが別。原本の変更が、すぐ全部に反映されるとは限らない")
ZERO_FILL, MAT_FILL, SRC_FILL, DEF_FILL = "D8EDE9", "ECE6F6", "E5EBEE", "FFFFFF"


def flow_node(slide, x, y, w, h, title, body, fill, accent, dashed=False):
    shp = box(slide, x, y, w, h, fill, line=accent if fill != SRC_FILL else None, rounded=True)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if fill != SRC_FILL:
        shp.line.width = Pt(1.75)
    text(slide, title, x + .18, y + .12, w - .36, .32, 15, accent, True)
    text(slide, body, x + .18, y + .5, w - .36, h - .58, 11.5, INK)


flow_node(s, .62, 2.62, 2.5, 3.0, "子会社の原本", "8つのLakehouse\n会社ごとに分離\n書き換えない\n\n親を含め計9個", SRC_FILL, NAVY)
flow_node(s, 3.65, 2.62, 3.3, 1.35, "全社分析用の表（12表）", "semantic_*の11表とkpi_snapshot。\n8社分を1表にまとめて、コピーして作る", MAT_FILL, PURPLE)
flow_node(s, 3.65, 4.27, 3.3, 1.35, "業種拡張の表（8表）", "raw_*_extension。\n原本の表を、コピーせず参照する", ZERO_FILL, TEAL, dashed=True)
arrow(s, 3.14, 3.3, 3.6, 3.3, PURPLE)
arrow(s, 3.14, 4.95, 3.6, 4.95, TEAL, True)
flow_node(s, 7.5, 2.62, 2.35, 1.35, "Ontology", "実体と関係の定義、\n表への対応付け。\nデータは持たない", DEF_FILL, NAVY)
flow_node(s, 7.5, 4.27, 2.35, 1.35, "GraphModel", "Ontologyの定義から作った\n実体のグラフ。\nコピーして作る", MAT_FILL, PURPLE)
arrow(s, 6.98, 3.3, 7.45, 3.3, PURPLE)
arrow(s, 6.98, 4.78, 7.72, 4.0, TEAL, True)
arrow(s, 8.68, 3.99, 8.68, 4.25, NAVY)
flow_node(s, 10.4, 2.62, 2.3, 1.35, "Fabric DA", "金額集計と説明", DEF_FILL, NAVY)
flow_node(s, 10.4, 4.27, 2.3, 1.35, "Notebook", "GQLで経路をたどり、\nDAGを表示", DEF_FILL, NAVY)
arrow(s, 9.88, 3.3, 10.35, 3.3, NAVY)
arrow(s, 9.88, 4.95, 10.35, 4.95, NAVY)
text(s, "凡例", .65, 5.78, .6, .26, 12, NAVY, True)
for lx, fill, accent, label, dashed, lw in [
    (1.25, SRC_FILL, NAVY, "原本（そのまま残す）", False, 2.2),
    (3.85, ZERO_FILL, TEAL, "ゼロコピー（参照だけ）", True, 2.2),
    (6.45, MAT_FILL, PURPLE, "実体化（コピーして作る）", False, 2.4),
    (9.3, DEF_FILL, NAVY, "定義と利用する側（データなし）", False, 3.3),
]:
    shp = box(s, lx, 5.8, .3, .24, fill, line=accent if fill != SRC_FILL else None, rounded=True)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    text(s, label, lx + .4, 5.8, lw, .26, 11.5, INK)
text(s, "ゼロコピー：OneLakeのショートカットで、元の場所のデータを複製せずに参照する。"
        "実体化：分析用に統合や集計をして、表やグラフとして作る。", .65, 6.14, 12.0, .26, 11, MUTED)


# 05
s = new_slide(
    "4つの用語を食品の例で区別",
    "「Ontology」と「画面に見える図」は同じものではありません。\n定義、照会対象、表示方法を分けると説明の範囲が明確になります。",
    "構成の解説 3  |  用語は名前より役割で説明",
    """【目安】45秒。
【話す】「Ontologyは、契約や活動が何を意味し、どう関係するかの共通定義です。GraphModelには、今回の契約や活動といった実際のレコードがノードと関係として入っています。」
【指す】左上から右上、左下、右下へ。グラフ一般は点と線からなる構造で、線のつながりをたどれます。DAGは向きがあり、ぐるっと一周して戻る閉路がない図です。
【このデモ】Ontologyに17種類の実体と39種類の関係を定義しています。GraphModelの実体照会では33,960ノードと65,684エッジを検証していますが、表示はその全体ではありません。食品の代表1活動に関連する8ノード、7矢印を選んでDAGにしています。
【避ける】GraphModel全体がDAGである、Ontologyが因果推論モデルである、グラフを描けば因果が証明できる、という説明はしません。
【つなぎ】「次に、このデモで定義した実体と関係を、全体のマップで見ます。」""",
    ["reports/step3_live_audit.json:diagram_relations", "config/deployment_manifest.json"],
    "画面のDAGは、取得した関係のうち代表1活動を選んだ部分図")
card(s, .62, 2.73, 5.86, 1.59, "Ontology：共通定義", "例：「活動は契約に属する」という関係")
card(s, 6.83, 2.73, 5.86, 1.59, "GraphModel：実体の照会対象", "例：ある販売活動と小麦契約のレコード")
card(s, .62, 4.67, 5.86, 1.59, "グラフ：点と線の構造", "例：活動、契約、根拠を関係で結ぶ")
card(s, 6.83, 4.67, 5.86, 1.59, "DAG：向きがあり閉路のない図", "例：食品の8ノード／7矢印を表示")

prompt_slide(
    "入力① 年間と月次の稼ぎ",
    Q1,
    "まず全社の稼ぎを見て、増減のあった月を確認します。\n"
    "グラフの崩れと集計の取り違いを防ぐ指定を含みます。そのままコピーできます。",
    "見る点：2本の月次線、年間actualの全額447.57億円と持分管理335.48億円",
    ["Q1: 本ファイルのコピー用質問", "reports/da_qa_published.json"],
    notes_extra="\n【Q1の文面】期間合計をグラフに含めず12か月だけにする、縦軸を億円にする、横軸を年月の文字列にする指定は、"
    "実行時の崩れ（期間合計の混入、縦軸のB表記、時刻軸で1点だけ）を避けるためです。"
    "説明者の実行で、期待どおりの表示を確認済みです（2026-10-26）。")

# 06
monthly = defaultdict(lambda: defaultdict(Decimal))
for r in dataset["semantic_financial_metrics_daily"]:
    if r["metric_code"] == "EBITDA":
        monthly[r["metric_date"][:7]][r["scenario_tag"]] += Decimal(r["amount_group"])
s = new_slide(
    "Step01は基準値と実績の比較",
    "増減の理由を質問する前に、対象年度と指標を固定します。\n持分比率を掛けた金額は、全額連結とは別の管理指標です。",
    "入力①に対応する参考結果  |  保存データから再作成",
    """【目安】ライブ2分。資料の説明は45秒。
【話す】「まず対象年度の全社損益を見ます。baselineがイベントを加える前の基準値、actualが配賦済みの実績です。DAに月次の折れ線比較を依頼します。」
【操作】公開済みFabric DAでQ1を貼り付けます。質問全文は直前の「入力①」スライドとノートにあります。
【指す】月次の2系列を指して期間を確認し、年間actualの全額447.57億円と持分管理335.48億円を区別します。このページの線は保存された元データから再作成した編集可能な参考グラフで、DA画面のスクリーンショットではありません。
【金額】全額44,757,399,732.84円。持分管理33,547,898,914.09円。どちらもactualのみです。51%子会社も全額側では100%を取り込みます。
【避ける】持分管理の差額をそのまま非支配株主利益と呼ぶ、actualにイベント差分をもう一度加える、baselineとactualを合算すること。
【DAの出力が崩れた場合】よくある崩れは3つです。①期間合計がグラフの13個目の点として入り、月次の線が平らになる。②縦軸が億円ではなく「B（10億円）」表記になる。③横軸が時刻軸になり、1点しか表示されない。Q1はこれらを避ける文面にしてありますが、実行ごとに結果は変わります。崩れた場合の追い指示は「期間合計をグラフから外し、12か月だけ、縦軸を億円にして作り直してください」です。それでも崩れる場合は、このページの参考グラフを使い、DAの応答は数値の確認に使います。
【つなぎ】「この変化にどのイベントが配賦されているか、10件を比較します。」
【コピー用Q1】\n""" + Q1,
    ["data/demo_dataset.json:semantic_financial_metrics_daily",
     "reports/da_qa_published.json"],
    "参考グラフは元データから再作成。DAの応答や表示形式はその場で確認")
chartdata = CategoryChartData()
chartdata.categories = [m[2:] for m in sorted(monthly)]
for scenario, label in [("baseline", "baseline"), ("actual", "actual")]:
    chartdata.add_series(label, [float(monthly[m][scenario] / Decimal(10**8))
                                for m in sorted(monthly)])
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(.64), Inches(2.85),
                        Inches(7.45), Inches(3.3), chartdata).chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.font.size = Pt(12)
for series, color in zip(ch.series, [MUTED, TEAL]):
    series.format.line.color.rgb = RGBColor.from_string(color)
    series.format.line.width = Pt(2.5)
ch.category_axis.tick_labels.font.size = Pt(11)
ch.value_axis.tick_labels.font.size = Pt(11)
ch.value_axis.has_title = True
ch.value_axis.axis_title.text_frame.text = "億円"
text(s, "年間actual／全額", 8.65, 2.94, 3.82, .4, 18, TEAL, True)
text(s, "447.57 億円", 8.65, 3.51, 3.9, .65, 33, NAVY, True)
text(s, "年間actual／持分管理", 8.65, 4.51, 3.82, .4, 18, TEAL, True)
text(s, "335.48 億円", 8.65, 5.08, 3.9, .65, 33, NAVY, True)

prompt_slide(
    "入力② 増減の理由を比較",
    Q2,
    "全社の稼ぎを押し上げた出来事と、押し下げた出来事を比べます。\n"
    "イベントとは、損益に影響を与える「出来事」です。",
    "見る点：全10件、プラスとマイナス、食品損失FOD_NEGの約−1.14億円",
    ["Q2: 本ファイルのコピー用質問", "reports/qa_expected_event_totals.json"])

# Per-effect average impact (10,000 JPY units), computed from the verified totals
_per_effect = {k: abs(Decimal(v["full_jpy"])) / v["effects"] / 10**4 for k, v in events.items()}
assert (events["CROSS_NEG"]["effects"], events["FOD_NEG"]["effects"]) == (336, 84)
P19_REASON = (f"【大きく見える理由】1効果あたりの平均は、横断混乱で約{_per_effect['CROSS_NEG']:.0f}万円です。"
              f"10イベントでは約{min(_per_effect.values()):.0f}〜{max(_per_effect.values()):.0f}万円の範囲です。"
              "合計が大きいのは、336効果に配賦されているためです（固有イベントは84効果）。")
# 07
order = ["MTL_NEG", "MTL_POS", "ENE_NEG", "ENE_POS", "FOD_NEG", "FOD_POS",
         "FIN_NEG", "FIN_POS", "CROSS_NEG", "CROSS_POS"]
labels = ["製錬停止", "再生合金", "LNG遅延", "配分最適化", "食品損失",
          "調達改善", "与信費増", "ヘッジ手数料", "横断混乱", "横断回復"]
s = new_slide(
    "Step02は10イベントの増減",
    "損失だけを見ると、好影響や横断イベントを取りこぼします。\n短い項目名を使い、正負の符号と億円の単位で比較します。",
    "入力②に対応する参考結果  |  保存データから再作成",
    """【目安】ライブ2分。資料の説明は45秒。
【話す】「全10イベントの配賦額です。マイナスは損益の押し下げ、プラスは押し上げです。食品損失は約-1.14億円なので、この事例を詳しく確認します。」
【指す】ゼロ線、食品損失、横断混乱と横断回復の順です。固有8件と横断2件の計10件であることを確認します。
【操作】Q2は短い日本語ラベルと符号付き億円を指定し、色や縦横は指定しません。この環境では赤青や横棒を厳密に指定した要求が失敗し、制約を外すとネイティブな縦棒が生成されました。長い名称は切れたため短いラベルにしています。製品の普遍的制約という意味ではありません。
【この図】出典の確定額で再作成した参考図です。棒と文字は編集可能なPowerPoint図形です。DAの実画面キャプチャではなく、表示色や配置を保証するものでもありません。ラベルに対応するevent_idと円単位の値は「10イベントの照合用金額」にあります。ローカルのLibreOfficeではネイティブ棒グラフの負数描画に問題があったため、符号に忠実な図形描画へ切り替えています。
【避ける】横断回復÷横断混乱を回復率と呼ぶこと。異なる時期の設計イベントです。棒の値をactualへ再加算しません。
【横断とは】「他の8件は1つの業種（2社）だけに起きる出来事です。この2件は、4業種すべて（8社）に同時に影響します。業種をまたぐので『横断』と呼んでいます。」横断混乱は、港の混乱と資金繰りの悪化が業種を問わず起きた想定（2026年5月）で、コストが増え、EBITDAを約6.80億円押し下げます。横断回復は、物流と資金管理が改善し、全業種で取扱量が増えた想定（2026年9月）で、約5.84億円押し上げます。「全社に波及するケースも扱えるか」と聞かれたら、第4章の「""" + T_RIPPLE + """」（""" + RIPPLE_PAGES + """）を使います。
""" + P19_REASON + """
【回復の意味】横断回復は、5月の混乱からの回復ではなく、別に設計した独立の好影響イベントです。
【与信費増の訳】FIN_NEGの原文はcounterparty servicing cost shockで、取引先対応コストの急増を指す設計です。ラベルの「与信費増」は短くした呼び名で、貸倒関連の与信費用を意味しません。金融に詳しいお客様には、この点を一言添えます。
【つなぎ】「次のページは、DAの金額を照合する表です。そのあと、入力③で、食品損失がどの記録や販売活動と結び付くかを確認します。」
【コピー用Q2】\n""" + Q2,
    ["reports/qa_expected_event_totals.json", "data/demo_dataset.json:semantic_risk_events"],
    "各effect_idは一度だけ集計。primary_sector='all'の横断2件を除外しない")
text(s, "全額EBITDA影響（億円）／元データから再作成", .92, 2.65, 8.8, .3, 14, MUTED)
zero, scale, chart_x, chart_w = 4.46, .17, 1.06, 8.64
for tick in [-8, -4, 0, 4, 8]:
    y = zero - tick * scale
    box(s, chart_x, y, chart_w, .012, MUTED if tick == 0 else GRAY)
    text(s, f"{tick:+d}".replace("-", "−") if tick else "0", .65, y - .1, .35, .22, 10, MUTED)
for i, (eid, label) in enumerate(zip(order, labels)):
    value = float(Decimal(events[eid]["full_jpy"]) / Decimal(10**8))
    x = chart_x + i * chart_w / 10 + .13
    height = abs(value) * scale
    y = zero - height if value > 0 else zero
    box(s, x, y, .5, height, CORAL if value < 0 else TEAL)
    text(s, f"{value:+.2f}".replace("-", "−"), x - .125,
         y - .27 if value > 0 else y + height + .035,
         .75, .23, 10, CORAL if value < 0 else TEAL, True, PP_ALIGN.CENTER)
    category = label if label != "ヘッジ手数料" else "ヘッジ\n手数料"
    text(s, category, x - .16, 6.01, .82, .43, 9.5, INK, False, PP_ALIGN.CENTER)
text(s, "選ぶ事例", 10.2, 3.05, 2.4, .4, 18, TEAL, True)
text(s, "食品損失", 10.2, 3.64, 2.4, .4, 23, NAVY, True)
text(s, "−1.14", 10.2, 4.25, 2.4, .7, 38, CORAL, True)
text(s, "億円／FOD_NEG", 10.2, 5.04, 2.5, .4, 17, MUTED)
frame_x = chart_x + 8 * chart_w / 10 - .02
dashed_frame(s, frame_x, 3.0, 1.66, 3.3, PURPLE)
text(s, "横断", frame_x, 2.7, .9, .24, 11, PURPLE, True)
text(s, "点線の枠は横断。\n4業種の8社すべてに影響。\n他の8件は1業種の2社。", 10.2, 5.58, 2.5, .85, 12.5, INK)

prompt_slide(
    "入力③ 食品損失の根拠",
    Q3,
    "食品損失について「どの記録、販売、契約とつながるか」を確認します。\n"
    "図はDAに作らせず、この後のNotebookで示します。",
    "見る点：全額−1.1446億円、持分管理−0.8763億円、84効果、根拠と取引のID",
    ["Q3: 本ファイルのコピー用質問", "reports/step3_live_audit.json"],
    notes_extra="\n【代表効果の固定】Q3は代表効果をeffect_idの昇順で先頭に固定します。固定しないと、DAが「最も負の1件」を代表に選び、"
    "Notebookの図と別の効果（SUB_FOD_US_01_A135_1_FOD_NEG_F、2026-02-13）を返すことがあります。"
    "全社KPI IDも、代表効果と同じ日付のものを指定しています（指定しないと、最新日のKPIを返した例があります）。"
    "2026-10-26に公開済みDAで、この文面が図と同じ代表効果（SUB_FOD_JP_01_A133_0_FOD_NEG_F、2026-02-11）、ID、金額を返すことを確認しました。")

s = new_slide(
    "食品損失の回答を読む",
    "食品損失によって、期間全体の稼ぎがどれだけ押し下げられたかを見ます。\n"
    "「84効果」は損益への影響明細の数で、84件の事故という意味ではありません。",
    "入力③に対応する参考結果  |  保存監査値を要約",
    """【目安】60秒。
【回答の位置づけ】このページは保存されたGraph照会とDeltaの金額照合結果から再構成した参考結果です。
Q3を今回DAで実行した応答の転載ではありません。金額とIDの照合に使用します。
【説明】「食品の損失イベント全体で、EBITDAを約1.14億円押し下げています。
その説明に使う代表例として、2月11日の日本食品会社の販売活動と小麦契約を確認します。」
【効果】84件のeffect_idを一度ずつ集計しています。代表1明細の金額を84倍して全体を求めません。
【根拠】根拠は合成業務ログです。温度と廃棄率の観測はありますが、温度が腐敗を起こしたという
実世界の因果の証明ではありません。IDは同じ記録を照合するための目印として示します。
【次の操作】Notebookで保存済みHTMLを開き、日時と図の関係を確認します。
【DAでの確認】Q3の最終版は、2026-10-26に公開済みDAで実行し、このページと同じ代表効果、ID、金額が返ることを確認しました。""",
    ["reports/step3_live_audit.json:exact_*,representative_path"],
    "営業の説明例：「稼ぎの減少額から、関連する販売と根拠記録まで確認できます」")
card(s, .65, 2.8, 3.78, 1.75, "全額の影響", "−1.1446億円\n約1億1,446万円の押し下げ", accent=CORAL)
card(s, 4.77, 2.8, 3.78, 1.75, "持分管理の影響", "−0.8763億円\n所有比率を反映した額", accent=PURPLE)
card(s, 8.89, 2.8, 3.78, 1.75, "影響明細の件数", "84効果\n重複のないeffect_id", accent=TEAL)
box(s, .65, 4.75, 12.02, 1.65, WHITE, rounded=True)
text(s, "代表例：2026-02-11／日本食品会社の外部販売／製粉用小麦", .9, 4.88, 11.5, .35, 16, NAVY, True)
text(s, "根拠：SUB_FOD_JP_01_FOD_NEG_0    活動：SUB_FOD_JP_01_A133_0\n"
     "契約：SUB_FOD_JP_01_C04    効果：SUB_FOD_JP_01_A133_0_FOD_NEG_F\n"
     "日次指標：SUB_FOD_JP_01_2026-02-11_ebitda_actual\n"
     "全社KPI：2026-02-11_EBITDA_FULL_CONSOLIDATED_actual",
     .9, 5.26, 11.5, 1.1, 12, MUTED)

# 08
s = new_slide(
    "Step03で食品損失の経路を表示",
    "DAでは食品の金額とIDを短く確認し、図は専用Notebookへ移ります。\n最初に見せるのは、実行日時を持つ保存済みのHTML出力です。",
    "ライブ操作  |  Step03 Notebookの先頭",
    """【目安】60秒。
【話す】「この図は、Ontologyから作ったGraphModelへ実際にGQLを実行して得た関係を表示しています。根拠、対象活動、契約、財務指標を一緒に確認できます。」
【操作】DAにQ3を入力し、要約とIDだけを確認します。その後、Step03 Notebookタブへ一度切り替えます。名前はTradingDemo_03_食品損失のOntology経路とDAG。先頭の保存済みHTMLを開き、上部3カードとDAGを表示します。表示だけなら実行不要です。
【指す】上段の左から根拠、イベント、1効果、会社日次EBITDA、次に右下の全社日次EBITDA。下段の温度記録、活動、契約にも注目してもらいます。8ノードと7矢印ですが、GQLで得た全関係を全部描いているわけではありません。
【画像の由来】このスライドは実行済み出力として保存されたローカルSVGをブラウザ表示し、そのまま撮影した画像です。FabricのUIを模造したスクリーンショットではありません。
【保存時点】2026-10-26 01:25:34 UTC、10:25:34 JST。いま画面を開いた時刻と同一ではありません。保存結果の閲覧と、現在のGraphの再照会を区別します。
【避ける】「この画像を開いた瞬間にGraphを再取得している」「DAがDAGを生成した」とは説明しません。
【コピー用Q3】\n""" + Q3 + "\n【Notebook】\n" + NOTEBOOK,
    ["reports/step3_food_loss_dag.svg", "reports/step3_live_audit.json",
     "reports/step3_ui_verification.json"],
    "図は実行済みSVGのローカル撮影。NotebookのクラウドUI画像ではない")
image = HERE / "actual_saved_dag.png"
from PIL import Image
iw, ih = Image.open(image).size
w = 8.7
h = w * ih / iw
s.shapes.add_picture(str(image), Inches(.65), Inches(2.57), width=Inches(w), height=Inches(h))
text(s, "イベント全体", 9.79, 2.81, 2.83, .36, 18, TEAL, True)
text(s, "-1.1446", 9.79, 3.34, 2.83, .65, 32, CORAL, True)
text(s, "億円／84効果", 9.79, 4.03, 2.83, .4, 17, MUTED)
text(s, "図は代表1活動\n8ノード／7矢印", 9.79, 4.74, 2.83, .9, 18, NAVY, True)
pill(s, "Notebookを開く", 9.79, 5.89, 2.64, NOTEBOOK)

# 09
s = new_slide(
    "矢印は3種類の意味を持つ",
    "つながった線をすべて「原因から結果」と読むと、説明を誤ります。\nこのDAGには根拠の関連、設計上の配賦、指標への対応が混在します。",
    "Step03の読み方  |  関係名を日本語で説明",
    """【目安】45～60秒。
【話す】「灰色は記録や対象との関連、紫はこの架空デモで設計した配賦、緑は日次指標への対応や集約先です。すべての線が実証された因果を表しているわけではありません。」
【指す】根拠からイベント、イベントから効果、効果から会社日次、会社日次から全社日次へと上段を読みます。
【方向】図の実際の向きを守ります。効果から活動へEffectOnActivity、活動から契約へActivityUnderContractです。「根拠→活動→契約」は説明の確認順としては使えますが、その通りの矢印列が図にあるとは言いません。
【省略】GQLではEffectExplainedByEvidence、EffectOnContract、EventImpactsActivityも確認済みですが、図では非循環で読みやすい部分だけを示しています。取得クエリは複数の条件を同時に検証しており、図はその要約です。
【財務】効果の差分、会社の日次総額、全社の日次総額が線でつながっても、3つを加算するという意味ではありません。
【つなぎ】「まず財務金額以外の記録を読み、何が言えて何が言えないかを確認します。」""",
    ["reports/step3_live_audit.json:diagram_relations",
     "scripts/step3_food_loss_runtime.py:197-267"],
    "説明用の再構成図。矢印の方向は保存された実関係に合わせて記載")
nodes(s, [
    ("根拠記録", "架空の業務ログ", MUTED),
    ("食品イベント", "FOD_NEG", PURPLE),
    ("財務効果", "代表1件の差分", PURPLE),
    ("会社日次指標", "actualの総額", TEAL)],
    xs=[.65, 3.84, 7.03, 10.22], w=2.45, h=1.5,
    edge_colors=[MUTED, PURPLE, TEAL])
text(s, "根拠の関連", 2.58, 2.8, 1.95, .3, 13, MUTED)
text(s, "設計した配賦", 5.73, 2.8, 1.95, .3, 13, PURPLE)
text(s, "指標との対応", 8.93, 2.8, 2.1, .3, 13, TEAL)
box(s, .66, 5.04, 12.02, 1.06, WHITE, rounded=True)
text(s, "灰色：関連", .9, 5.13, 3.25, .4, 20, MUTED, True)
text(s, "記録や対象との関連。\n原因を示すとは限らない", .9, 5.6, 3.4, .44, 12, INK)
text(s, "紫：影響の配賦", 4.82, 5.13, 3.5, .4, 20, PURPLE, True)
text(s, "このデモで設計した、影響額の割り当て", 4.82, 5.6, 3.4, .44, 12, INK)
text(s, "緑：対応と集約", 8.81, 5.13, 3.5, .4, 20, TEAL, True)
text(s, "日次指標への対応と、全社の合計への集約", 8.81, 5.6, 3.5, .44, 12, INK)

# 10
s = new_slide(
    "代表記録でわかること、取れる行動",
    "食品損失の記録を1件見ても、経営者は、それで何をしてよいのかが分からないことがあります。\n"
    "この記録で言えること、言えないことを分けて、取れる行動を示します。",
    "",
    """【目安】60秒。「経営者は、この記録から何ができるのか」と聞かれたときに使います。
【話す】「この代表記録から言えるのは、いつ、どこで、何が記録され、いくらの影響が出たかまでです。原因は、この記録だけでは決められません。経営者が取れる行動は、調べる先を絞って、担当部門に事実確認を依頼するところまでです。」
【わかること】2026年2月11日、日本の食品会社の外部販売、製粉用小麦（商品コードmilling_wheat）、契約SUB_FOD_JP_01_C04。観測は、温度マイナス18.0度、廃棄率3.5%です。この1効果は−3,003,488.94円（約−300万円）です。
【わからないこと】許容温度や契約上の閾値は記録にありません。マイナス18度が異常かどうかは判断できず、「高温障害」「小麦が冷凍必須」といった現実世界の知識で補いません。根拠は、SIM://SUB_FOD_JP_01/FOD_NEG/operations_logという合成業務ログで、原文にSYNTHETICと書かれています。実在の監査証憑ではなく、実世界の因果の証明でもありません。代表1件だけでは、同じ事象が他の取引にも起きたかはわかりません。
【取れる行動】該当の契約と活動のIDで調べる先を指名し、担当部門に事実確認を依頼できます。たとえば、温度ログ、出荷記録、クレームや保険対応の有無です（説明用の例で、このデモのデータにはありません）。影響が大きい契約の比較は、Step04で行います。
【まだ取れない行動】原因の断定と、再発防止策の決定です。許容温度などの追加の記録が要ります。
【ID】活動SUB_FOD_JP_01_A133_0、契約SUB_FOD_JP_01_C04、根拠SUB_FOD_JP_01_FOD_NEG_0です。長いIDは口頭で読み上げず、「食品損失の回答を読む」のページのID表を指します。
【日付】イベント期間は2月9日から3月1日、代表活動と効果は2月11日です。イベント開始日と個別の効果日を同一視しません。
【つなぎ】「同じ図に出てくる4つの金額は、測っている範囲が違います。」""",
    ["reports/step3_live_audit.json:representative_path,cold_chain_observation",
     "reports/da_qa_published.json"],
    "調べる先は、契約と活動のIDで特定できる。原因の断定には、許容温度などの追加の記録が要る")
rep_cols = [
    ("記録からわかること", TEAL, [
        ("いつ", "2026-02-11"),
        ("どこで", "日本の食品会社の外部販売（製粉用小麦）\n契約：SUB_FOD_JP_01_C04"),
        ("何が記録されたか", "温度 −18.0℃、廃棄率 3.5%"),
        ("いくら", "この1効果で約−300万円\n（−3,003,488.94円）"),
    ]),
    ("記録だけではわからないこと", CORAL, [
        ("許容温度と閾値", "未記録。−18℃が異常かは判断できない"),
        ("腐敗の原因", "設備か運用かは特定できない。根拠は架空の業務ログで、因果の証明ではない"),
        ("広がり", "代表1件では、他の取引の状況はわからない"),
    ]),
    ("経営者が取れる行動", PURPLE, [
        ("取れる", "該当の契約と活動を指名し、担当部門に事実確認を依頼する（温度ログ、出荷記録、クレーム対応の有無）"),
        ("取れる", "影響が大きい契約を比べて、優先順位を付ける（Step04）"),
        ("まだ取れない", "原因の断定と、再発防止策の決定"),
    ]),
]
for c, (head, accent, items) in enumerate(rep_cols):
    x = .62 + c * 4.09
    box(s, x, 2.6, 3.89, 3.78, WHITE, rounded=True)
    text(s, head, x + .22, 2.74, 3.45, .34, 17, accent, True)
    stack_items(s, x + .22, 3.2, 3.45, items, accent, size=12, gap=.16)

# 11
s = new_slide(
    "4つの金額は範囲が違う",
    "食品損失の説明には、似た4つの金額が出てきますが、何の合計かが違います。\n"
    "足したり、直接比べたりせず、何の合計かを添えて話します。",
    "",
    """【目安】60秒。質疑で最も間違えやすい箇所です。
【話す】「4つの金額は、どれも正しい数字です。ただし、何の合計かが違います。一番外側が、その日の全8社のEBITDA。その中に日本食品会社のEBITDA。さらにその中に、食品損失の代表1効果があります。右の食品損失の全体は、代表1効果を含む84件の合計です。足したり、1つを別の1つと比べたりしません。」
【指す】外側の枠から内側へ、④、③、②の順に指し、最後に右の①を指します。
【金額】①FOD_NEG全体（84効果、期間全体）：-114,461,186.25円（持分管理なら-87,630,294.31円）。②代表効果SUB_FOD_JP_01_A133_0_FOD_NEG_F：-3,003,488.94円（所有比率51%、持分管理は-1,531,779.36円）。③2026-02-11の日本食品会社のactual EBITDA：+3,711,231.19円。④同日の全8社のactual EBITDA：+113,069,201.90円。
【包含関係】②の押し下げは、③のactualにすでに反映されています。2026-02-11の日本食品会社には、食品損失の効果が3件あり（-3,003,488.94円、-2,466,144.46円、-1,422,653.08円）、合計は-6,892,286.48円です。損失がなければ、この日のbaselineは+10,603,517.67円で、actualは+3,711,231.19円です（差は-6,892,286.48円で、3件の合計と一致します）。③は④に含まれます。①は、②を含む84効果の合計で、期間はイベント期間（2月9日から3月1日）に広がります。
【よくある読み違い】「食品損失があるので、2月11日の全社は1.13億円の損失」と言う、「-1.14億円と+1.13億円を足す」、「84効果×代表効果＝全体」と計算する。
【精度】画面の億円表記は丸めています。厳密な照合には円の表示や監査JSONを使います。効果の粒度と日次の粒度では、換算の丸めに微差があり得るので、別粒度の差を機械的に不整合と断定しません。
【つなぎ】「関係の取得と金額の照合を別処理にすることで、この違いを保っています（第6章）。」""",
    ["reports/step3_live_audit.json:exact_*",
     "data/demo_dataset.json:semantic_financial_metrics_daily、semantic_event_financial_effects（2026-02-11の照合）"],
    "画面の億円表示は丸めた値。厳密に照合するときは、円単位の値を使う")
box(s, .62, 2.6, 7.6, 3.82, "D8EDE9", rounded=True)
text(s, "④ 同じ日の全8社のEBITDA", .88, 2.72, 4.6, .3, 15, TEAL, True)
text(s, "+1.13億円", 5.6, 2.68, 2.4, .36, 22, TEAL, True, PP_ALIGN.RIGHT)
text(s, "2026-02-11の全取引の合計　+113,069,201.90円", .88, 3.06, 7.0, .24, 11, INK)
box(s, .98, 3.5, 6.88, 2.78, "EAF5F2", rounded=True)
text(s, "③ その日の日本食品会社のEBITDA", 1.22, 3.6, 4.4, .3, 15, TEAL, True)
text(s, "+371万円", 5.5, 3.56, 2.1, .36, 22, TEAL, True, PP_ALIGN.RIGHT)
text(s, "2026-02-11の全取引の合計　+3,711,231.19円", 1.22, 3.94, 6.3, .24, 11, INK)
text(s, "食品損失3効果（計−689万円）を反映した後の値。反映前は+1,060万円", 1.22, 4.18, 6.3, .24, 11, INK)
box(s, 1.34, 4.74, 6.18, 1.3, WHITE, rounded=True)
text(s, "② 食品損失の代表1効果", 1.58, 4.84, 3.8, .3, 15, CORAL, True)
text(s, "−300万円", 5.2, 4.8, 2.1, .36, 22, CORAL, True, PP_ALIGN.RIGHT)
text(s, "SUB_FOD_JP_01_A133_0_FOD_NEG_F　−3,003,488.94円", 1.58, 5.18, 5.7, .24, 11, INK)
text(s, "③に反映されている押し下げ分の1つ", 1.58, 5.5, 5.7, .24, 11, INK)
box(s, 8.5, 2.6, 4.2, 1.95, WHITE, rounded=True)
text(s, "① 食品損失の全体（84効果）", 8.74, 2.72, 3.8, .3, 15, CORAL, True)
text(s, "−1.14億円", 8.74, 3.08, 3.8, .42, 26, CORAL, True)
text(s, "期間全体（イベント期間は2月9日から3月1日）の合計　−114,461,186.25円", 8.74, 3.62, 3.75, .4, 11, INK)
text(s, "②を含む84件の合計", 8.74, 4.12, 3.75, .24, 11, INK)
box(s, 8.5, 4.68, 4.2, 1.74, WHITE, rounded=True)
text(s, "読み違いの例", 8.74, 4.78, 3.8, .28, 13, PURPLE, True)
stack_items(s, 8.74, 5.1, 3.78, [
    ("誤", "食品損失があるので、2/11の全社は1.13億円の損失"),
    ("正", "1.13億円は損失ではなく、その日の全取引の合計"),
], PURPLE, size=11, gap=.05, label_h=.2)


# 12
s = new_slide(
    "関係はGraph、金額はDeltaで照合",
    "Graphの数値は浮動小数点で、金額の合計を会計の根拠にはできません。\n"
    "それでも関係はGraphでたどるので、2つを別々に確認して照合しています。",
    "",
    """【目安】60秒。「なぜGraphだけで金額を出さないのか」「図の数字はどう保証しているのか」と聞かれたときに使います。非技術者にはカードの見出しだけで足ります。
【話す】「NotebookはGraphで本物の関係を読み、その84効果が、Delta側の84効果と一致するかを調べます。金額の合計は、小数を正確に扱うDeltaのDECIMALで計算します。2つが一致したときだけ、図を保存します。」
【なぜ分けるか】Graphの数値は浮動小数点で、会計の合計の唯一の根拠にはしません。一方、Deltaの表は、効果、会社の日次、全社の日次を区別して、重複なく合計できます。Microsoft Learn（Ontology概要、2026-10-16更新）も、Graphは関係の経路が重要な場面のための任意の実行層で、通常の参照や集計、メトリクスは元のデータで実行できる、と説明しています。このデモの分担は、その説明に沿っています。
【指す】左のGraph、右のDelta、下の一致の検査の順です。
【Graph】OntologyからマテリアライズされたGraphModelに対して、公開Graph APIのbeta GQLを実行しています。定義そのものに未提供の実体APIを呼んだり、SQL JOINをGQLと呼び換えたりしていません。
【金額】Sparkでsemantic_event_financial_effects、semantic_financial_metrics_daily、kpi_snapshotのDelta DECIMALを読取専用で照合します。84 effect_idの集合が一致し、全84金額が小数点以下2桁で一致しました。
【選択】代表はeffect_id順で決定しており、都合のよい値を手作業で書いた図ではありません。食品観測の実関係も追加GQLで取得します。
【停止条件】認証失敗、空結果、取得上限到達、ID不一致、金額不一致では停止します。古い静的な図をライブ結果の代わりに表示して成功扱いにはしません。
【認証】Notebook実行時にnotebookutils.credentials.getToken('pbi')を利用。デッキやソースにトークンは埋め込みません。原本とGraphは読取専用ですが、HTML出力は自分自身の新Notebookへ保存します。
【つなぎ】「構成の説明は以上です。新しいOntologyの機能は、第7章で説明します。」""",
    ["scripts/step3_food_loss_runtime.py:85-182",
     "reports/step3_live_audit.json",
     "Microsoft Learn：fabric/iq/ontology/overview（2026-10-16更新、Optional graph execution）"],
    "このページを使う場面：Graphだけで金額を出さない理由や、図の数字をどう保証しているかを聞かれたとき")
for k, (head, body, accent) in enumerate([
    ("Graph（GQL）：つながりをたどる",
     "食品損失の84効果と、その関係（活動、契約、会社）を取得する\n代表1活動の観測値（温度、廃棄率）を取得する\n向かないこと：金額の合計（数値が浮動小数点）", TEAL),
    ("Delta（Spark）：金額を正確に合計する",
     "84効果の金額を、DECIMALで重複なく合計する\n効果、会社の日次、全社の日次を区別して読む\n向かないこと：関係の経路をたどること", PURPLE),
]):
    x = .62 + k * 6.22
    box(s, x, 2.6, 5.86, 2.0, WHITE, rounded=True)
    text(s, head, x + .24, 2.72, 5.4, .34, 16, accent, True)
    text(s, body, x + .24, 3.14, 5.4, 1.4, 12.5, INK)
arrow(s, 3.55, 4.64, 3.55, 4.92)
arrow(s, 9.77, 4.64, 9.77, 4.92)
box(s, .62, 4.98, 5.86, .62, PALE, rounded=True)
text(s, "Graphで取得した84効果のID", .86, 5.12, 5.4, .3, 14, TEAL, True)
box(s, 6.84, 4.98, 5.86, .62, PALE, rounded=True)
text(s, "Deltaで合計した84効果のIDと金額", 7.08, 5.12, 5.4, .3, 14, TEAL, True)
text(s, "＝", 6.5, 5.06, .3, .4, 20, NAVY, True, PP_ALIGN.CENTER)
text(s, "一致したときだけ、代表経路のHTMLとSVGを保存する。不一致、認証失敗、空の結果では停止する。",
     .64, 5.74, 12.0, .3, 12.5, NAVY, True)
text(s, "Learn：Graphは、関係の経路が重要な場面のための任意の実行層。通常の参照、集計、メトリクスは、元のデータで実行できる。",
     .64, 6.08, 12.0, .3, 10.5, MUTED)


# 13
s = new_slide(
    "Notebookは保存済みの結果から始める",
    "Notebookは実行済みの結果（HTML）を先頭に置いているので、再実行を待たずに説明できます。\n最新のGraphを読み直す操作は、保存結果を見せる操作とは別に、必要なときだけ行います。",
    "",
    """【目安】60秒。実演前は必ず手元で一度確認します。
【通常表示】①正しいNotebookを開く。②先頭の保存済みHTMLを表示。③読取日時を確認。④上部カード、DAG、根拠原文、必要ならID表の順に指します。表示だけなら「すべて実行」を押しません。
【保存とライブ】最初の表示セルはDEMO_HTMLがセッションにあればそれを表示します。ない場合のコード上の案内と、ファイルに保存された過去の出力は別です。後段runtimeが実GQLを照会し、自Notebookの先頭出力へ結果を保存する設計です。最初のセルが先に走ることは、この配置を変える理由にはしません。
【更新】更新が必要な場合だけ「すべて実行」を1回押し、完了を待ちます。完了後、画面の状態に応じて先頭表示セルを再表示するかNotebookを開き直し、日時が更新されたかを確認します。自動で必ず先頭が最新表示になるとは約束しません。
【失敗時】エラーを確認し、再実行の連打を避けます。読取認証、Graphの利用可否、データ照合などを切り分けます。保存結果を続けて使うなら「こちらは保存時点の結果です」と明示します。権限変更やGraph再生成、DA再公開をその場で試みません。
【現状】2026-10-26の実行完了後、セルを走らせずに開き直した画面で8ノード7矢印と金額、日時の可視性を検証済みです。Notebookそのもののフロントエンド状態まで毎回同じとは保証しません。
【つなぎ】「次に、図の矢印の意味を説明します。」
【直接リンク】\n""" + NOTEBOOK,
    ["scripts/build_step3_notebook.py:39-64", "scripts/step3_food_loss_runtime.py",
     "reports/step3_ui_verification.json"],
    "更新操作は必要時だけ1回。保存時点の結果を現在時点のライブ結果と呼ばない")
card(s, .63, 2.78, 3.78, 2.99, "通常の説明", "Notebookを開く\n先頭HTMLと日時を確認\nそのまま図を説明", number="1")
card(s, 4.77, 2.78, 3.78, 2.99, "最新値への更新", "「すべて実行」を1回\n完了後に再表示／再読込\n日時と結果を確認", number="2")
card(s, 8.91, 2.78, 3.78, 2.99, "エラーの場合", "実行の連打を避ける\n未検証範囲を明示\n保存結果なら日時を告知", accent=CORAL, number="3")
for cx, label, when, accent in [
        (.63, "使う場面", "普段のデモ。保存済みの結果を、そのまま説明する。", TEAL),
        (4.77, "使う場面", "保存日時より新しい値を、見せたいとき。", TEAL),
        (8.91, "使う場面", "実行が失敗した、結果が出ないとき。", CORAL)]:
    text(s, label, cx + .25, 4.93, 3.3, .24, 11.5, accent, True)
    text(s, when, cx + .25, 5.2, 3.3, .5, 12, INK)

# Step04: impact points and action ideas
con_by_id = {c["contract_id"]: c for c in dataset["semantic_contracts"]}
co_names = {c["entity_id"]: c["entity_name"] for c in dataset["semantic_company_profile"]}
neg_by_contract = defaultdict(lambda: {"amount": Decimal(0), "events": set()})
for r in dataset["semantic_event_financial_effects"]:
    if r["event_id"].endswith("_NEG"):
        neg_by_contract[r["contract_id"]]["amount"] += Decimal(r["delta_ebitda_full_jpy"])
        neg_by_contract[r["contract_id"]]["events"].add(r["event_id"])
ranked = sorted(neg_by_contract, key=lambda cid: neg_by_contract[cid]["amount"])
scope = [cid for cid in ranked
         if cid in ranked[:5] or len(neg_by_contract[cid]["events"]) >= 2]
da_tier = {"SUB_FIN_US_01_C08": "高", "SUB_FIN_US_01_C07": "高", "SUB_FIN_JP_01_C08": "高",
           "SUB_FIN_JP_01_C07": "高", "SUB_MTL_JP_01_C07": "中",
           "SUB_FIN_JP_01_C09": "低", "SUB_FIN_US_01_C09": "低"}
assert set(scope) == set(da_tier), scope
neg_total = sum(v["amount"] for v in neg_by_contract.values())
top5_share = abs(sum(neg_by_contract[c]["amount"] for c in ranked[:5]) / neg_total) * 100
assert (len(con_by_id), len(neg_by_contract), round(top5_share)) == (96, 48, 16)

prompt_slide(
    "入力④ 影響点の特定と対応案",
    Q_ACT,
    "影響額を確認しただけでは、どこから手を打つかは決まりません。\n"
    "影響額に加え、複数のイベントの影響を受けた契約を見て、優先順位を付けます。",
    "見る点：7契約、うち6契約が2回影響、最大は金融の−0.43〜−0.45億円、優先度の基準",
    ["Q4: 本ファイルのコピー用質問", "data/demo_dataset.json:semantic_event_financial_effects"],
    size=15,
    result_note=(
        "次ページの表は元データから再作成した参考結果、その次のページの対応案は2026-10-26に公開済みDAでQ4を実行した応答の要約です。"
        "実行ごとに文面や優先度の分け方は変わります。"),
    notes_extra=(
        "\n【期待する表（7行）】SUB_FIN_US_01_C08 −0.4535（28件、2イベント）、SUB_FIN_US_01_C07 −0.4365（27件、2イベント）、"
        "SUB_FIN_JP_01_C08 −0.4345（28件、2イベント）、SUB_FIN_JP_01_C07 −0.4283（28件、2イベント）、"
        "SUB_MTL_JP_01_C07 −0.3714（14件、1イベント）、SUB_FIN_JP_01_C09 −0.2371（16件、2イベント）、"
        "SUB_FIN_US_01_C09 −0.2356（17件、2イベント）。単位は億円です。"
        "\n【優先度】高、中、低はDAの判断です。判定基準が文章で出ているか確認します。"
        "\n【取引先】データに取引先名はありません。DAには作らせず、契約IDと商品で示します。"
        "取引先名まで見せるには、架空の取引先マスタを追加してOntologyとグラフを更新する必要があります（今回は未実施）。"
        "\n【順序】Q4はNotebookの前後どちらでも実行できます。画面の切替を減らしたい場合は、Q3の直後にDAでQ4を実行してからNotebookへ移ります。"
        "\n【Q4の文面】primary_sector='all'の除外防止、effect_idを一度だけ合計する指定、取引先名を作らせない指定は、"
        "過去の誤り（横断イベントの除外、重複集計、名称の捏造）を避けるためです。"))

s = new_slide(
    "影響が大きい契約と繰り返し影響を受けた契約",
    f"ネガティブ5件の影響は{len(con_by_id)}契約中{len(neg_by_contract)}契約に広がり、上位5件でも全体の約{top5_share:.0f}%です。\n"
    "2回影響を受けた契約を加えると、優先して見る先を絞れます。",
    "入力④に対応する参考結果  |  表は元データから再作成、優先度はDA実行例",
    """【目安】60秒。
【話す】「食品損失の次に、どの契約から手を打つかを考えます。5つのマイナスのイベントを契約ごとに通算すると、影響は96契約中48契約に広く薄く散り、上位5件でも全体の約16%です。そこで、複数のイベントの影響を受けた契約にも注目します。」
【指す】優先度の列と、右端の「回数」の列です。金融子会社2社のtrade_finance契約（C07、C08）は、与信費増（FIN_NEG）と横断混乱（CROSS_NEG）の2回の影響を受け、各−0.43〜−0.45億円と最大です。
【時期】2回の影響は同時期ではありません。FIN_NEGは4月中旬〜5月初め、CROSS_NEGは5月中旬〜下旬に、約2週間あけて続けて当たっています。
【優先度】高、中、低はDAの判断です。判定基準が画面に出ているか確認します。2026-10-26の指示文の更新後は、C09の2契約を「中」とし、低を該当なしとする実行もありました（金額と契約は同じです）。実行例では、額が大きく2回影響を受けた4契約を高、額は大きいが1回だけのSUB_MTL_JP_01_C07を中、2回影響を受けたが額が小さいC09の2契約を低としました。3位と4位の差は約63万円で、実質同じ水準です。
【金額】SUB_FIN_US_01_C08 −0.4535、SUB_FIN_US_01_C07 −0.4365、SUB_FIN_JP_01_C08 −0.4345、SUB_FIN_JP_01_C07 −0.4283、SUB_MTL_JP_01_C07 −0.3714、SUB_FIN_JP_01_C09 −0.2371、SUB_FIN_US_01_C09 −0.2356（億円、全額連結EBITDAへの影響）です。
【避ける】取引先名を口にしません（データにありません）。優先度を、取るべき対応の確定順位と説明しません。与信費増を、貸倒関連の与信費用と説明しません。
【つなぎ】「では、それぞれの契約に対して、どんな対応が考えられるでしょうか。」""",
    ["data/demo_dataset.json:semantic_event_financial_effects,semantic_contracts",
     "説明者の公開済みDA実行例（優先度の列）、2026-10-26"],
    "取引先名はデータにない。契約IDで特定し、実務では取引先マスタにつなぐ")
tier_color = {"高": CORAL, "中": PURPLE, "低": MUTED}
cols = [("優先度", .83, .8), ("契約ID", 1.7, 2.5), ("子会社", 4.25, 2.7), ("商品", 7.0, 1.5),
        ("影響したイベント", 8.55, 2.0), ("影響額（億円）", 10.55, 1.5), ("回数", 12.1, .55)]
for label, x, w in cols:
    text(s, label, x, 2.68, w, .3, 13, TEAL, True)
for i, cid in enumerate(scope):
    y = 3.05 + i * .4
    if i % 2 == 0:
        box(s, .67, y - .07, 12.0, .38, WHITE, flat=True)
    info = neg_by_contract[cid]
    tier = da_tier[cid]
    vals = [tier, cid, co_names[con_by_id[cid]["entity_id"]], con_by_id[cid]["commodity"],
            ", ".join(sorted(info["events"], reverse=True)),
            f"{info['amount'] / Decimal(10**8):.4f}".replace("-", "−"), str(len(info["events"]))]
    for j, ((label, x, w), val) in enumerate(zip(cols, vals)):
        text(s, val, x, y, w, .3, 12.5, tier_color[tier] if j == 0 else (CORAL if j == 5 else INK), j == 0)
text(s, "表は、上位5件と、2件以上のイベントの影響を受けた契約を重複なく並べた7契約です。",
     .85, 6.02, 11.7, .3, 13, MUTED)

s = new_slide(
    "優先度別の対応案（DA実行例）",
    "影響点が分かっても、打つ手は事実だけからは決まりません。\n"
    "DAは、データから言えることと、仮説としてのアイデアを分けて出します。",
    "入力④に対応する応答  |  説明者の実行例（2026-10-26）を要約",
    """【目安】60秒。
【位置づけ】このページの文面は、説明者が2026-10-26に公開済みDAでQ4を実行した応答（対応案の部分）の要約です。実行ごとに文面は変わります。指示文の更新後は、答えの書き出しが「結論→前提→表」の順になり、C09の2契約の優先度が「低」でなく「中」になる実行もありました。
【話す】「データから言えることと、仮説としてのアイデアを分けて出しています。アイデアはあくまで仮説で、実行を決めるのは担当者です。」
【指す】高、中、低の順に、各カードの「データから言えること」、次に「仮説としてのアイデア」を指します。
【取引先】データには取引先名がないため、契約IDで示しています。実務では取引先マスタを契約につなぎ、取引先単位で見ます。
【実運用】過去の対応事例（Foundry IQ）や担当者とのやり取り（Work IQ）と組み合わせると、一般論ではなく自社の対応の型になります。通知やタスク化は実運用の発展として説明し、このデモの範囲外とします。
【避ける】アイデアを実証済みの対策や因果の証明と言いません。「与信費上昇」は与信費増（FIN_NEG）の言い換えで、貸倒関連の与信費用ではありません。
【実行例の原文】\n""" + DA_ACTION_TEXT,
    ["説明者の公開済みDA実行例（対応案の部分）、2026-10-26",
     "data/demo_dataset.json:semantic_event_financial_effects"],
    "次の1歩：高の4契約について、信用枠、担保・ヘッジ条件の現状を担当者に確認する")


def action_card(slide, x, y, w, h, tier, ids, fact, idea, accent):
    box(slide, x, y, w, h, WHITE, rounded=True)
    text(slide, tier, x + .22, y + .16, w - .44, .42, 20, accent, True)
    text(slide, ids, x + .22, y + .62, w - .44, .5, 10.5, MUTED)
    text(slide, "データから言えること", x + .22, y + 1.14, w - .44, .26, 12.5, TEAL, True)
    text(slide, fact, x + .22, y + 1.42, w - .44, .85, 13)
    text(slide, "仮説としてのアイデア", x + .22, y + 2.3, w - .44, .26, 12.5, PURPLE, True)
    text(slide, idea, x + .22, y + 2.58, w - .44, 1.0, 13)


for k, (tier, ids, fact, idea, accent) in enumerate([
    ("高：金融の4契約", "SUB_FIN_US_01_C08, SUB_FIN_US_01_C07\nSUB_FIN_JP_01_C08, SUB_FIN_JP_01_C07",
     "4契約ともtrade_finance\nCROSS_NEGとFIN_NEGの両方の影響\n各0.43〜0.45億円の押し下げ",
     "契約別の信用枠、担保・ヘッジ条件、例外承認ルールを優先点検し、早期エスカレーション条件を設定する。", CORAL),
    ("中：金属の1契約", "SUB_MTL_JP_01_C07",
     "copper_alloy契約\nCROSS_NEGのみ\n0.37億円の押し下げ",
     "代替調達、在庫バッファ、納期・価格条件の見直しなど、横断的な混乱に耐える契約運用を検討する。", PURPLE),
    ("低：金融の2契約", "SUB_FIN_JP_01_C09, SUB_FIN_US_01_C09",
     "trade_finance契約で2イベントの影響\n各0.24億円の押し下げ\n上位4契約より小さい",
     "高優先契約で有効性を確認した監視ルールを横展開し、影響拡大時に優先度を引き上げる基準を設ける。", MUTED),
]):
    action_card(s, .65 + k * 4.07, 2.68, 3.9, 3.65, tier, ids, fact, idea, accent)

# 14
s = new_slide(
    "デモで見せた3つの価値",
    "経営の数字は、会社や部署によって言葉の意味が違い、根拠を確認するにも手間がかかります。\n"
    "Ontologyは、意味とつながりを一度定義して、質問のたびに作り直さずに使えるようにします。",
    "",
    """【目安】45秒。10分枠の最後に使います。
【話す】「ここまでで見せたのは、金額そのものよりも、金額の意味と、金額どうしのつながりと、根拠の記録を、質問のたびに作り直さずに使える状態です。全額と持分管理、基準値と実績の区別は、定義として決めてあります。食品損失から契約、会社、全社の指標までは、定義した関係をたどりました。答えには、根拠と契約のIDを添えています。」
【指す】3つのカード。上から「ふだんの困りごと」と「このデモでは」を順に読みます。
【価値の根拠】Microsoft Learn（Ontology概要、2026-10-16更新）は、Ontologyを、データの物理的な保存の形と、業務での理解の差を埋める層と説明しています。利用者ごとに、表、列、メジャー、結合を解釈する代わりに、再利用できる業務の概念と、データへの対応付けを与えます。
【SQLとの比較】SQLでも同じ数値は作れます。比較が出たら、第4章の「SQLとOntologyの違い」のページを使います。OntologyだけがJOINできる、SQLでは不可能、という説明はしません。
【現状】このデモでOntologyが担っているのは、実体と関係の定義（17実体、39関係）とGraphです。全額連結などの言葉の定義の多くは、いまはDAの指示文にあります。
【追加の質問】時間があれば、ENE_POSの好影響と、横断2件の業種別比較を聞きます。入力文は、第3章の「4つの入力の使い分け」のページにあります。
【避ける】画面操作を自動業務判断や本番会計の完成形と結び付けること。因果の証明、全面ゼロコピー、自動最新化とも言いません。""",
    ["docs/architecture.md:112-143", "reports/qa_expected_event_totals.json",
     "Microsoft Learn：fabric/iq/ontology/overview（2026-10-16更新）"],
    "次の一歩：お客様のイベントを1件選び、同じ形で定義できるかを確かめる（第5章）")
value_cards = [
    ("意味をそろえる", TEAL,
     "全額か持分か、基準値か実績かの意味が、人によって違うことがある。",
     "全額連結と持分管理、baselineとactualを区別して定義し、8社をJPYの同じ尺度で比べた。"),
    ("つながりをたどる", PURPLE,
     "「この損失はどの契約に効いたか」と聞くたびに、詳しい人が表をつなぎ直すことが多い。",
     "イベント、効果、活動、契約、会社、全社KPIのつながりを関係として定義し、食品損失から契約までたどった。"),
    ("根拠を添える", CORAL,
     "数字は出ても、元の記録を確かめるのに時間がかかることが多い。",
     "答えに、根拠、契約、効果のIDと金額を添えた。金額は別の集計で照合した。"),
]
for k, (head, accent, trouble, demo) in enumerate(value_cards):
    x = .62 + k * 4.09
    box(s, x, 2.6, 3.89, 3.8, WHITE, rounded=True)
    text(s, head, x + .24, 2.74, 3.45, .38, 19, accent, True)
    text(s, "ふだんの困りごと", x + .24, 3.32, 3.45, .24, 11.5, MUTED, True)
    text(s, trouble, x + .24, 3.6, 3.45, 1.0, 14, INK)
    text(s, "このデモでは", x + .24, 4.62, 3.45, .24, 11.5, accent, True)
    text(s, demo, x + .24, 4.9, 3.45, 1.4, 14, INK)


# 15
s = new_slide(
    "4つの入力の使い分け",
    "実演でDAに4つの質問を入れる順番と、うまくいかないときの代わりの手段を、先に決めておきます。\n"
    "各入力の全文は、このあとの「入力①」から「入力④」のページにあります。",
    "",
    """【目安】60秒。実演の前に、この表で順序と代わりの手段を確認します。
【使い方】PowerPointのノートペイン／発表者ツールでこのページを開きます。各入力の全文は、本編の「入力①」から「入力④」のページ（編集できる文字）と、このノートにあります。
【実行順】Q1 → Q2 → Q3 → Q4。Q4はNotebookの前後どちらでも実行できます。画面の切替を減らしたい場合は、Q3の直後にDAでQ4を実行してからNotebookへ移ります。
【うまくいかないとき】Q1とQ2は、グラフの崩れに対する追い指示が、各入力の参考結果のページにあります。Q3は、Notebookが動かない場合に保存済みのHTMLを使い、「保存時点の結果です」と断ります。Q4は、参考結果のページを使い、ライブの回答とは呼びません。
【追加の質問】時間があれば、ENE_POSの好影響（全額+210,155,445.28円、約+2.10億円）と、CROSS_NEGとCROSS_POSの業種別の比較を聞きます。Step03はFOD_NEG専用なので、イベントIDを変えてエネルギーの図は作りません。横断2件は、primary_sector='all'を落とさずに会社の業種で分けます。負と正の時期が違うので、回復率とは説明しません。
【表示の経験】現在の環境ではMermaidはコードブロックのまま、Pythonによる任意DAG画像の生成表示も利用できない旨の回答でした。DA一般や将来の機能を一律に否定する意味ではありません。棒グラフは色や横向きを強く指定すると失敗し、制約を外した縦棒が成功しました。短い項目名を使います。グラフの成功は、毎回確認が必要です。
【実演先】Fabric DAの公開済み画面です。M365 CopilotはDA結果を再構成するため、同じグラフが通過表示される保証はありません。このデモではM365側の表示成功を主張しません。
【Q1全文】\n""" + Q1 + "\n\n【Q2全文】\n" + Q2 + "\n\n【Q3全文】\n" + Q3 + "\n\n【Q4全文】\n" + Q_ACT
    + "\n\n【追加A全文】\n" + Q4 + "\n\n【追加B全文】\n" + Q5 + "\n\n【DA直接リンク】\n" + DA,
    ["prompts/demo_questions_ja.txt", "prompts/presenter_prompts_ja.txt", "prompts/da_instructions_ja.txt",
     "説明者の現在環境での表示確認（色と配置は未固定）"],
    "迷ったら①②③の順。うまくいかないときは、保存済みの参考結果に切り替え、ライブの回答とは呼ばない")
input_rows = [
    ["①", "全社の稼ぎは、基準に対してどう動いた？", "月次のbaselineとactual、\n年間の全額と持分管理", "DA", "必須",
     "期間合計を外し、12か月、\n縦軸を億円にして作り直す"],
    ["②", "何が稼ぎを押し上げ、押し下げた？", "10イベントの符号付きの影響。\n食品損失は約−1.14億円", "DA", "必須",
     "色や縦横の指定を外す"],
    ["③", "食品損失は、どの記録、取引、\n契約につながる？", "金額とIDをDAで確認し、\n関係の図をNotebookで見せる", "DAと\nNotebook", "必須",
     "Notebookが動かなければ、\n保存済みHTMLを「保存時点」と断って使う"],
    ["④", "どの契約から手を打つ？", "契約別の影響額と、優先度つきの対応案", "DA", "任意",
     "参考結果のページを使う\n（ライブの回答とは呼ばない）"],
    ["追加", "エネルギーの好影響は？\n横断の2件を業種別に見ると？", "ENE_POSの+2.10億円と、\n横断2件の業種別比較", "DA", "任意",
     "時間がなければ実行しない"],
]
grid_table(s, .62, 2.58, [.6, 3.3, 3.1, 1.05, .6, 3.43],
           ["入力", "営業向けの聞き方", "見せること", "画面", "区分", "うまくいかないとき"],
           input_rows, [.68, .68, .68, .68, .68], size=11,
           colors=[TEAL, INK, INK, INK, CORAL, MUTED], bolds=[True, True, False, False, True, False],
           cell_colors={(3, 4): MUTED, (4, 4): MUTED})
pill(s, "Fabric DAを開く", 10.1, .3, 2.55, DA)


# 16
s = new_slide(
    "DAの金額を照合する正解表",
    "DAが示した10イベントの金額が合っているかは、この表で確認します。\n"
    "億円の表示は丸めた値なので、厳密に照合するときは円の列を使います。",
    "",
    """【準備用】ライブ回答が異なる場合の照合表です。取得に失敗した回答を、この表で差し替えて成功したとは言いません。
【使う場面】入力②の回答、棒グラフの数値、短いラベルとevent_idの対応を確認するとき。入力③の食品損失（FOD_NEG）も、この表の金額と一致します。
【読み方】各行はeffect_idを一度だけ集計した全額JPYのdeltaです。actualに追加する金額ではありません。固有イベントは各84効果、横断イベントは各336効果、合計1,344効果です。
【DAの金額が違うとき】①対象期間が2025-10-01から2026-09-30か。②全額連結か持分管理か。③primary_sector='all'の横断2件（CROSS_NEGとCROSS_POS）を除外していないか。④同じeffect_idを二重に合計していないか。この4点を順に確認します。
【短いラベル】製錬停止=smelter outage、再生合金=recycled alloy premium contract、LNG遅延=LNG terminal maintenance delay、配分最適化=storage dispatch optimization、食品損失=cold chain spoilage incident、調達改善=grain yield and procurement improvement、与信費増=counterparty servicing cost shock、ヘッジ手数料=hedge execution fee uplift、横断混乱=cross-sector port and liquidity disruption、横断回復=cross-sector logistics and treasury recovery。いずれも架空の設計です。
【重複防止】evidenceをevent_idだけで結合すると金額が根拠件数倍になります。effect.evidence_idで結ぶか、効果を重複なく確定してから合計します。日次指標を活動件数倍に加算することも避けます。
【食品の持分管理】-87,630,294.31円。表の全額-114,461,186.25円とは指標が異なります。
【日付】CROSS_NEGは2026年5月、CROSS_POSは9月のイベントです。両者の比率を施策の回復率や同条件の効果比較とは呼びません。
【与信費増の訳】FIN_NEGの原文はcounterparty servicing cost shockで、取引先対応コストの急増を指す設計です。ラベルの「与信費増」は短くした呼び名で、貸倒関連の与信費用を意味しません。""",
    ["reports/qa_expected_event_totals.json", "data/demo_dataset.json:semantic_risk_events"],
    "固有イベントは84効果×8件、横断は336効果×2件。合計1,344効果")
colxs = [.75, 2.5, 4.2, 6.75, 7.95]
widths = [1.7, 1.7, 2.5, 1.2, .8]
for x, w, label in zip(colxs, widths, ["event_id", "短いラベル", "全額影響（円）", "億円", "効果数"]):
    text(s, label, x, 2.66, w, .3, 13, TEAL, True)
for i, (eid, label) in enumerate(zip(order, labels)):
    y = 3.05 + i * .33
    if i % 2 == 0:
        box(s, .65, y - .03, 8.35, .33, WHITE)
    amount = Decimal(events[eid]["full_jpy"])
    vals = [eid, label, f"{amount:+,.2f}", f"{amount / Decimal(10**8):+.4f}", str(events[eid]["effects"])]
    for x, w, val in zip(colxs, widths, vals):
        text(s, val, x, y, w, .28, 12.5, CORAL if amount < 0 else INK, eid == "FOD_NEG")
box(s, 9.3, 2.6, 3.4, 3.78, WHITE, rounded=True)
text(s, "DAの金額が違うとき", 9.52, 2.72, 3.0, .3, 15, TEAL, True)
stack_items(s, 9.52, 3.12, 3.0, [
    ("1 対象期間", "2025-10-01から2026-09-30か"),
    ("2 全額か持分か", "全額と持分を取り違えていないか"),
    ("3 横断の2件", "横断の2件を除外していないか"),
    ("4 二重の合計", "同じeffect_idを2回足していないか"),
], PURPLE, size=11.5, gap=.1)


# 17
s = new_slide(
    "想定質問への短い回答",
    "細部を推測して答えると、実装が確認した範囲を超えてしまいます。\n短い回答のあと、必要な場合だけノートの補足を使います。",
    "",
    """【Q：SQLでもできる？】A：「この集計だけなら、SQLでも同じ数字を出せます。違いは、意味とつながりを一度定義して、質問やツールをまたいで使い回せることです。」詳しくは、次のページ（SQLとOntologyの違い）を使います。OntologyだけがJOINできる、とは言いません。
【Q：線があるなら原因と証明できる？】A：「できません。合成のログと設計した配賦を結んでいます。線には関連や集約先の意味もあります。」現実で因果を論じるには別のデータや検証設計が必要です。
【Q：51%なら半分だけ連結？】A：「このデモは全社が支配子会社なので、全額指標には100%を取り込みます。比率を掛けた方は別の持分管理寄与です。」正式な会計処理一式を再現しているわけではありません。
【Q：食品の温度が原因？】A：「観測値は-18℃ですが、許容温度や閾値は未記録です。原因として断定せず、架空の設計例として説明します。」
【Q：今の画面は最新？】A：「表示しているのは読取日時付きの保存結果です。最新のGraphを読む場合は再実行し、完了と表示日時を確認します。」GraphModel自体の更新時点も別です。
【Q：金融にも出荷がある？】A：「金融はexposureや手数料の活動を使います。存在しない物理出荷は補完しません。」
【Q：なぜマイナスの損失とプラスのKPIが同じ図に？】A：「前者はイベントの差分、後者は他の取引も含む当日の総額だからです。」詳しくは、第3章の「4つの金額は範囲が違う」を使います。
【Q：Notebookはデータを書き換える？】A：「原本とGraphは読取専用です。自Notebookの定義には、実行したHTML/SVGと監査出力を保存します。」全操作が完全な無書込みという説明もしません。
【Q：イベントのデータは社内にない？】A：「整った形では、ないのが普通です。差分の理由コメントや事故や契約の記録、業務データの異常から拾い、小さな台帳から始めます。」詳細は第5章にあります。
【Q：取引先の名前は出せる？】A：「このデモのデータに取引先名はありません。契約IDで特定し、実務では取引先マスタを契約につなぎます。」
【Q：日本語は公式にサポートされている？】A：「公式には、英語以外は現在サポートされていません（Microsoft Learn、2026-10-21更新）。このデモは日本語で入力し、日本語で答えが返ることを確認しています。」詳細は、第1章の「実演前の確認事項」にあります。
【Q：全社に波及するケースは扱える？】A：「全社に同時に及ぶ形は扱えます。横断の2件が、8社すべてに当たります。会社から会社へ連鎖する形は、会社同士を結ぶ関係をまだ持っていないので、このデモでは扱えていません。」詳細は、""" + RIPPLE_PAGES + """にあります。
【Q：もっとカジュアルに聞ける？】A：「短い質問でも答えられるよう、言い回しの読み替えをDAの指示文に書きました。ただし、LLMなので毎回同じ答えにはなりません。」詳細は、""" + CASUAL_PAGES + """にあります。
【Q：全部ゼロコピー？】A：「いいえ。ゼロコピーは業種拡張8表のショートカットです。分析用の投影とGraphModelは、コピーして作っています。」詳しくは、第6章の「ゼロコピーと実体化の使い分け」を使います。
【Q：業務アプリでは、Ontologyが金額を計算しているの？】A：「いいえ。つながりはOntologyのグラフから取得し、金額はLakehouseの値を、決まった手順で計算しています。」詳細は、p.""" + str(PAGE_OF[T_APP_QA]) + """にあります。""",
    ["docs/architecture.md", "reports/step3_live_audit.json",
     "scripts/step3_food_loss_runtime.py",
     "Microsoft Learn：fabric/data-science/concept-data-agent（2026-10-21更新）"],
    "答えが未確認なら「未検証」と明示。温度の閾値や因果を補完しない")
faqs = [
    ("SQLでもできる？", "集計は可能。違いは、意味とつながりを使い回せること（次ページ）"),
    ("矢印は因果の証明？", "いいえ。関連、設計上の配賦、集約先の区別"),
    ("51%なら半分を連結？", "全額は100%で集計。持分管理は別の指標"),
    ("温度が損失の原因？", "許容温度や閾値は未記録。観測値から断定しない"),
    ("開くだけで最新？", "保存時点の結果。更新は再実行して、日時を確認"),
    ("イベントのデータは社内にない？", "差分の理由や事故の記録から拾い、小さく始める（第5章）"),
    ("取引先名は出せる？", "データに取引先名はない。契約IDで特定する"),
    ("日本語は公式サポート？", "公式には英語以外は未サポート。このデモは日本語で動作確認済み"),
    ("全社に波及するケースは？", f"同時に全社へ及ぶ形は扱える。連鎖は未対応（{RIPPLE_PAGES}）"),
    ("もっとカジュアルに聞ける？", f"短い質問も可。ただし結果は毎回変わりうる（{CASUAL_PAGES}）"),
]
for i, (q, a) in enumerate(faqs):
    y = 2.58 + i * .39
    box(s, .65, y, 12.04, .34, WHITE, rounded=True)
    text(s, q, .86, y + .045, 4.0, .26, 14.5, TEAL, True)
    text(s, a, 4.9, y + .055, 7.7, .26, 13.5)

# 17b SQL and ontology: how to answer "can't SQL do this?"
s = new_slide(
    "SQLとOntologyの違い",
    "このデモの集計だけなら、SQLでも同じ数字を出せます。\n"
    "差が出やすいのは、質問、利用者、ツールが増えて、定義を使い回す場面です。",
    "",
    """【目安】90秒。「SQLでもできるのでは」「何が違うのか」と聞かれたときに使います。
【話す】「このデモの集計だけなら、SQLでも同じ数字を出せます。SQLのビューやPower BIのセマンティックモデルで、共通化することもできます。そのうえで、Ontologyが加えるのは、業務の言葉、つながり、複数のデータ元への対応付けを、ひとつの定義として持たせることです。人にも、エージェントにも、同じ定義を渡せます。」
【言葉の意味】ビューやセマンティックモデルでは、意味が、列名、計算式、メジャーに持たされます。利用者は、表、列、メジャーの意味を読み解いて使います。Ontologyでは、実体、関係、ルールとして、業務の言葉で定義します。
【つながり】ビュー、クエリ、セマンティックモデルのリレーションに、結合条件を書きます。表をたどる質問では、利用者やツールが、結合の意味を理解して使います。条件を誤ると、金額が二重に数えられます（このデモでも、根拠をevent_idだけで結ぶと金額が根拠件数倍になる落とし穴を、指示文に書いています）。Ontologyでは、関係に名前を付けて定義します（このデモは39種類）。
【複数のデータ元】複数のデータ元をまたぐ共通化は、データ元ごと、ツールごとの作業になりやすくなります。Ontologyでは、実体の定義を、複数のデータ元の表に対応付けられます（Learn：Ontologyは、データの上の仮想的な意味と文脈の層で、データはそれぞれの元の場所にあります）。エージェントに同じ文脈を渡せますが、対応状況は、利用するエージェントごとに違います（Learn：Availability and supported functionality can vary by agent experience）。
【既存のセマンティックモデルとの関係】OntologyはPower BIのセマンティックモデルから生成でき、既存のテーブル、リレーション、計算列、メジャーを引き継げます。セマンティックモデルを置き換えるものではなく、その上に業務の概念と関係を重ねる層です（Learn）。
【公式の説明】Microsoft Learn（Ontology概要、2026-10-16更新）は、利用者ごとに、表、列、メジャー、結合を解釈する代わりに、再利用できる業務の概念とデータへの対応付けを与えることを、Ontologyの役割と説明しています。また、エージェントには、生のスキーマやプロンプトだけに頼らず、定義、関係、メトリック、ルール、データへの対応付けで文脈を与えられる、と説明しています。
【このデモの現状】Ontologyが担っているのは、実体と関係の定義（17実体、39関係）と、そこから作ったGraphです。全額連結などの言葉の定義の多くは、いまはDAの指示文に書いてあり、Ontologyのルールへ移す余地があります（第7章）。「Ontologyが、このデモの言葉の定義をすべて担っている」とは言いません。
【避ける】「SQLでは不可能」「OntologyがJOINを置き換える」と言いません。定義を直したとき、参照する側にどう反映されるかは、バインディングの再適用やグラフの更新が要る場合があります（第6章）。このデモの4つの質問だけを見れば、差は小さいと認めます。""",
    ["Microsoft Learn：fabric/iq/ontology/overview（2026-10-16更新）",
     "reports/ontology_runtime_current.json", "prompts/da_instructions_ja.txt"],
    "答え方：SQLでも作れると認め、質問、利用者、ツールが増えたときの、定義の使い回しを示す")
sql_rows = [
    ["言葉の意味",
     "意味は、列名、計算式、メジャーに持たせる。利用者は、表、列、メジャーの意味を読み解いて使う。",
     "実体、関係、ルールとして、業務の言葉で定義する。\n人もエージェントも、同じ定義を参照する。"],
    ["つながり",
     "結合条件は、ビュー、クエリ、セマンティックモデルのリレーションに書く。表をたどる質問では、利用者やツールが、結合の意味を理解して使う。",
     "関係に名前を付けて定義する（このデモは39種類）。\n経路をたどる質問を、定義した関係で扱える。"],
    ["複数のデータ元",
     "複数のデータ元をまたぐ共通化は、データ元ごと、ツールごとの作業になりやすい。",
     "実体の定義を、複数のデータ元の表に対応付けられる。\nエージェントに同じ文脈を渡せる\n（対応状況はエージェントごとに違う）。"],
]
grid_table(s, .62, 2.58, [1.75, 5.2, 5.13], ["違い", "SQLのビューやセマンティックモデルで共通化した場合", "Ontologyで定義した場合"],
           sql_rows, [.93, .93, .93], size=12.5, colors=[TEAL, INK, INK], bolds=[True, False, False])
box(s, .62, 5.78, 12.08, .66, PALE, rounded=True)
text(s, "このデモの現状：Ontologyが担っているのは、実体と関係の定義（17実体、39関係）とGraphです。\n"
        "全額連結などの言葉の定義の多くは、いまはDAの指示文にあり、Ontologyのルールへ移す余地があります（第7章）。",
     .82, 5.84, 11.7, .52, 11.5, NAVY)


# Group-wide ripple, two pages after the SQL page. The first shows how far one event reaches, the second
# separates the forms of ripple the demo covers from the ones it does not. Every figure is computed from the
# saved data and asserted against the verified reports; nothing is typed in by hand.
graph_paths = {r["event_id"]: int(r["effects_with_complete_path"]) for r in json.loads(
    (PROJECT / "reports/graph_all_event_paths.json").read_text())["result"]["data"]}
sector_of = {s["entity_id"]: s["sector_code"] for s in subsidiaries}
assert [s["sector_code"] for s in subsidiaries] == ["metals"] * 2 + ["energy"] * 2 + ["foods"] * 2 + ["finance"] * 2
reach = {}
for _eid in ("FOD_NEG", "CROSS_NEG", "CROSS_POS"):
    _rows = [r for r in dataset["semantic_event_financial_effects"] if r["event_id"] == _eid]
    reach[_eid] = {
        "companies": {r["entity_id"] for r in _rows}, "contracts": {r["contract_id"] for r in _rows},
        "activities": {r["activity_id"] for r in _rows}, "effects": len(_rows),
        "full": sum((Decimal(r["delta_ebitda_full_jpy"]) for r in _rows), Decimal(0)),
        "mgmt": sum((Decimal(r["delta_management_jpy"]) for r in _rows), Decimal(0))}
    assert reach[_eid]["full"] == Decimal(events[_eid]["full_jpy"])
    assert reach[_eid]["mgmt"] == Decimal(events[_eid]["management_jpy"])
    assert reach[_eid]["effects"] == events[_eid]["effects"] == graph_paths[_eid]
    assert len(reach[_eid]["activities"]) == reach[_eid]["effects"]
assert [len(reach["FOD_NEG"][k]) for k in ("companies", "contracts")] == [2, 6]
assert [len(reach["CROSS_NEG"][k]) for k in ("companies", "contracts")] == [8, 26]
assert reach["CROSS_NEG"]["companies"] == reach["CROSS_POS"]["companies"] == set(sector_of)
assert (reach["FOD_NEG"]["effects"], reach["CROSS_NEG"]["effects"]) == (84, 336)
assert sum(graph_paths.values()) == env_counts["effects"]


# Optional live demo of the Fabric App (Rayfin) and the questions it draws. Every figure comes from the saved
# live-check report and is asserted against numbers this deck already derives; nothing is typed in by hand.
app = json.loads((PROJECT / "reports/whatif_app_live_verification_2026-10-28.json").read_text())
APP_SC = app["scenarios"]
APP_BOTH = APP_SC["food_loss_50_and_yen_up_10"]
assert app["status"] == "verified_live_with_limits" and app["environment"]["identifiers_recorded"] is False
assert app["probe"]["reconciliation"]["ok"] is True
APP_BASE = {k_: Decimal(v_) for k_, v_ in app["probe"]["base_annual_ebitda_jpy"].items()}
assert APP_BASE["full"] == actual_total["EBITDA"].quantize(Decimal("0.01"))
APP_FOOD_ONLY = Decimal(APP_SC["food_loss_50"]["delta_full_jpy"])
APP_STEPS = {k_: Decimal(v_) for k_, v_ in APP_BOTH["steps_full_jpy"].items()}
APP_FOD = app["graph_paths_per_event"]["FOD_NEG"]
assert (APP_FOD["FinancialEffect"], APP_FOD["Activity"], APP_FOD["Contract"], APP_FOD["FinancialMetric"],
        APP_FOD["GroupKPI"]) == (84, 84, 6, 28, 15)
assert APP_FOD["FinancialEffect"] == reach["FOD_NEG"]["effects"] and APP_FOD["Contract"] == len(reach["FOD_NEG"]["contracts"])
assert "59" in app["ui_observations"]["first_open_after_idle_seconds_observed"]
assert "63" in app["ui_observations"]["first_open_after_idle_seconds_observed"]
assert "Sign in required" in app["ui_observations"]["signed_out_access"]
assert APP_STEPS["events"] + APP_STEPS["fx"] == Decimal(APP_BOTH["delta_full_jpy"])


def app_oku(v, sign=True):
    """Hundred-million yen with one decimal, rounded half up, the way the app screen shows it."""
    q = (Decimal(v) / OKU).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
    return (f"{q:+.1f}" if sign else f"{q:.1f}").replace("-", "−")


def app_mil(v):
    """Million yen with one decimal and an explicit sign (the unit of the app's breakdown)."""
    q = (Decimal(v) / Decimal(10) ** 6).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
    return f"{q:+,.1f}".replace("-", "−")


A_BASE_FULL, A_BASE_MGMT = app_oku(APP_BASE["full"], False), app_oku(APP_BASE["management"], False)
A_FOOD = app_oku(APP_FOOD_ONLY)
A_BOTH_FULL, A_BOTH_MGMT = app_oku(APP_BOTH["delta_full_jpy"]), app_oku(APP_BOTH["delta_management_jpy"])
A_EV, A_FX, A_TOT = app_mil(APP_STEPS["events"]), app_mil(APP_STEPS["fx"]), app_mil(APP_BOTH["delta_full_jpy"])
assert (A_BASE_FULL, A_BASE_MGMT, A_FOOD, A_BOTH_FULL, A_BOTH_MGMT) == ("447.6", "335.5", "+0.6", "−22.7", "−13.7")
assert (A_EV, A_FX, A_TOT) == ("+57.2", "−2,331.8", "−2,274.6")
APP_FOOT = "架空の業務データ／公開ECB為替を使用  |  アプリの確認：2026-10-28（Fabric AppsとOntologyはPreview、クエリ用APIはbeta）"
APP_SOURCES = ["reports/whatif_app_live_verification_2026-10-28.json",
               "Microsoft Learn：fabric/apps/overview（2026-09-29更新）、fabric/apps/functions（2026-09-29更新）",
               "Microsoft Learn：fabric/graph/gql-query-api（2026-09-29更新）"]


def dot(slide, label, cx, cy, r, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    text(slide, label, cx - r, cy - r + .035, 2 * r, 2 * r, 13, WHITE, True, align=PP_ALIGN.CENTER)


# Chapter 3, the last page: an optional live demo after the 10 minutes
s = new_slide(
    T_APP_DEMO,
    "10分の実演のあとに、画面を持つアプリでの使い方を求められる場合があります。\n"
    "5つの操作で、つながりはOntology、金額はLakehouseと計算、という分担を、画面で示せます。",
    "",
    f"""【目安】5分（任意）。10分の実演のあとで、お客様が「画面を持つアプリで使えるのか」に関心を示したときだけ使います。時間がなければ、お客様向け資料の第4章「業務アプリで使う」の実画面のページで代えます。
【前準備】アプリは、見せる直前に開き、経路の表示まで待ちます。放置したあとの最初の呼び出しは、グラフの起動に約1分かかります（2026年10月27日と28日の実測、59秒から63秒）。2回目以降は、グラフへの問い合わせが約1.1〜1.6秒でした（一度だけ6.1秒）。アプリがFunctionsに保持するデータは、約10分で入れ替わります（アプリの実装の値）。グラフ側が、どのくらい放置すると再び遅くなるかは、未確認です。10分の実演の前に開いておくと、見せるときに、再び待つ場合があります。サインインが必要で、未サインインでは、アプリの内容にもFunctionsにも届かないことを確認済みです。利用者ごとの見え方は、未検証です。
【手順】①画面を開き、実績を読みます（全額連結{A_BASE_FULL}億円、持分管理{A_BASE_MGMT}億円、差は0）。画面の「実績」は、変えない状態の値で、第2章で見た実績と同じです。基準値（予算に相当する値）とは別です。②「食品損失を 50% 軽減」を押します（全額連結の差は{A_FOOD}億円）。③続けて「円高 10%」を押します（全額連結{A_BOTH_FULL}億円、持分管理{A_BOTH_MGMT}億円）。④「差の内訳」を指します（イベント{A_EV}、為替{A_FX}、合計の差{A_TOT}。いずれも百万円）。⑤「Ontologyの経路」で、イベントに食品損失（FOD_NEG）を選びます（影響額{APP_FOD['FinancialEffect']}件、取引活動{APP_FOD['Activity']}件、契約{APP_FOD['Contract']}件、日次指標{APP_FOD['FinancialMetric']}件、全社KPI{APP_FOD['GroupKPI']}件。「影響額の件数が一致しています」と表示されます）。⑥終わったら「実績に戻す」を押します。
【話す】「条件を画面で動かすと、差額と内訳がその場で変わります。つながりはOntologyのグラフから取得し、金額はLakehouseの値を、決まった手順で計算しています。」
【待ち時間が出たら】最初の画面に、「オントロジーのグラフと Lakehouse に接続しています…」と、「使われていなかった後の初回は、グラフの起動に最大 1 分ほどかかります。」（経過秒数つき）が出ます。イベントを選び直したときは、経路の欄に「グラフに問い合わせています…（初回は最大 1 分ほどかかります）」と出ます。慌てずに、「グラフの起動を待っています」と伝え、その間に、お客様向け資料の第4章の最初のページ（役割分担）を使って、つながりと金額の担当の違いを話します。
【応答しないとき】再試行を1回だけ押します。直らなければ、お客様向け資料の第4章の実画面と経路のページ（保存した画像）に切り替え、「保存した画像で、いまのライブの結果ではありません」と告げます。権限の変更や再デプロイを、その場で試みません。
【避ける】「Ontologyが金額を計算した」「本番で使える」「利用者ごとに見える範囲を変えられる」とは言いません。円高10%は、USD建ての子会社の円換算レートだけを10%下げる仮定で、為替ヘッジ、価格への転嫁、取引量の変化は反映していません。金額は、設計した配賦と仮定の為替による試算で、実際の損益の予測ではありません。
【現状】確認した日は2026年10月28日です。Fabric AppsとOntologyはPreview、グラフのクエリ用APIはbetaです。アプリのURLは、この資料に載せていません。事前に、説明者のブックマークに保存します。
【つなぎ】「アプリについての質問は、第4章の最後から2ページ目（p.{PAGE_OF[T_APP_QA]}）に、答え方をそろえてあります。」続けて「来期の見通しは」と聞かれたら、将来タブの実演（p.{PAGE_OF[T_APP_FDEMO]}）に進みます。""",
    APP_SOURCES,
    f"質問が出たら、p.{PAGE_OF[T_APP_QA]}の答え方で答える。見通しを聞かれたら、p.{PAGE_OF[T_APP_FDEMO]}へ", foot=APP_FOOT)
box(s, .62, 2.62, 7.62, 3.78, WHITE, rounded=True)
text(s, "手順（約5分）", .87, 2.74, 5.0, .3, 16, TEAL, True)
for i, (act, shows) in enumerate([
        ("画面を開き、実績を読む", f"全額連結 {A_BASE_FULL}億円、持分管理 {A_BASE_MGMT}億円。差は0"),
        ("「食品損失を 50% 軽減」を押す", f"全額連結の差は {A_FOOD}億円"),
        ("続けて「円高 10%」を押す", f"全額連結 {A_BOTH_FULL}億円、持分管理 {A_BOTH_MGMT}億円"),
        ("「差の内訳」を指す", f"イベント {A_EV}、為替 {A_FX}、合計の差 {A_TOT}（百万円）"),
        ("経路で、イベントに食品損失を選ぶ",
         f"影響額{APP_FOD['FinancialEffect']}件、契約{APP_FOD['Contract']}件。「影響額の件数が一致」と表示")]):
    y = 3.2 + i * .62
    dot(s, str(i + 1), 1.1, y + .26, .2, TEAL)
    text(s, act, 1.5, y, 6.5, .27, 14, NAVY, True)
    text(s, shows, 1.5, y + .28, 6.5, .27, 12, INK)
box(s, 8.42, 2.62, 4.28, 3.78, WHITE, rounded=True)
text(s, "準備と、困ったとき", 8.67, 2.74, 3.8, .3, 16, TEAL, True)
stack_items(s, 8.67, 3.2, 3.8, [
    ("開くタイミング", "見せる直前に開き、経路の表示まで待つ。空きが長いと、再び待つ場合がある"),
    ("待ち時間が出たら", "「グラフの起動を待っています」と伝え、お客様向け資料の役割分担のページで補う"),
    ("応答しないとき", "再試行は1回だけ。直らなければ、保存した実画面に切り替え、保存した画像と告げる"),
    ("見られる人", "サインインが必要。利用者ごとの見え方は未検証"),
], TEAL, size=11.5, gap=.12)

# Chapter 4, the last page: the questions the app draws
s = new_slide(
    T_APP_QA,
    "アプリの画面を見せたあとは、「Ontologyが計算したのか」という質問が出やすくなります。\n"
    "確認した範囲を超えて答えないよう、6つの答えをそろえておきます。",
    "",
    f"""【目安】質疑のとき。必要な質問だけを、1つ30秒ほどで答えます。業務アプリの実演（第3章の最後から2ページ目、p.{PAGE_OF[T_APP_DEMO]}）のあとに使います。予測や将来の仮定の質問は、p.{PAGE_OF[T_APP_FQA]}にあります。
【Q：Ontologyが金額を計算したの？】A：「いいえ。つながり（経路）は、Ontologyのグラフから取得しています。金額は、Lakehouseの値を、決まった手順で計算しています。」理由：Data Agentの円高の試算は、実行ごとの答えが安定しなかったので、条件を動かす用途では、金額を言語モデルに任せていません。
【Q：FabricのアプリからOntologyに、標準でつなげる？】A：「現在のコネクタの対象は、Lakehouse、Warehouse、SQLデータベース、セマンティックモデルで、Ontologyは含まれません（Microsoft Learn、2026-09-29更新）。このデモでは、FunctionsからGraphのクエリ用API（GQL Query API）を直接呼んでいます。」このAPIはbetaです。
【Q：利用者ごとに、見える範囲を変えられる？】A：「Functionsが外部を呼ぶときのIDは、現在、アプリの所有者です（Learn）。画面を見る利用者の権限は、反映されません。利用者ごとに変える設計は、このデモでは検証していません。PoCの最初に決めます。」お客様向け資料の第3章の権限のページは、Ontologyのアイテムとデータ側の権限の話で、別です。
【Q：最初に開くとき、遅いのはなぜ？】A：「放置したあとの最初の呼び出しは、グラフの起動に約1分かかりました（実測59秒から63秒）。2回目以降は、グラフへの問い合わせが約1.1〜1.6秒でした。」放置の時間の目安は、未確認です。
【Q：本番で使える？】A：「Fabric AppsとOntologyはPreview、グラフのクエリ用APIはbetaです。Learnは、クエリ用APIの本番利用を推奨していません。性能、容量の消費量、コストは未確認で、PoCで確認します。」
【Q：東日本のリージョンで使える？】A：「Fabric Appsはリージョンが限定的です（Learn）。このデモはWest US 3で確認しました。東日本など、他のリージョンは未確認です。」
【避ける】「Ontologyが金額を計算した」「本番で使える」「利用者の権限が自動で効く」とは言いません。性能やコストの数値を、推測で答えません。答えが未確認なら、「未検証」と明示します。""",
    APP_SOURCES,
    "未確認の質問は、お客様向け資料の確認範囲のページ（第4章末）で、どの列かを示す", foot=APP_FOOT)
app_qa = [
    ("Ontologyが金額を計算したの？",
     "いいえ。つながり（経路）はOntologyのグラフから取得し、金額はLakehouseの値を、決まった手順で計算しています。"),
    ("FabricのアプリからOntologyに、標準でつなげる？",
     "標準のコネクタの対象はLakehouseなどで、Ontologyは含まれません（Learn）。このデモはGraphのクエリ用API（beta）を直接呼んでいます。"),
    ("利用者ごとに、見える範囲を変えられる？",
     "呼び出しは、アプリの所有者のIDです（Learn）。画面を見る利用者の権限は反映されず、利用者ごとの設計は未検証です。PoCの最初に決めます。"),
    ("最初に開くとき、遅いのはなぜ？",
     "放置後の最初の呼び出しは、グラフの起動に約1分（実測59〜63秒）。2回目以降は、グラフへの問い合わせが約1.1〜1.6秒です。"),
    ("本番で使える？",
     "Fabric AppsとOntologyはPreview、クエリ用APIはbetaです（Learnは、クエリ用APIの本番利用を推奨していません）。性能とコストは、PoCで確認します。"),
    ("東日本のリージョンで使える？",
     "Fabric Appsはリージョンが限定的です（Learn）。このデモはWest US 3で確認しました。東日本など、他のリージョンは未確認です。"),
]
for i, (q, a) in enumerate(app_qa):
    y = 2.58 + i * .65
    box(s, .62, y, 12.08, .6, WHITE, rounded=True)
    text(s, q, .85, y + .05, 4.1, .52, 14.5, TEAL, True)
    text(s, a, 5.1, y + .06, 7.4, .5, 12.5, INK)



# Optional live demo of the future tab (forecast + hypothetical event) and the questions it draws. Every figure comes
# from the saved live-check report and is tied to numbers this deck already derives; nothing is typed in by hand.
fut = json.loads((PROJECT / "reports/whatif_future_live_verification_2026-10-28.json").read_text())
assert fut["status"] == "verified_live_with_limits" and fut["identifiers_recorded"] is False
FUT_BASE = Decimal(fut["forecast"]["baseline_jpy"])
FUT_SC = fut["scenarios"]
FUT_A, FUT_B, FUT_C = FUT_SC["fod_neg_nov15_100"], FUT_SC["fod_neg_dec20_100"], FUT_SC["fod_neg_dec31_100"]
assert (FUT_A["event_id"], FUT_A["event_date"], FUT_A["scale"]) == ("FOD_NEG", "2026-11-15", 1)
assert (FUT_B["event_date"], FUT_B["included_effects"], FUT_B["excluded_effects"]) == ("2026-12-20", 57, 27)
assert (FUT_C["event_date"], FUT_C["included_effects"], Decimal(FUT_C["delta_jpy"])) == ("2026-12-31", 0, 0)
for _sc in (FUT_A, FUT_B, FUT_C):
    assert Decimal(_sc["baseline_jpy"]) == FUT_BASE
    assert Decimal(_sc["scenario_jpy"]) - FUT_BASE == Decimal(_sc["delta_jpy"])
    assert _sc["included_effects"] + _sc["excluded_effects"] == _sc["historical_effects"] == 84
# the food-loss effect that the 11/15 scenario replays is the one this deck reports for FOD_NEG, to the yen
assert Decimal(FUT_A["delta_jpy"]) == Decimal(events["FOD_NEG"]["full_jpy"]).quantize(Decimal("0.01")) == Decimal("-114461186.25")
assert sum(Decimal(q["total_jpy"]) for q in fut["actual_by_quarter"]) == Decimal(fut["annual_actual_ebitda_full_jpy"]) == APP_BASE["full"]
assert all(e["ids_match_exactly"] and e["graph_paths"] == e["lakehouse_effects"] == graph_paths[e["eventId"]]
           for e in fut["graph_vs_lakehouse_per_event"]) and len(fut["graph_vs_lakehouse_per_event"]) == 10
_FM = {m["model"]: m for m in fut["notebook"]["models"]}
assert (_FM["AutoML + Graph"]["q4_error_percent"], _FM["AutoML"]["q4_error_percent"], _FM["直近91日日平均"]["q4_error_percent"]) == (
    1.6, -5.44, -1.73)
assert Decimal(str(_FM["AutoML + Graph"]["next_q_forecast_oku"])) == (FUT_BASE / OKU).quantize(Decimal("0.001")) == Decimal("114.895")
assert fut["ui_observations"]["console_errors"] == 0 and "25" in fut["ui_observations"]["first_open_seconds_observed"]
_FNV = " ".join(fut["not_verified"])
assert all(w in _FNV for w in ("一般化", "二重計上", "将来の為替", "持分管理", "権限", "性能"))


def fut_oku(v, places="0.01", sign=False):
    q = (Decimal(v) / OKU).quantize(Decimal(places), rounding=ROUND_HALF_UP)
    return (f"{q:+.{len(places.split('.')[1])}f}" if sign else f"{q:.{len(places.split('.')[1])}f}").replace("-", "−")


F_BASE2 = fut_oku(FUT_BASE)
F_AFTER_A = fut_oku(FUT_A["scenario_jpy"])
F_DELTA_A = fut_oku(FUT_A["delta_jpy"], sign=True)
F_DELTA_B = fut_oku(FUT_B["delta_jpy"], sign=True)
assert (F_BASE2, F_AFTER_A, F_DELTA_A, F_DELTA_B) == ("114.90", "113.75", "−1.14", "−0.93")
F_ERR = {k_: f"{Decimal(str(v_['q4_error_percent'])).quantize(Decimal('0.01')):+.2f}".replace("-", "−") for k_, v_ in _FM.items()}
assert F_ERR == {"直近91日日平均": "−1.73", "AutoML": "−5.44", "AutoML + Graph": "+1.60"}
FUT_FOOT = "架空の業務データ／公開ECB為替を使用  |  将来タブの確認：2026-10-28（予測は探索的。Fabric AppsとOntologyはPreview、クエリ用APIはbeta）"
FUT_SOURCES = ["reports/whatif_future_live_verification_2026-10-28.json",
               "Microsoft Learn：fabric/data-science/how-to-use-automated-machine-learning-fabric（2026-07-22更新）",
               "Microsoft Learn：fabric/apps/overview（2026-09-29更新）、fabric/graph/gql-query-api（2026-09-29更新）"]

# Chapter 3, the last page: an optional demo of the future tab, three minutes
s = new_slide(
    T_APP_FDEMO,
    "「来期の見通しは」と続けて聞かれたときに使う、任意の3分です。\n"
    "予測は保存済みの結果で、イベントは仮定です。Ontologyが予測した、とは言いません。",
    "",
    f"""【目安】3分（任意）。業務アプリの実演（p.{PAGE_OF[T_APP_DEMO]}）のあとで、「来期の見通しは」「このイベントが起きたら」と聞かれたときだけ使います。時間がなければ、お客様向け資料の第4章「次の四半期の予測にOntologyの関係を使う」と「将来のイベントを仮定して予測と比べる」の2ページで代えます。
【前準備】アプリは、見せる直前に開き、「将来の予測」タブを一度開いておきます。モデルの読み込み後、将来タブの初回の表示は約25〜30秒でした（2026年10月28日の実測）。放置したあとの最初の呼び出しは、グラフの起動に約1分かかります（p.{PAGE_OF[T_APP_DEMO]}の【前準備】と同じ）。予測は、Notebookが保存した結果を読んでいるだけで、画面を開くたびに予測を計算し直してはいません。
【手順】①「将来の予測」タブを開きます。折れ線は、実績4四半期（FY26 Q1〜Q4。年度は10月始まり）と、予測の1四半期（FY27 Q1。点線）です。棒は、業種別の積み上げで、右の軸（0から始まる別の目盛り）です。基準予測は{F_BASE2}億円です。②タブを開いたときの初期値は、「過去イベントのパターン」が「コールドチェーンの劣化（食品ロス）」、「仮の発生日」が2026年11月15日、「追加する影響の大きさ」が100%です。「イベント想定後」は{F_AFTER_A}億円、「予測との差」は{F_DELTA_A}億円です（四捨五入のため、引き算とは0.01億円ずれます）。この差は、第2章で見た食品損失の全期間の合計と同額ですが、同じ84件の影響額を足し合わせているので当然で、予測の良し悪しを確かめた結果ではありません。「照合 84 件。期間内 84 件、期間外 0 件」と表示されます。③日付を2026年12月20日に変えます。差は{F_DELTA_B}億円になり、四半期に収まる影響額は{FUT_B['included_effects']}件、{FUT_B['excluded_effects']}件は四半期の外です。④日付を2026年12月31日にすると、差は0になります（収まる影響額は0件）。画面で選べる日付は、予測の四半期の最終日である12月31日までです。⑤「参照した過去のOntology経路」を開いて、過去の同じイベントの影響額と、Graphの経路を指します。
【話す】「基準予測は、Notebookで動かしたAutoMLにGraphの情報を加えた方法が、1年分のデータで学習し直して出した、次の四半期の保存済みの値です。そこに、過去に起きた同じイベントの影響を、仮定した日に移して足しています。影響額は、Ontologyのグラフの経路と、Lakehouseの値を、IDで突き合わせて取り出しています。Ontologyが予測したわけではありません。」
【待ち時間が出たら】将来タブの初回は、表示まで25〜30秒ほどかかります。画面に「Notebookの保存済み予測を読み込み中…」と出ます。「保存済みの予測を読み込んでいます」と伝え、その間に、お客様向け資料の第4章「次の四半期の予測にOntologyの関係を使う」のページで、AutoMLとOntologyの担当の違いを話します。
【応答しないとき】再試行を1回だけ押します。直らなければ、お客様向け資料の第4章の2ページ（保存した画像）に切り替え、「保存した画像で、いまのライブの結果ではありません」と告げます。権限の変更や再デプロイを、その場で試みません。
【避ける】「Ontologyが予測した」「予測は当たる」「Graphを使うと精度が上がる」「この見通しは確定」とは言いません。逆に、「Graphは精度に効かない」とも言いません。Graphの効果を聞かれたら、「合成データなので、Graphを加えると精度が上がるかは、貴社のデータを使ったPoCでの実証が必要です」と答えます（p.{PAGE_OF[T_APP_FQA]}）。予測は、合成データの最後の1四半期の1回で検証しただけで、その四半期の実績は、比べる前にすでに見ています（探索的）。基準予測には、過去のイベントの影響がすでに含まれている可能性があり（二重計上）、将来の為替は予測していません。イベントも予測せず、利用者が仮定します。全額連結の金額だけに対応し、持分管理には対応していません。
【現状】確認した日は2026年10月28日です。Fabric AppsとOntologyはPreview、グラフのクエリ用APIはbetaです。アプリのURLは、この資料に載せていません。
【つなぎ】「予測や将来の仮定についての質問は、第4章の最後のページ（p.{PAGE_OF[T_APP_FQA]}）に、答え方をそろえてあります。」""",
    FUT_SOURCES,
    f"質問が出たら、p.{PAGE_OF[T_APP_FQA]}の答え方で答える。予測の精度は「未確認」、Graphの効果は「PoCでの実証が必要」と言う", foot=FUT_FOOT)
box(s, .62, 2.62, 7.62, 3.78, WHITE, rounded=True)
text(s, "手順（約3分）", .87, 2.74, 5.0, .3, 16, TEAL, True)
for i, (act, shows) in enumerate([
        ("「将来の予測」タブを開く", f"基準予測 {F_BASE2}億円。折れ線は実績4四半期と予測1四半期"),
        ("初期値のまま（食品ロス、11月15日、100%）", f"イベント想定後 {F_AFTER_A}億円。予測との差 {F_DELTA_A}億円"),
        ("日付を2026年12月20日に変える", f"差 {F_DELTA_B}億円。期間内 {FUT_B['included_effects']}件、期間外 {FUT_B['excluded_effects']}件"),
        ("日付を2026年12月31日にする", "差は0。四半期に収まる影響額が0件"),
        ("「参照した過去のOntology経路」を開く", "過去の同じイベントの影響額を、Graphの経路とLakehouseで照合")]):
    y = 3.2 + i * .62
    dot(s, str(i + 1), 1.1, y + .26, .2, TEAL)
    text(s, act, 1.5, y, 6.5, .27, 14, NAVY, True)
    text(s, shows, 1.5, y + .28, 6.5, .27, 12, INK)
box(s, 8.42, 2.62, 4.28, 3.78, WHITE, rounded=True)
text(s, "準備と、困ったとき", 8.67, 2.74, 3.8, .3, 16, TEAL, True)
stack_items(s, 8.67, 3.2, 3.8, [
    ("開くタイミング", "見せる直前に開き、将来タブを一度表示しておく。初回は約25〜30秒"),
    ("待ち時間が出たら", "「保存済みの予測を読み込んでいます」と伝え、お客様向け資料の予測のページで補う"),
    ("応答しないとき", "再試行は1回だけ。直らなければ、保存した画像に切り替え、保存した画像と告げる"),
    ("言わないこと", "「Ontologyが予測した」「予測は当たる」「Graphで精度が上がる」"),
], TEAL, size=11.5, gap=.12)

# Chapter 4, the last page: the questions the forecast and the future tab draw
s = new_slide(
    T_APP_FQA,
    "予測を見せたあとは、「当たるのか」「Ontologyは何をしたのか」という質問が出やすくなります。\n"
    "確認した範囲を超えて答えないよう、6つの答えをそろえておきます。",
    "",
    f"""【目安】質疑のとき。必要な質問だけを、1つ30秒ほどで答えます。将来の実演（第3章の最後のページ、p.{PAGE_OF[T_APP_FDEMO]}）のあとに使います。
【Q：予測は当たるの？】A：「確認できていません。検証は、デモのデータの最後の1四半期の1回だけで、その四半期の実績は、比べる前にすでに見ています（探索的）。最初の3四半期で学習し、最後の四半期を予測させたとき、誤差は、直近91日の日平均が{F_ERR['直近91日日平均']}%、AutoMLが{F_ERR['AutoML']}%、AutoMLにGraphの情報を加えたものが{F_ERR['AutoML + Graph']}%でした。貴社のデータで、複数の四半期を使って確かめます。」データは合成データで、季節性などの形も、設計したものです。
【Q：Ontologyが予測しているの？／Graphを使うと精度が上がるの？】A：「予測を計算しているのは、AutoMLです。Ontologyのグラフは、過去のイベントが、どの影響額、会社、日の指標につながるかを返します。Notebookが、その結果とLakehouseの金額から、予測の入力を作ります。Graphを加えると精度が上がるかは、今回のデータが合成データのため、このデモでは判断できません。貴社のデータを使ったPoCでの実証が必要です。」デモの比べ方では、選ばれたモデルの種類も異なり（AutoMLはrf、Graphを加えたものはextra_tree）、Graphの有無だけを比べた結果ではありません。「Graphに効果がない」という意味でもありません。PoCでは、モデルの種類と学習の期間をそろえ、Graphの情報の有無だけを変えて、複数の四半期で比べる進め方が考えられます。
【Q：イベントを足すと、二重に数えない？】A：「その可能性があります。基準予測は、過去の実績から作っており、過去のイベントの影響が、すでに含まれている可能性があります。そこに、同じ種類のイベントの影響を足すと、二重に数えるおそれがあります。分離は確かめていません。」
【Q：為替や、将来のイベントは予測しているの？】A：「いいえ。将来の為替は予測していません。イベントも、予測せず、利用者が仮定します。仮定は、過去の同じイベントの影響額を、指定した発生日に移し、倍率を掛けて、予測の四半期に収まる分だけを足したものです。」
【Q：なぜ3四半期で学習したの？】A：「デモのデータが1年分（4四半期）だけなので、方法の比べ方では、最初の3四半期で学習し、最後の1四半期を検証に使いました。アプリの基準予測は、4四半期すべてで学習し直した値です（学習の行数は、検証用が176、基準予測用が288）。1年分のデータでは、季節性を複数年にわたって確かめられません。学習に使う92日の期間は、週ごとに重なっており、独立した観測の数ではありません。」
【Q：持分管理や、他の指標（売上など）も出せる？】A：「将来タブは、全額連結のEBITDAだけに対応しています。持分管理は、未対応です。売上などの指標は、このデモでは扱っていません。」
【避ける】「予測は当たる」「Graphで精度が上がる」「Ontologyが予測した」「二重計上はない」とは言いません。また、「Graphは精度に効かない」とも言いません。Graphの効果は、「合成データのため判断できず、PoCでの実証が必要」と言います。精度の数字（誤差のパーセント）を、一般的な性能として示しません。答えが未確認なら、「未検証」と明示します。""",
    FUT_SOURCES,
    "未確認の質問は、お客様向け資料の確認範囲のページ（第4章末）と、予測のページで、どの列かを示す", foot=FUT_FOOT)
fut_qa = [
    ("予測は当たるの？",
     f"未確認です。検証は1四半期の1回で、実績は先に見ています（探索的）。誤差：日平均 {F_ERR['直近91日日平均']}%、AutoML {F_ERR['AutoML']}%、AutoML + Graph {F_ERR['AutoML + Graph']}%。"),
    ("Ontologyが予測しているの？\nGraphで精度は上がる？",
     "予測を計算するのはAutoMLです。Graphは過去のイベントの影響を返します。Graphで精度が上がるかは、合成データでは判断できず、PoCでの実証が必要です。"),
    ("イベントを足すと、二重に数えない？",
     "その可能性があります。基準予測に、過去のイベントの影響が含まれている場合があります。分離は未確認です。"),
    ("為替や将来のイベントは予測している？",
     "いいえ。将来の為替は予測せず、イベントは利用者が仮定します（過去の影響を、仮定した日に移して足します）。"),
    ("なぜ3四半期で学習したの？",
     "検証用に、最初の3四半期で学習しました。アプリの値は1年分で学習し直しています。1年分では季節性を確かめられません。"),
    ("持分管理や売上も出せる？",
     "将来タブは、全額連結のEBITDAだけです。持分管理は未対応で、売上などは、このデモでは扱っていません。"),
]
for i, (q, a) in enumerate(fut_qa):
    y = 2.58 + i * .65
    box(s, .62, y, 12.08, .6, WHITE, rounded=True)
    text(s, q, .85, y + .05, 4.1, .52, 14.5, TEAL, True)
    text(s, a, 5.1, y + .06, 7.4, .5, 12.5, INK)



def oku(v, sign=True):
    return (f"{v / OKU:+.2f}" if sign else f"{v / OKU:.2f}").replace("-", "−")


def pct_of_annual(v):
    return f"{v / actual_total['EBITDA'] * 100:.2f}%".replace("-", "−")


ebitda_by_entity, cross_by_entity = defaultdict(Decimal), defaultdict(Decimal)
for r in dataset["semantic_financial_metrics_daily"]:
    if r["scenario_tag"] == "actual" and r["metric_code"] == "EBITDA":
        ebitda_by_entity[r["entity_id"]] += Decimal(r["amount_group"])
for r in dataset["semantic_event_financial_effects"]:
    if r["event_id"] == "CROSS_NEG":
        cross_by_entity[r["entity_id"]] += Decimal(r["delta_ebitda_full_jpy"])
assert sum(cross_by_entity.values()) == reach["CROSS_NEG"]["full"]
_shares = [-cross_by_entity[k] / ebitda_by_entity[k] * 100 for k in cross_by_entity]
share_lo, share_hi = min(_shares), max(_shares)
assert Decimal("1.4") <= share_lo and share_hi <= Decimal("1.6")
cross_pct_1 = f"{-reach['CROSS_NEG']['full'] / actual_total['EBITDA'] * 100:.1f}"
cross_vs_food = reach["CROSS_NEG"]["full"] / reach["FOD_NEG"]["full"]
mgmt_ratio = reach["CROSS_NEG"]["mgmt"] / reach["CROSS_NEG"]["full"] * 100


def sector_line(eid, field):
    sums = defaultdict(Decimal)
    for r in dataset["semantic_event_financial_effects"]:
        if r["event_id"] == eid:
            sums[sector_of[r["entity_id"]]] += Decimal(r[field])
    return "、".join(f"{SECTOR_JP[k]} {oku(sums[k])}" for k in ("metals", "energy", "foods", "finance"))


interco = dataset["intercompany_transactions"]
interco_pairs = len({r["pair_id"] for r in interco})
assert (len(interco), interco_pairs) == (48, 24)
assert {(r["entity_id"], r["counterparty_entity_id"]) for r in interco if r["side"] == "sale"} == {
    ("SUB_FIN_JP_01", "SUB_ENE_JP_01"), ("SUB_FIN_US_01", "SUB_ENE_US_01")}
interco_oku = sum((Decimal(r["amount_group"]) for r in interco if r["side"] == "sale"), Decimal(0)) / OKU
assert not [r for r in ontology_relationships if r["from"] == "Company" and r["to"] == "Company"]
assert "intercompany_transactions" in dataset
_usd = [r for r in dataset["semantic_financial_metrics_daily"]
        if r["scenario_tag"] == "actual" and r["metric_code"] == "EBITDA" and r["currency_code"] == "USD"]
usd_local = sum((Decimal(r["amount_local"]) for r in _usd), Decimal(0))
usd_group = sum((Decimal(r["amount_group"]) for r in _usd), Decimal(0))
fx10 = usd_local * 10

# The Data Agent was asked these questions on the published version on 2026-10-27. The slides say the
# returned amounts equal the saved data, so every value is compared here at the same rounding.
wide = json.loads((PROJECT / "reports/da_group_wide_2026-10-27.json").read_text())
wide_runs = {r["id"]: r for r in wide["runs"]}
live_seconds = {k: v["response_seconds"] for k, v in wide_runs.items()}
# The same eight questions were asked again after the ontology relationships were set active.
fixed = json.loads((PROJECT / "reports/da_traversal_fix_2026-10-27.json").read_text())
fixed_runs = {r["id"]: r for r in fixed["runs"]}


def wide_rows(run_id, column):
    ex = next(x for x in reversed(wide_runs[run_id]["sql_executions"]) if x["is_success"] and column in x["columns"])
    assert ex["rows_total"] == len(ex["rows"])
    return [dict(zip(ex["columns"], row)) for row in ex["rows"]]


def oku2(v, places="0.01"):
    return (v / OKU).quantize(Decimal(places), rounding=ROUND_HALF_UP)


_sector_full, _sector_mgmt = defaultdict(Decimal), defaultdict(Decimal)
_cross_mgmt, _cross_contracts = defaultdict(Decimal), defaultdict(set)
for r in dataset["semantic_event_financial_effects"]:
    _sector_full[r["event_id"], sector_of[r["entity_id"]]] += Decimal(r["delta_ebitda_full_jpy"])
    _sector_mgmt[r["event_id"], sector_of[r["entity_id"]]] += Decimal(r["delta_management_jpy"])
    if r["event_id"] == "CROSS_NEG":
        _cross_mgmt[r["entity_id"]] += Decimal(r["delta_management_jpy"])
        _cross_contracts[r["entity_id"]].add(r["contract_id"])
_sector_rows = wide_rows("B_sector", "sector_code")
assert len(_sector_rows) == 8
for r in _sector_rows:
    k = (r["event_id"], r["sector_code"])
    assert Decimal(str(r["delta_ebitda_full_oku_jpy"])) == oku2(_sector_full[k])
    assert Decimal(str(r["delta_management_oku_jpy"])) == oku2(_sector_mgmt[k])
    assert r["distinct_effect_count"] == 84
_entity_of_name = {s["entity_name"]: s["entity_id"] for s in subsidiaries}
_company_rows = wide_rows("company", "subsidiary_name")
assert len(_company_rows) == len(reach["CROSS_NEG"]["companies"]) == 8
assert sum(r["impacted_contract_count"] for r in _company_rows) == len(reach["CROSS_NEG"]["contracts"])
for r in _company_rows:
    e = _entity_of_name[r["subsidiary_name"]]
    assert Decimal(str(r["full_consolidated_impact_oku_jpy"])) == oku2(cross_by_entity[e])
    assert Decimal(str(r["management_weighted_impact_oku_jpy"])) == oku2(_cross_mgmt[e])
    assert r["impacted_contract_count"] == len(_cross_contracts[e])
    assert Decimal(str(r["impact_pct_of_annual_actual_full_ebitda"])) == (
        cross_by_entity[e] / ebitda_by_entity[e] * 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
_largest = wide_rows("casual_largest_event", "event_id")[0]
assert _largest["event_id"] == "CROSS_NEG"
assert Decimal(str(_largest["delta_ebitda_full_jpy_oku"])) == oku2(reach["CROSS_NEG"]["full"], "0.0001")
assert Decimal(str(_largest["annual_ebitda_full_consolidated_oku"])) == oku2(actual_total["EBITDA"], "0.001")
assert Decimal(str(_largest["event_impact_to_annual_ebitda_percent"])) == (
    reach["CROSS_NEG"]["full"] / actual_total["EBITDA"] * 100).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
_fx_row = wide_rows("fx_what_if_1", "yen_appreciation_10_jpy_sensitivity_oku_jpy")[0]
assert Decimal(str(_fx_row["yen_appreciation_10_jpy_sensitivity_oku_jpy"])) == oku2(-fx10, "0.0001")
assert "算定できません" in wide_runs["fx_what_if_2"]["answer"]
_traversals = [v["ontology_traversal"] for v in wide_runs.values() if isinstance(v["ontology_traversal"], list)]
assert len(_traversals) == 6 and all(x[0]["chosenPath"] == "no_path_found" for x in _traversals)
_fixed_traversals = [v["ontology_traversal"] for v in fixed_runs.values() if isinstance(v["ontology_traversal"], list)]
assert len(_fixed_traversals) == 6
_max_hops = max(len(t["hops"]) for v in _fixed_traversals for t in v)
assert all(t["found"] and t["chosenPath"] == "single_engine_join" for v in _fixed_traversals for t in v)
assert all(f"{oku2(fx10)}億円" in fixed_runs[k]["answer"] for k in ("fx_what_if_1", "fx_what_if_2"))
assert all("内部取引" in fixed_runs[k]["answer"] for k in ("chain_1", "chain_2"))
for _rid in ("chain_1", "chain_2"):
    assert "内部取引" in wide_runs[_rid]["answer"] or "intercompany" in wide_runs[_rid]["answer"]
fx_live = -Decimal(str(_fx_row["yen_appreciation_10_jpy_sensitivity_oku_jpy"]))

s = new_slide(
    T_RIPPLE,
    "実演で追う食品損失は、影響が一部の会社に限られた出来事です。\n"
    "聞かれやすいのは、全社に同時に及ぶ出来事も、同じ経路で追えるかどうかです。",
    "",
    f"""【目安】90秒。「全社に及ぶ話も扱えるのか」と聞かれたときに使います。ライブの10分には含めません。
【話す】「実演で追った食品損失は、影響が2社に限られた出来事です。このデモには、全社に同時に及ぶ出来事も2件入れてあります。港の混乱と資金繰りの悪化が、業種を問わず起きた想定の『横断混乱』です。8社すべてに当たり、EBITDAを約{-reach['CROSS_NEG']['full'] / OKU:.1f}億円押し下げます。食品損失と同じ経路で、根拠から全社KPIまでたどれます。」
【指す】左の表を、上から読みます。会社は2社から8社、影響明細は84件から336件に増えています。次に、右の図で、8社すべてに影響が出ている点と、各社の年間EBITDAに対する比率が約{share_lo:.1f}〜{share_hi:.1f}%でそろっている点を指します。
【全額と持分管理】全額は、8社を100%で足した額（{oku(reach['CROSS_NEG']['full'])}億円）です。持分管理は、所有比率を掛けた額（{oku(reach['CROSS_NEG']['mgmt'])}億円）で、51%所有の4社の分が小さくなるため、全額の約{mgmt_ratio:.0f}%です。
【年間EBITDA比】年間のEBITDA（実績、{oku(actual_total['EBITDA'], False)}億円）に対する比率です。横断混乱は{pct_of_annual(reach['CROSS_NEG']['full'])}、食品損失は{pct_of_annual(reach['FOD_NEG']['full'])}で、前者は後者の約{cross_vs_food:.1f}倍です。ただし、横断混乱でも年間の約{cross_pct_1}%で、これより大きな規模の出来事は、このデモに収録していません。この割合を「小さい」「軽微」とは評価しません。重大性を判定する基準は、このデモにありません。次のページで、収録していない範囲も伝えます。
"""
    f"""【数字の出どころ】
DAが返したもの：横断2件を含む全10イベントの、イベントごとの金額（横断混乱 {reach['CROSS_NEG']['full']:+,.2f}円、横断回復 {reach['CROSS_POS']['full']:+,.2f}円）。さらに、2026-10-27の実測で、業種別、会社別、年間EBITDA比、最大のイベントも、公開済みDAが返しました。
Graphで確認したもの：全{env_counts['effects']:,}効果が、根拠から全社KPIまでの経路でたどれること。横断混乱は{reach['CROSS_NEG']['effects']}件すべてです。2026-10-23のGraph更新後に確認しました。
保存データから集計したもの：食品損失の年間EBITDA比（{pct_of_annual(reach['FOD_NEG']['full'])}）。DAには、年間EBITDA比を、最大のイベント（横断混乱）についてだけ聞きました。
【2026-10-27の実測（修正前の8回）】公開済みDAで、チャットを空にして、1問ずつ聞きました（記録：reports/da_group_wide_2026-10-27.json）。①追加B（横断2件の業種別比較、{live_seconds['B_sector']}秒）、②横断混乱の会社別（年間EBITDA比つき、{live_seconds['company']}秒）、③「全社に同時に響いたマイナスの出来事は？会社ごとにどれくらい影響した？」（{live_seconds['casual_all_companies']}秒）、④「このデータで、グループのEBITDAにいちばん大きく響いた出来事は？年間のEBITDAの何%くらい？」（{live_seconds['casual_largest_event']}秒）です。表の金額は、すべて保存データの集計と、小数第2位まで一致しました。③は、聞き方が短くても、横断混乱を選び、8社を返しました。④は、横断混乱の{oku(reach['CROSS_NEG']['full'])}億円、年間{oku(actual_total['EBITDA'], False)}億円の{pct_of_annual(reach['CROSS_NEG']['full'])}と答えました。
【DAの文章の注意】表の数字は一致しましたが、表のまわりの文章には誤りが出ることがあります。追加Bの答えは、持分管理の金融と食品の大小を説明しましたが、横断混乱の値（食品{oku(_sector_mgmt['CROSS_NEG', 'foods'])}億円、金融{oku(_sector_mgmt['CROSS_NEG', 'finance'])}億円）と合いませんでした。会社別の答えは、DAが自分で置いた基準（1億円超、または年間EBITDA比−1.55%超）で、優先度が「高」の会社を挙げましたが、基準に当てはまる3社が漏れていました。この基準は、デモの定義ではありません。実演では、表の数字を指して説明し、補足の文章は読み上げません。
【DAの経路】この日の最初の8回では、DAがOntologyの経路探索を試した6回すべてで「経路なし」が返り、答えは、Lakehouseの表へのSQLで求めたものでした。原因は、Ontologyの42のリレーションシップがすべて無効に設定されていたことです。DAに渡る文脈で、どの関係も結合条件が空欄でした。\
同じ日に、リレーションシップを有効にし、接続情報を再適用し、Graphを更新したあと、同じ8問を聞き直しました。経路探索を試した{len(_fixed_traversals)}回すべてで経路が見つかり（最大で{_max_hops}つの結合をたどる道筋）、金額は変わりませんでした（記録：reports/da_traversal_fix_2026-10-27.json）。\
ただし、答えは今もLakehouseの表へのSQLで求めたものです。Ontologyは、どの表をどの列で結ぶかの道筋をDAに示す役目です。「DAがGraphをたどって答えた」とは言いません。関係をたどる場面は、Step03のNotebook（Graph）で見せます。
【照合用の業種別の値（億円）】横断混乱の全額：{sector_line('CROSS_NEG', 'delta_ebitda_full_jpy')}。持分管理：{sector_line('CROSS_NEG', 'delta_management_jpy')}。横断回復の全額：{sector_line('CROSS_POS', 'delta_ebitda_full_jpy')}。持分管理：{sector_line('CROSS_POS', 'delta_management_jpy')}。丸めのため、足すと合計と0.01ずれることがあります。
【避ける】横断回復÷横断混乱を、回復率と呼ぶこと。別の時期に設計した、独立のイベントです。各社の年間EBITDA比がそろっているのは、データの設計によるもので、現実の波及の姿を示すものではありません。実世界の因果の証明でもありません。
【つなぎ】「同時に全社へ及ぶ形は、ここまでです。次のページで、連鎖する形や、大きな規模の想定まで扱えるかを整理します。」
【コピー用Q5】\n""" + Q5,
    ["reports/qa_expected_event_totals.json", "reports/graph_all_event_paths.json",
     "reports/da_qa_published.json", "reports/da_group_wide_2026-10-27.json",
     "reports/da_traversal_fix_2026-10-27.json",
     "data/demo_dataset.json:semantic_event_financial_effects、semantic_financial_metrics_daily（actual）",
     "config/event_generation_rules.yaml:event_mix"],
    "公開済みDAで聞いた業種別と会社別の金額は、元データの集計と一致した（2026-10-27）。説明では表の数字を示す")
reach_rows = [
    ["イベントID", "FOD_NEG", "CROSS_NEG"],
    ["影響を受けた会社", f"{len(reach['FOD_NEG']['companies'])}社（食品の日米）",
     f"{len(reach['CROSS_NEG']['companies'])}社（4業種すべて）"],
    ["契約", f"{len(reach['FOD_NEG']['contracts'])}件", f"{len(reach['CROSS_NEG']['contracts'])}件"],
    ["活動と影響明細", f"{reach['FOD_NEG']['effects']}件", f"{reach['CROSS_NEG']['effects']}件"],
    ["全額EBITDA", f"{oku(reach['FOD_NEG']['full'])}億円", f"{oku(reach['CROSS_NEG']['full'])}億円"],
    ["持分管理", f"{oku(reach['FOD_NEG']['mgmt'])}億円", f"{oku(reach['CROSS_NEG']['mgmt'])}億円"],
    ["年間EBITDA比", pct_of_annual(reach["FOD_NEG"]["full"]), pct_of_annual(reach["CROSS_NEG"]["full"])],
    ["全社KPIまでの経路", f"{reach['FOD_NEG']['effects']}件すべてたどれる",
     f"{reach['CROSS_NEG']['effects']}件すべてたどれる"],
]
grid_table(s, .62, 2.58, [1.7, 2.05, 2.3], ["比べる点", "食品損失", "横断混乱"], reach_rows, .42, size=11.5,
           colors=[TEAL, INK, INK], bolds=[True, False, True],
           fills=[WHITE if i % 2 == 0 else PALE for i in range(len(reach_rows))],
           cell_colors={(i, j): CORAL for i in (4, 5, 6) for j in (1, 2)})
chart_x0, chart_w, zero_y, scale = 7.0, 5.7, 3.5, 2.0
text(s, "横断混乱：会社別の全額EBITDA影響（億円）", chart_x0, 2.62, chart_w, .28, 13, TEAL, True)
text(s, f"各社の年間EBITDA比は約{share_lo:.1f}〜{share_hi:.1f}%です。DAの実測と一致しました。",
     chart_x0, 2.94, chart_w, .24, 11, MUTED)
box(s, chart_x0, zero_y, chart_w, .012, MUTED)
col_w = chart_w / 8
for i, sub in enumerate(subsidiaries):
    v = cross_by_entity[sub["entity_id"]]
    h = abs(float(v / OKU)) * scale
    x = chart_x0 + i * col_w + (col_w - .46) / 2
    box(s, x, zero_y + .012, .46, h, CORAL)
    text(s, oku(v), x - .12, zero_y + .012 + h + .03, .7, .22, 10, CORAL, True, PP_ALIGN.CENTER)
    text(s, "日本" if sub["country_code"] == "JP" else "米国", x - .12, 5.98, .7, .2, 10, INK, False,
         PP_ALIGN.CENTER)
for k, sector in enumerate(("metals", "energy", "foods", "finance")):
    gx = chart_x0 + k * 2 * col_w
    box(s, gx + .08, 6.2, 2 * col_w - .16, .012, GRAY)
    text(s, SECTOR_JP[sector], gx, 6.24, 2 * col_w, .22, 11, NAVY, True, PP_ALIGN.CENTER)

s = new_slide(
    T_RIPPLE_SCOPE,
    "「全社に波及する」が指す出来事は、聞き手によって違います。\n"
    "指すものが違えば、このデモで見せられるかどうかの答えも変わります。",
    "",
    f"""【目安】90秒。「全社に波及するリスクも扱えるのか」「連鎖はどうなのか」と聞かれたときに使います。
【話す】「『全社に波及する』には、大きく2つの形があります。1つは、港の混乱のように、全社に同時に当たる形です。これはこのデモに入っていて、食品損失と同じ経路で追えます。もう1つは、ある会社の問題が、取引を通じて別の会社に伝わる連鎖です。こちらは今のデモでは扱えていません。会社と会社を結ぶ関係を、Ontologyにまだ持っていないからです。」
【指す】表を上から指します。「このデモ」の列で、扱えるのは1行目だけです。2行目以降は、足りない点を、そのまま伝えます。
【連鎖が扱えない理由】Ontologyの関係は、{len(ontology_entities)}種類の実体の間の{len(ontology_relationships)}種類で、会社同士を結ぶものはありません。内部取引の表（{len(interco)}行、{interco_pairs}組）は親Lakehouseにあり、DAはSQLで読めますが、Ontologyの関係にはしていません。内部取引は、金融の会社とエネルギーの会社の間の2経路（日本同士、米国同士）だけで、売り手側の年間の額は約{interco_oku:.1f}億円です。横断の2件は、8社の各活動に直接配賦した記録で、1社の問題が別の社へ伝わる過程を計算したものではありません。
【DAの実測：連鎖】2026-10-27に、「食品の損失が、ほかの子会社にも連鎖した影響はある？グループ内の取引を通じた影響も見たい。」と、チャットを空にして2回聞きました（{live_seconds['chain_1']}秒、{live_seconds['chain_2']}秒）。どちらも、食品以外の子会社への配賦はなく、内部取引経由の影響は定量できないと答えました。\
1回目は、内部取引経由の影響を「未検証」としました。2回目は、内部取引の表に活動と契約の対応がなく、食品損失の効果に結べないためと説明し、「影響なし」とは判定しませんでした。ただし、2回目は、食品以外の子会社について、対象0件、0.00億円と表に出しました。\
聞き手には、「データに配賦がない」という意味だと添えます。Ontologyのリレーションシップを有効にした後の聞き直し（2回、{fixed_runs['chain_1']['response_seconds']}秒、{fixed_runs['chain_2']['response_seconds']}秒）も、同じ結論でした。内部取引の表に、食品の会社の行がないためです（金融とエネルギーの組だけ）。\
「連鎖は起きない」とは言いません。
"""
    f"""【Microsoft Learnの記述】Graphは、複数のつながりを順にたどる問い、依存関係の分析、到達可能性、最短経路の問いに使う実行層です。集計や指標は、元のデータでの実行を続けられます（Ontology概要、2026-10-16更新）。関係型は、始点と終点の実体の種類を指定して作り、対応表で結べます（関係型の作成、2026-09-04更新）。同じ種類の実体の間（会社と会社）の関係を作れるかは、この資料では確認していません。実行する前に、Previewの画面で確認します。
【連鎖を扱う場合の一案】内部取引の表を、会社と会社の関係として足し、金融の会社の問題が、取引でつながるエネルギーの会社へ及ぶ様子を、2社の小さな範囲で試します。ただし、伝わる大きさを決める規則とデータは、いまありません。関係を足すだけでは、金額は出ません。
【規模】最大の横断混乱で、全額{oku(reach['CROSS_NEG']['full'])}億円、年間EBITDA（実績）の約{cross_pct_1}%です。これより大きな規模の出来事は、収録していません。8社それぞれの年間EBITDAに対する比率は、約{share_lo:.1f}〜{share_hi:.1f}%でそろっています。会社ごとの差が大きい例や、より大きな規模の例を見せる場合は、影響データの追加とGraphの更新が要ります。実体や関係の種類を変える場合は、Ontologyの定義も更新します。DAに「このデータで、グループのEBITDAにいちばん大きく響いた出来事は？年間のEBITDAの何%くらい？」と聞くと、横断混乱の{oku(reach['CROSS_NEG']['full'])}億円、年間{oku(actual_total['EBITDA'], False)}億円の{pct_of_annual(reach['CROSS_NEG']['full'])}と答えました（2026-10-27の実測）。
【もしもの試算】2026-10-27に、「もし円が10円高くなったら、グループのEBITDAはいくら変わる？」と、チャットを空にして2回聞きました（{live_seconds['fx_what_if_1']}秒、{live_seconds['fx_what_if_2']}秒）。1回目は、ドル建てEBITDAに−10円を掛けた単純な感応度で、{oku(-fx_live * OKU)}億円と答えました（保存データからの概算と同じ値です）。\
2回目は、「ショック後のレートや感応度の定義がデータにないため、算定できない」と答え、ドル建てEBITDAの実績（約{usd_local / 10**6:.0f}百万ドル、{usd_group / OKU:.2f}億円）だけを示しました。\
Ontologyのリレーションシップを有効にした後に、同じ質問をさらに2回聞くと（{fixed_runs['fx_what_if_1']['response_seconds']}秒、{fixed_runs['fx_what_if_2']['response_seconds']}秒）、2回とも{oku(-fx_live * OKU)}億円と答えました。同じ日の4回のうち3回が数字、1回が算定不可で、答えが分かれるので、実演では使いません。\
聞かれたら、「このデモのDAは、記録済みのデータを集計して答えます。もしもの計算は、前提を決めて別に行います」と答えます。\
目安は、米国の4社のEBITDAが年間で約{usd_local / 10**6:.0f}百万ドル（全額で約{usd_group / OKU:.0f}億円、全社の{usd_group / actual_total['EBITDA'] * 100:.0f}%）なので、円高が10円進むと、換算だけで約{fx10 / OKU:.0f}億円（全社の{fx10 / actual_total['EBITDA'] * 100:.1f}%）の押し下げです。ヘッジや価格への転嫁は考慮していません。\
数字を出すときは「概算」と断ります。
【避ける】「全社波及もすべて扱える」「連鎖も追える」と言いません。DAが答えた円高の概算を、シミュレーションの結果として紹介することも、しません。年間の約{cross_pct_1}%を「小さい」「軽微」と評価することも、しません（重大性を判定する基準は、このデモにありません）。「Ontologyでは連鎖を扱えない」とも言いません。会社間の関係を足せば扱える見込みはありますが、このデモでは未検証です。いずれも設計した配賦で、実世界の因果の証明ではありません。
【つなぎ】「お客様のデータで連鎖を見たい場合は、会社間の取引表から、小さく始めます。進め方は、第5章にあります。」""",
    ["config/subsidiaries.json", "ontology_tmdl/entityRelationships.tmdl（会社同士の関係なし）",
     "data/demo_dataset.json:intercompany_transactions、semantic_financial_metrics_daily（actual）",
     "config/data_agent_create.json（DAが参照する表）", "reports/graph_all_event_paths.json",
     "reports/da_group_wide_2026-10-27.json", "reports/da_traversal_fix_2026-10-27.json",
     "Microsoft Learn：fabric/iq/ontology/overview（2026-10-16更新）",
     "Microsoft Learn：fabric/iq/ontology/how-to-create-relationship-types（2026-09-04更新）"],
    "聞き返す例：「全社に同時に及ぶ話ですか。会社から会社へ連鎖する話ですか」")
scope_rows = [
    ("同時に全社へ及ぶ", "例：港の混乱、資金繰りの悪化", "扱える", TEAL,
     "横断の2件が8社に同時に当たり、全社KPIまでの経路もすべてたどれます。\n"
     "DAは、イベント別、業種別、会社別の金額を、保存データと同じ値で返しました（実測）。"),
    ("会社から会社へ連鎖する", "例：金融の問題がエネルギーに及ぶ", "未対応", CORAL,
     "会社同士の関係がOntologyになく、内部取引の表も金融とエネルギーの2経路だけです。\n"
     "DAは4回とも「他の子会社への配賦なし、内部取引経由は確認できない」と答えました。\n"
     "表の0.00億円は、データに配賦がないことを示し、影響なしの確認にはなりません。"),
    ("もっと大きな規模の出来事", "例：全社の稼ぎが大きく動く出来事", "未収録", PURPLE,
     f"最大の横断混乱で、年間EBITDAの約{cross_pct_1}%です。これより大きな出来事は、収録していません。\n"
     "大きな出来事の追加には、影響データの追加とGraphの更新が要ります。"),
    ("もしもの試算", "例：円高が10円進んだら", "要注意", MUTED,
     f"同じ日に4回聞くと、3回は「円高10円で全額EBITDAは約{fx_live:.1f}億円減る」と単純計算で答え、\n"
     "1回は「感応度の定義がない」と答え、算定しませんでした。実演では使いません。"),
]
text(s, "聞かれ方", .8, 2.58, 2.6, .26, 12, TEAL, True)
text(s, "このデモ", 3.55, 2.58, 1.2, .26, 12, TEAL, True)
text(s, "根拠と、足りない点", 4.95, 2.58, 7.5, .26, 12, TEAL, True)
ry = 2.92
for label, example, status, color, body in scope_rows:
    h = max(.72, .1 + est_lines(body, 7.6, 11.5, .9) * 11.5 * 1.25 / 72 + body.count("\n") * .07 + .12)
    box(s, .62, ry, 12.08, h, WHITE, rounded=True)
    text(s, label, .8, ry + .1, 2.6, .27, 13, NAVY, True)
    text(s, example, .8, ry + .4, 2.65, .22, 10.5, MUTED)
    box(s, 3.55, ry + (h - .36) / 2, 1.05, .36, color, rounded=True)
    text(s, status, 3.55, ry + (h - .36) / 2 + .07, 1.05, .26, 13, WHITE, True, PP_ALIGN.CENTER)
    text(s, body, 4.95, ry + .1, 7.6, h - .15, 11.5, INK)
    ry += h + .08


# 18
s = new_slide(
    "実演前の確認事項",
    "デモが動く環境の前提、製品の状態、日本語の扱いを、実演の前に確認します。\n"
    "日本語の対応状況と、DAのランタイムは、お客様から聞かれやすい項目です。",
    "",
    """【目安】60秒。実演の前日までに一度確認します。
【環境の前提】Fabric Data Agentは、有料のFabric容量（F2以上）、またはPower BI Premium（P1以上）でFabricを有効にした容量が前提です（Microsoft Learn、Fabric data agent concepts、2026-10-21更新）。このデモのワークスペースの容量はF16です。管理者のテナント設定で、Copilotの利用、Azure OpenAI、Ontology（Preview）アイテムの作成を有効にします。反映まで最大1時間かかります。利用者には、DA、Ontology、Lakehouseの閲覧権限が必要です。
【製品の状態】Data Agentは一般提供（GA）です。実行基盤（ランタイム）には標準（GA）とプレビューの2種類があり、このデモのDAはプレビューで動いています（Learn、Fabric data agent runtime、2026-10-22更新）。公開済みのDAは、公開した時点のランタイムで動き続けます。標準ランタイムでの動作は、このデモでは確認していません。OntologyはPreviewで、Graph（GQL）の公開APIはbetaです。
【日本語】Learnは「The Fabric data agent doesn't currently support non-English languages. For optimal performance, provide questions, instructions, and example queries in English.」と記載しています（2026-10-21更新）。日本語は公式のサポート対象ではありません。このデモは、日本語で入力し、日本語で答えが返ることを、公開済みDAで確認しました。お客様には「日本語でも動作を確認しているが、公式のサポート対象ではない」と伝えます。
【確認済みと保証しない範囲】確認済みなのは、公開済みDAの日本語QA、実GraphのGQL照会と金額照合、Notebookの保存出力の再表示です。保証しないのは、日本語の公式サポート、標準ランタイムで同じ結果が出ること、M365 Copilot経由で同じグラフが表示されること、正式会計と実世界の因果の証明です。
【会計】全額と所有比率付き管理寄与を区別します。CFはEBITDAの80%を起点にイベントCFを足した簡易シミュレーションで、正式な営業CFやフリーCFではありません。税、非支配持分、資本連結、在庫未実現利益などは対象外です。
【チャネル】公開済みFabric DAを実演先にします。M365 Copilotへの公開はOFFで、結果を再構成する場合のグラフ表示は未保証です。
【PowerPoint表示】このファイルは編集可能な図形とネイティブグラフを使い、日本語にYu Gothicを指定しています。ローカル描画で点検していますが、Windows PowerPointなど実際に使う環境で、フォント、改行、ノート表示を最終確認します。
【リンク】\nDA：""" + DA + "\nNotebook：" + NOTEBOOK + "\nOntology：" + ONTOLOGY
    + "\nWorkspace：" + WORKSPACE,
    ["config/deployment_manifest.json:workspace.capacity_sku、data_agent.runtime", "config/step3_deployment.json",
     "Microsoft Learn：fabric/data-science/concept-data-agent（2026-10-21更新）",
     "Microsoft Learn：fabric/data-science/data-agent-runtime（2026-10-22更新）",
     "Microsoft Learn：fabric/data-science/data-agent-tenant-settings、fabric/iq/ontology/overview-tenant-settings"],
    "本番前に、使用環境のPowerPointで日本語フォント、改行、ノート、リンクを最終確認する")
check_cols = [
    ("必要な環境", TEAL, [
        ("容量", "有料のFabric容量（F2以上、またはP1以上でFabricを有効化）。このデモの容量はF16"),
        ("管理者の設定", "テナント設定で、Copilotの利用、Azure OpenAI、Ontology（Preview）の作成を有効にする。反映まで最大1時間"),
        ("利用者", "DA、Ontology、Lakehouseの閲覧権限"),
    ]),
    ("製品の状態（2026-10時点）", PURPLE, [
        ("Data Agent", "一般提供（GA）。ランタイムは標準（GA）とプレビューの2種類で、このデモはプレビュー"),
        ("OntologyとGraph", "OntologyはPreview。Graph（GQL）の公開APIはbeta"),
        ("日本語", "公式には英語以外は未サポート（Learn）。このデモは日本語で動作を確認済み"),
    ]),
    ("説明できる範囲", CORAL, [
        ("確認済み", "公開済みDAの日本語QA、実GraphのGQL照会と金額照合、Notebookの保存出力の再表示"),
        ("保証しない", "日本語の公式サポート、標準ランタイムで同じ結果になること、M365 Copilot経由で同じグラフが出ること、正式会計と実世界の因果の証明"),
    ]),
]
for c, (head, accent, items) in enumerate(check_cols):
    x = .62 + c * 4.09
    box(s, x, 2.58, 3.89, 3.12, WHITE, rounded=True)
    text(s, head, x + .22, 2.7, 3.45, .32, 16, accent, True)
    stack_items(s, x + .22, 3.1, 3.45, items, accent, size=11.5)
for k, (label, link) in enumerate([("Fabric DA", DA), ("Step03 Notebook", NOTEBOOK),
                                   ("Ontology", ONTOLOGY), ("Workspace", WORKSPACE)]):
    pill(s, label, .62 + k * 3.03, 5.9, 2.85, link)



s = new_slide(
    "社内にないイベントを集める方法",
    "イベントを一覧にしたデータは、多くの会社でまだ整っていません。\n"
    "既にある文書と業務データから拾い、小さな台帳として育てます。列の例は次の2ページです。",
    "",
    """【目安】準備用。お客様から「そんなデータは社内にない」と言われたときに使います。
【話す】「整ったイベントデータが最初から社内にないのは普通です。デモの10件は構造を見せるための設計データで、実際は既にある文書や業務データから拾います。」
【拾い先】①差分の理由コメント：月次の予算や見通しとの差異分析、経営会議や営業会議の資料に「何が起きたか」が文章で書かれていることが多く、最も手軽です。②事故や契約の記録：事故やクレームの報告、納期遅延、フォースマジュールや契約変更の通知。③業務データの異常：納期遅延、価格の乖離、歩留り、与信格付の変更などをルールで検知し、出来事の候補にします。④外部情報：ニュース、港湾や気象、地政学リスクの情報。
【進め方】差分から始める：デモのStep01のように基準値と実績の差を出し、差が大きい月と会社だけ「何が起きたか」を聞いて登録します。全部は網羅しません。実務ではbaselineを予算や見通しに置き換えます。下書きはAI、承認は人：文書から候補をAIで抽出し、担当者が確認して登録します。元の文書の権限や秘密度ラベルを引き継ぐ設計が前提です。最初は20〜30件：過去1〜2年の主要な出来事を、既存の資料から棚卸しします。
【単独と横断】別の種類として管理せず、1つの台帳に、影響を受ける会社、商品、地域、契約を紐づけます。複数の業種に及ぶものを横断と判定します。このデモはprimary_sector='all'で簡略化しています。
【金額の注意】実務でも、台帳の記録だけでは、イベントが原因だと証明できません。「根拠との紐づけ」と「影響額の推定」を分けて持ち、推定には方法（予算差異の配賦、前後比較など）と確度を添えます。デモの金額は設計値です。
【実装の例（案）】台帳はFabricのLakehouseかSharePointリスト。登録はTeamsやForms。異常検知はFabricのEventhouseとActivator。文書からの候補抽出はCopilotやエージェント。いずれも提案段階の構成で、このデモでは実装していません。
【避ける】「このデモのイベントが実際に社内にある」「AIが原因を特定する」と言いません。""",
    ["提案段階の構成（このデモでは未実装）", "data/demo_dataset.json:semantic_risk_events"],
    "影響額は推定とし、方法と確度を添える。実務ではbaselineを予算や見通しに置き換える")
text(s, "拾い先", .65, 2.62, 5.9, .34, 17, TEAL, True)
for i, (a, b) in enumerate([
    ("差分の理由コメント", "予実差異の分析、経営会議や営業会議の資料"),
    ("事故と契約の記録", "事故やクレームの報告、納期遅延、\nフォースマジュール、契約変更の通知"),
    ("業務データの異常", "納期遅延、価格の乖離、歩留り、\n与信格付の変更をルールで検知"),
    ("外部情報", "ニュース、港湾や気象、地政学リスクの情報"),
]):
    y = 3.02 + i * .68
    box(s, .65, y, 5.9, .6, WHITE, rounded=True)
    text(s, a, .82, y + .16, 1.95, .3, 14, TEAL, True)
    text(s, b, 2.78, y + (.17 if len(b) <= 21 else .08), 3.65, .45, 12, INK)
text(s, "進め方", 6.85, 2.62, 5.8, .34, 17, TEAL, True)
for i, (n_, a, b) in enumerate([
    ("1", "差分から始める", "基準値と実績の差が大きい月と会社だけ、理由を聞いて登録する。"),
    ("2", "下書きはAI、承認は人", "文書から候補をAIで抽出し、担当者が確認して登録する。"),
    ("3", "最初は20〜30件", "過去1〜2年の主要な出来事を、既存の資料から棚卸しする。"),
]):
    y = 3.02 + i * .9
    box(s, 6.85, y, 5.84, .8, WHITE, rounded=True)
    text(s, n_, 7.03, y + .15, .4, .45, 22, TEAL, True)
    text(s, a, 7.55, y + .07, 5.0, .3, 14, INK, True)
    text(s, b, 7.55, y + .37, 5.0, .42, 12, INK)
box(s, .65, 5.9, 12.04, .5, WHITE, rounded=True)
text(s, "単独と横断：影響を受ける会社、商品、地域、契約を台帳に紐づけ、複数の業種に及ぶものを横断と判定する。",
     .85, 6.02, 11.65, .3, 13, NAVY)

# Event ledger columns (two pages). The "demo" column marks columns that exist in the demo tables.
s = new_slide(
    "イベントの列の例",
    "台帳は、イベント、影響額、根拠の3つに分けて持ちます。このページはイベントの台帳です。\n"
    "列名は、このデモの表（risk_events）に合わせ、デモにない列は「デモの表：なし」と示します。",
    "",
    """【目安】60秒。「台帳には、何を書けばよいのか」と聞かれたときに使います。
【話す】「イベントの台帳は、1つの出来事につき1行です。最低限必要なのは、ID、何が起きたか、期間、好影響か悪影響か、影響が及ぶ範囲の5つです。重大度や確度、担当者は、あとから足せます。」
【指す】表の左から、列、入れる内容、このデモでの例（食品損失）、必須か推奨か、このデモに同じ列があるか、の順です。
【例の出どころ】例の値は、このデモのリスクイベントの表（risk_events）の、FOD_NEGの行です。event_idはFOD_NEG、期間は2026-02-09から2026-03-01、severityは5、confidenceは0.95、causal_statusはdesigned_attribution_not_observational_proof、narrativeはSYNTHETIC: cold chain spoilage incidentです。
【デモにない列】event_name（何が起きたかを一言で）、status（調査中、確定、終了）、owner（担当者）は、実務で足す列の例です。このデモの表にはありません。
【因果の扱い】causal_statusは、因果を確認済みか、仮説か、設計上の配賦かを区別する列です。このデモの金額は設計上の配賦で、観測による因果の証明ではありません。実務でも、台帳の記録だけでは、イベントが原因だと証明できないので、この列で扱いを明示します。
【横断】影響が複数の事業に及ぶ場合は、別の種類として管理せず、事業の欄（primary_sector）に、横断を意味する値を入れます。このデモは、primary_sector='all'で表しています。
【つなぎ】「次のページは、影響額と根拠の台帳です。」
【列名と「デモの表」の列】列名は、このデモの表（risk_events）の列名に合わせています。「デモの表：あり」は、同じ名前の列がデモの表にあることを示し、「なし」は、実務で足す列の例です。""",
    ["ddl/10_common_subsidiary_tables.sql:risk_events", "data/demo_dataset.json:semantic_risk_events（FOD_NEG）"],
    "最初は必須の5列だけでよい。因果の扱いと確度は、早めに列として持つ")
ev_rows = [
    ["event_id", "イベントを区別するID", "FOD_NEG", "必須", "あり"],
    ["event_name", "何が起きたかを一言で", "食品の廃棄事故（コールドチェーン）", "必須", "なし"],
    ["start_date、end_date", "起きた期間。終わっていなければ空欄", "2026-02-09 から 2026-03-01", "必須", "あり"],
    ["direction", "好影響か悪影響か", "悪影響（negative）", "必須", "あり"],
    ["primary_sector、region_code", "影響が及ぶ事業と地域（複数なら横断）", "食品（日本と米国）", "必須", "あり"],
    ["severity", "重大度（1から5など）", "5", "推奨", "あり"],
    ["confidence", "影響の見立ての確からしさ", "0.95", "推奨", "あり"],
    ["causal_status", "因果の扱い（確認済み、仮説、配賦）", "設計上の配賦。観測による証明ではない", "推奨", "あり"],
    ["status、owner", "調査中、確定、終了と、担当者", "（このデモにはない）", "推奨", "なし"],
    ["narrative", "経緯の説明を1行から3行で", "SYNTHETIC: cold chain spoilage incident", "推奨", "あり"],
]
grid_table(s, .62, 2.5, [2.75, 3.3, 3.85, .8, 1.38],
           ["列", "入れる内容", "例（食品損失）", "区分", "デモの表"],
           ev_rows, [.35] * 10, size=11, head_size=11.5,
           colors=[TEAL, INK, INK, CORAL, MUTED], bolds=[True, False, False, True, False])

s = new_slide(
    "影響額と根拠の列の例",
    "影響額の台帳は、1つのイベントが、どの契約に、いつ、いくら効いたかを、1行ずつ持ちます。\n"
    "根拠の台帳は、その金額を裏付ける元の記録の所在を持ちます。",
    "",
    """【目安】75秒。前のページのイベントの台帳に続けて使います。
【話す】「影響額の台帳は、イベントと契約の組ごとに1行です。いつの影響か、いくらか、何と比べた差か、どう推定したかを持ちます。根拠の台帳は、元の文書やログの所在を持ち、金額から元の記録までたどれるようにします。」
【影響額の例】例の値は、このデモの食品損失の代表効果（effect_id SUB_FOD_JP_01_A133_0_FOD_NEG_F）です。契約SUB_FOD_JP_01_C04、効果日2026-02-11、影響額-3,003,488.94円（費用の増加3,003,488.94円で、EBITDAへの影響はマイナス）、attribution_methodはunique_event_activity_additive（イベントと活動の組ごとに1回だけ加算）です。経路は、cost、遅れは2日、減衰の半減期は7日です。
【根拠の例】evidence_id SUB_FOD_JP_01_FOD_NEG_0、元の記録の所在はSIM://SUB_FOD_JP_01/FOD_NEG/operations_log、記録日は2026-02-09、種類はoperations_logです。いずれも合成データです。
【デモにない列】baseline_ref（何と比べた差か。予算や見通し）と、reliability（信頼度）は、実務で足す列の例です。このデモでは、基準値（baseline）のシナリオとの差として金額を持っています。
【推定の注意】実務では、影響額は推定です。方法（予算差異の配賦、前後比較など）と確度を、必ず添えます。イベントが原因だと証明するものではありません。
【列名と「デモの表」の列】列名は、このデモの表（event_financial_effects、event_impact_bridge、event_evidence）の列名に合わせています。「あり」は同じ名前の列がデモの表にあること、「なし」は実務で足す列の例です。金額の列は、デモではEBITDAへの影響額（delta_ebitda_local）と通貨（currency_code）の2列です。
【つなぎ】「最初から全部は不要です。必須の列から、小さく始めます。」""",
    ["ddl/10_common_subsidiary_tables.sql:event_financial_effects,event_impact_bridge,event_evidence",
     "data/demo_dataset.json:semantic_event_financial_effects、semantic_event_evidence、semantic_event_impact_bridge（FOD_NEG）"],
    "影響額は推定として持つ。方法と確度を添え、因果の証明とは書かない")
led_rows = [
    ["影響額", "event_id、contract_id", "どのイベントが、どの契約に効いたか", "FOD_NEG、SUB_FOD_JP_01_C04", "必須", "あり"],
    ["影響額", "effect_date", "いつの影響か", "2026-02-11", "必須", "あり"],
    ["影響額", "delta_ebitda_local、currency_code", "EBITDAへの影響額（現地通貨）と通貨", "−3,003,488.94　JPY", "必須", "あり"],
    ["影響額", "baseline_ref", "何と比べた差か（予算、見通し）", "（baselineシナリオとの差）", "必須", "なし"],
    ["影響額", "attribution_method", "推定や配分の方法", "unique_event_activity_additive", "必須", "あり"],
    ["影響額", "impact_channel、lag_days", "影響の経路と、効くまでの遅れ", "cost、2日", "推奨", "あり"],
    ["根拠", "evidence_id、event_id", "根拠のIDと、結び付くイベント", "SUB_FOD_JP_01_FOD_NEG_0、FOD_NEG", "必須", "あり"],
    ["根拠", "source_record", "元の記録の所在（ID、URLなど）", "SIM://SUB_FOD_JP_01/FOD_NEG/operations_log", "必須", "あり"],
    ["根拠", "recorded_date、evidence_type", "記録された日と、記録の種類", "2026-02-09、operations_log", "推奨", "あり"],
    ["根拠", "reliability", "信頼度", "（このデモにはない）", "推奨", "なし"],
]
grid_table(s, .62, 2.5, [.8, 2.9, 2.95, 3.7, .72, 1.01],
           ["台帳", "列", "入れる内容", "例（食品損失）", "区分", "デモの表"],
           led_rows, [.33] * 10, size=10.5, head_size=11,
           colors=[PURPLE, TEAL, INK, INK, CORAL, MUTED], bolds=[True, True, False, False, True, False])

# Structure explanation: Lakehouses and the ontology
JP_ENTITY = {"Evidence": "根拠", "RiskEvent": "イベント", "FinancialEffect": "効果",
             "FinancialMetric": "日次指標", "GroupKPI": "全社KPI", "Activity": "活動",
             "Contract": "契約", "Shipment": "出荷", "Company": "会社"}
lh_tables, lh_rows = defaultdict(int), defaultdict(int)
for key, rows_n in cloud["counts"].items():
    lh_tables[key.split(".")[0]] += 1
    lh_rows[key.split(".")[0]] += rows_n
assert (cloud["tables"], cloud["rows"]) == (95, 90424) and lh_tables["LH_Group_SAP"] == 17

s = new_slide(
    "9つのLakehouseの役割分担",
    "会社ごとの原本は、子会社のLakehouseに分けたまま残します。\n"
    "全社の分析と問い合わせは、親のLakehouseに集めた表で行います。",
    "構成の解説 1  |  Lakehouse",
    """【目安】60秒。構成の質問が出たときに使います。
【話す】「会社ごとの原本は、それぞれのLakehouseに分けて残しています。親のLakehouseには、全社で同じ尺度で分析するための表と、各社の業種固有データへの参照を置いています。」
【指す】右の8つのタイル（4業種×日本、米国）、次に左の親Lakehouseの4つのまとまりの順です。
【数字】物理表は95、行数は90,424です。内訳は親が17表（58,018行）、子会社が78表です。出荷表を持つ6社は10表、金融の2社は出荷表がなく9表です。ショートカット8本は物理表に含めません。物理表95とショートカット8本を足して103表とは言いません。
【親の4つのまとまり】①SAP風の管理データ5表（会社マスタ、為替、仕訳、内部取引、連結消去）。②分析用の投影11表（semantic_*）。8社分を1表に統合したコピーで、source_lakehouse列に原本の場所を残します。③全社KPI（kpi_snapshot）。④業種拡張のショートカット8本（raw_*_extension）。子会社の拡張表をコピーせずに参照します。
【各子会社】共通9表（会社プロフィール、契約、活動、出荷、イベント、根拠、影響の橋渡し、財務効果、日次指標）と、業種拡張1表です。業種拡張は、金属が品質（純度、製錬所の稼働率）、エネルギーが配分（貯蔵率、配分量）、食品がコールドチェーン（温度、廃棄率）、金融がエクスポージャ（想定元本、信用スプレッド、ヘッジ比率）です。金融は、存在しない物理出荷を補完しません。
【実装上の注意】親と金属日本のLakehouseはスキーマ有効（Tables/dbo）、残り7つは従来形式（Tables/）でした。デプロイは、実際のLakehouseの属性を見て分岐します。
【避ける】すべてのデータを親へコピーしたと言いません。コピーしているのは分析用の投影だけで、業種拡張は参照です。
【つなぎ】「このデータが、どの順でOntology、DA、Notebookに流れるかを次に見ます。」""",
    ["reports/cloud_validation.json", "ddl/", "config/sector_shortcuts.json", "docs/architecture.md"],
    "物理表は95、行数は90,424。ショートカット8本は物理表に数えない")
box(s, .62, 2.62, 5.25, 3.84, WHITE, rounded=True)
text(s, "親：LH_Group_SAP", .85, 2.76, 4.8, .36, 18, TEAL, True)
text(s, "17表とショートカット8本", .85, 3.14, 4.8, .26, 12.5, MUTED)
for i, (head, desc, fill, accent, dashed) in enumerate([
    ("SAP風の管理データ　5表", "会社マスタ、為替、仕訳、内部取引、連結消去", SRC_FILL, None, False),
    ("分析用の投影　11表", "8社分を1表に統合し、原本の場所を列に残す", MAT_FILL, PURPLE, False),
    ("全社KPI　1表", "EBITDAとCF、全額と持分管理", MAT_FILL, PURPLE, False),
    ("業種拡張の参照　8本", "子会社の業種拡張表を、コピーせず参照", ZERO_FILL, TEAL, True),
]):
    y = 3.5 + i * .65
    shp = box(s, .8, y, 4.9, .57, fill, line=accent, rounded=True, flat=True)
    if accent:
        shp.line.width = Pt(1.25)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    text(s, head, .95, y + .04, 4.6, .26, 14, NAVY, True)
    text(s, desc, .95, y + .31, 4.6, .24, 11, INK)
for lx, fill, accent, label, dashed in [
    (.85, SRC_FILL, None, "そのまま", False),
    (2.2, MAT_FILL, PURPLE, "実体化（コピー）", False),
    (4.05, ZERO_FILL, TEAL, "ゼロコピー（参照）", True),
]:
    shp = box(s, lx, 6.17, .22, .18, fill, line=accent, rounded=True, flat=True)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    text(s, label, lx + .3, 6.14, 1.6, .24, 10.5, INK)
text(s, "子会社8社のLakehouse（原本）", 6.15, 2.66, 6.4, .34, 16, TEAL, True)
for c, (jp_sector, en_sector) in enumerate([("金属", "Metals"), ("エネルギー", "Energy"),
                                            ("食品", "Foods"), ("金融", "Finance")]):
    x = 6.15 + c * 1.67
    text(s, jp_sector, x, 3.08, 1.5, .28, 13, INK, True, PP_ALIGN.CENTER)
    for r, country in enumerate(["JP", "US"]):
        name = f"LH_{en_sector}_{country}_01"
        y = 3.4 + r * .92
        box(s, x, y, 1.5, .8, WHITE, rounded=True)
        text(s, name, x + .08, y + .12, 1.34, .24, 10, TEAL, True)
        text(s, f"{lh_tables[name]}表／{lh_rows[name]:,}行", x + .08, y + .44, 1.34, .26, 11.5, INK)
box(s, 6.15, 5.3, 6.52, 1.08, WHITE, rounded=True)
text(s, "各子会社は、共通9表と業種拡張1表を持つ（金融の2社は出荷表がなく9表）。\n"
        "業種拡張は、金属の品質、エネルギーの配分、食品の温度、金融のエクスポージャ。",
     6.32, 5.5, 6.25, .8, 11, INK)

s = new_slide(
    "Ontologyの全体マップ",
    "色は、お客様が既に持っていることの多いデータ（青緑）と、これから整備するデータ（赤）を表します。\n"
    "実体は17種類、関係は39種類で、整備が必要なのは、イベント、効果、根拠の3つです。",
    "",
    """【目安】90秒。Ontologyの説明で最も使うページです。
【話す】「Ontologyは、このデモの『もの』と『つながり』に名前を付けた設計図です。イベントは根拠に裏付けられ、効果を持ち、活動と契約に影響します。効果は日次指標に反映され、全社KPIに集約されます。色は、お客様のデータに当てはめたときの整備状況です。青緑は、基幹システムや会計、契約管理に既にあることが多いデータ。赤は、多くの会社でまだ整っていないデータで、これから整備します。」
【指す】上段を左から右へ（根拠、イベント、効果、日次指標、全社KPI）。次に中段（業種拡張、活動、契約）、下段（出荷、会社）の順です。各ノードの3行目は、お客様の元データの例です。
【整備が必要な3つ】イベント、効果、根拠です。多くの会社には、イベントを一覧にした台帳も、イベントごとの影響額の推定も、根拠の記録を元データにつなぐ表もありません。第5章の「社内にないイベントを集める方法」と「イベントの列の例」「影響額と根拠の列の例」で、台帳の作り方を説明します。
【既存のデータ】会社（会社マスタ）、契約（契約管理）、活動（受注や売上の明細）、出荷（物流の実績）、日次指標（管理会計や月次決算）、全社KPI（経営ダッシュボード）、業種拡張（現場のシステムの業種固有データ）です。元データの例は、一般的な想定で、お客様ごとに確認します。
【数】実体は17種類（本体9、業種拡張8）、関係は39種類（本体15、業種拡張24）です。
【色】矢印の灰色は関連、紫は影響の配賦、緑は対応と集約です。Step03の図（食品損失のDAG）と同じ色分けです。矢印の向きは関係の定義の向きで、因果の方向を証明するものではありません。
【図に矢印を描いていない関係】EffectExplainedByEvidence（効果→根拠）、会社に結ぶ4本（契約、活動、効果、日次指標→会社）、業種拡張のForCompanyとForContract。一覧は「39種類の関係と読み方」にあります。
【避ける】「Ontologyが因果を判定する」「OntologyがJOINを置き換える」と言いません。表のJOINでも同じ数値は作れます。Ontologyの価値は、意味と関係を名前付きで共有し、根拠までの経路を再利用できることです。
【つなぎ】「各実体は、どの表のどの列につながっているかを次に見ます。」""",
    ["ontology_tmdl/entities/", "ontology_tmdl/entityRelationships.tmdl", "reports/graph_edge_counts.json"],
    "整備が必要な3つ（イベント、効果、根拠）は、第5章の台帳で作る。矢印は主な関係だけを描いている")

HAVE_FILL, NEED_FILL = "E1F1EE", "FBE8EB"


def onode(slide, x, y, jp, en, color, hint, need=False, w=1.62, h=.82):
    shp = box(slide, x, y, w, h, NEED_FILL if need else HAVE_FILL, line=CORAL if need else TEAL, rounded=True)
    shp.line.width = Pt(1.25)
    text(slide, jp, x + .1, y + .05, w - .2, .3, 14, color, True)
    text(slide, en, x + .1, y + .35, w - .2, .2, 8.5, MUTED)
    text(slide, hint, x + .1, y + .55, w - .2, .22, 9.5, INK, True)


def olabel(slide, label, x, y, w=.95, align=PP_ALIGN.CENTER):
    text(slide, label, x, y, w, .24, 10, MUTED, False, align)


def olabel_on(slide, label, cx, cy, w):
    box(slide, cx - w / 2, cy - .13, w, .26, BG, flat=True)
    text(slide, label, cx - w / 2, cy - .1, w, .24, 10, MUTED, False, PP_ALIGN.CENTER)


A_Y, B_Y, C_Y = 2.7, 4.2, 5.6
XS = [.65, 3.22, 5.79, 8.36, 10.93]
onode(s, XS[0], A_Y, "根拠", "Evidence", MUTED, "文書に点在", need=True)
onode(s, XS[1], A_Y, "イベント", "RiskEvent", CORAL, "台帳を新設", need=True)
onode(s, XS[2], A_Y, "効果", "FinancialEffect", PURPLE, "影響額を推定", need=True)
onode(s, XS[3], A_Y, "日次指標", "FinancialMetric", TEAL, "管理会計、月次決算")
onode(s, XS[4], A_Y, "全社KPI", "GroupKPI", TEAL, "経営ダッシュボード")
onode(s, XS[0], B_Y, "業種拡張 ×8", "4業種 × 日米", NAVY, "現場システム")
onode(s, XS[1], B_Y, "活動", "Activity", NAVY, "受注、売上明細")
onode(s, XS[2], B_Y, "契約", "Contract", NAVY, "契約管理")
onode(s, XS[1], C_Y, "出荷", "Shipment", NAVY, "物流、出荷実績")
for i, (label, color) in enumerate([("裏付ける", MUTED), ("効果を持つ", PURPLE),
                                    ("指標に反映", TEAL), ("KPIに集約", TEAL)]):
    arrow(s, XS[i] + 1.62 + .03, A_Y + .41, XS[i + 1] - .06, A_Y + .41, color)
    olabel(s, label, XS[i] + 1.62, A_Y + .08)
arrow(s, XS[0] + 1.62 + .03, B_Y + .41, XS[1] - .06, B_Y + .41, MUTED)
olabel(s, "活動に紐づく", XS[0] + 1.62, B_Y + .5)
arrow(s, XS[1] + 1.62 + .03, B_Y + .41, XS[2] - .06, B_Y + .41, MUTED)
olabel(s, "契約に属する", XS[1] + 1.62, B_Y + .5)
arrow(s, 4.03, A_Y + .85, 4.03, B_Y - .06, MUTED)
olabel_on(s, "活動に影響", 4.03, 3.9, .85)
arrow(s, 6.6, A_Y + .85, 6.6, B_Y - .06, MUTED)
olabel_on(s, "契約への効果", 6.6, 3.9, 1.0)
arrow(s, 6.0, A_Y + .85, 4.87, B_Y + .12, MUTED)
olabel_on(s, "活動への効果", 5.43, 3.97, 1.0)
arrow(s, 4.03, C_Y - .04, 4.03, B_Y + .88, MUTED)
olabel_on(s, "活動の出荷", 4.03, 5.35, .9)
arrow(s, 4.87, C_Y + .13, 5.9, B_Y + .86, MUTED)
olabel_on(s, "契約に属する", 5.38, 5.42, 1.0)
box(s, 6.3, C_Y, 6.25, .82, HAVE_FILL, line=TEAL, rounded=True).line.width = Pt(1.25)
text(s, "会社", 6.5, C_Y + .08, 1.0, .3, 14, NAVY, True)
text(s, "Company", 6.5, C_Y + .4, 1.2, .2, 8.5, MUTED)
text(s, "契約、活動、効果、日次指標の持ち主（4種類の関係）", 7.75, C_Y + .1, 4.7, .26, 11, INK)
text(s, "元データの例：会社マスタ", 7.75, C_Y + .45, 4.7, .24, 9.5, INK, True)
box(s, 8.36, 4.2, 4.19, 1.2, WHITE, rounded=True)
text(s, "色と矢印", 8.56, 4.27, 3.8, .26, 12, TEAL, True)
box(s, 8.56, 4.6, .26, .2, HAVE_FILL, line=TEAL, rounded=True).line.width = Pt(1.25)
text(s, "既存のデータで用意しやすい", 8.92, 4.57, 3.5, .24, 11, INK, True)
box(s, 8.56, 4.88, .26, .2, NEED_FILL, line=CORAL, rounded=True).line.width = Pt(1.25)
text(s, "これから整備が必要（イベントの台帳）", 8.92, 4.85, 3.5, .24, 11, INK, True)
text(s, "矢印：灰＝関連、紫＝配賦、緑＝集約（因果の証明ではない）", 8.56, 5.12, 3.85, .22, 9, MUTED)


s = new_slide(
    "実体と元の表の対応",
    "実体ごとに、元になる表とキーの列を1つずつ指定します。\n"
    "その表の1行が、グラフの1つの点（ノード）になります。",
    "構成の解説 5  |  Ontologyのバインディング",
    """【目安】60秒。
【話す】「実体ごとに、元になる表とキーの列を1つずつ指定しています。その表の1行が、グラフの1つの点になります。点の合計は33,960です。」
【表の選び方】会社は会社マスタ（company_master）、全社KPIはkpi_snapshot、それ以外の本体の実体は、8社分を統合した分析用の投影（semantic_*）に結び付けています。業種拡張の8実体は、親にあるショートカット（raw_*_extension）に結び付けています。
【接続】各表は、親Lakehouseに対するDirectLakeのパーティションとして定義しています（定義はontology_tmdl/tablesにあります）。
【数の読み方】日次指標の17,520は、8社×365日×6系列（EBITDAとCFのactualとbaseline、売上と原価のactual）です。全社KPIの2,920は、8種類（EBITDAとCF、全額と持分管理、actualとbaseline）×365日です。点の数を、取引件数や事故件数として読みません。
【リンク表】semantic_event_impact_bridge、semantic_effect_metric_links、semantic_metric_kpi_linksの3表は、点にならず、関係を表すために使います。
【避ける】Ontologyが元の表を置き換える、と言いません。元の表がデータの正で、金額の厳密な集計はSQLのDECIMALで行います（グラフの数値は浮動小数点です）。
【つなぎ】「次に、実体どうしの関係を見ます。」""",
    ["ontology_tmdl/entities/", "ontology_tmdl/tables/", "reports/graph_node_counts.json"],
    "点は合計33,960。リンク表3本は点にならず、関係を表す")
cols = [("実体", .83, 2.0), ("意味", 2.9, 1.3), ("元の表（親Lakehouse）", 4.25, 4.3),
        ("キー", 8.6, 1.9), ("点の数", 10.5, 1.75)]
for label, x, w in cols:
    text(s, label, x, 2.66, w, .3, 13, TEAL, True, PP_ALIGN.RIGHT if label == "点の数" else None)
entity_rows = [(n, JP_ENTITY[n]) for n in ["Company", "RiskEvent", "Evidence", "Contract", "Activity",
                                            "Shipment", "FinancialEffect", "FinancialMetric", "GroupKPI"]]
table_rows = []
for name, jp in entity_rows:
    e = ontology_entities[name]
    table_rows.append((name, jp, e["table"], e["key"], f"{node_counts[name]:,}"))
ext_entities = [n for n in ontology_entities if n.endswith("Extension")]
assert len(ext_entities) == 8 and {node_counts[n] for n in ext_entities} == {365}
table_rows.append(("業種拡張（8実体）", "業種固有の観測", "raw_<業種>_<jp/us>_extension", "extension_id", "365 × 8"))
for i, vals in enumerate(table_rows):
    y = 3.02 + i * .335
    if i % 2 == 0:
        box(s, .67, y - .04, 12.0, .31, WHITE, flat=True)
    for j, ((label, x, w), val) in enumerate(zip(cols, vals)):
        text(s, val, x, y, w, .26, 12.5, NAVY if j == 0 else INK, j == 0,
             PP_ALIGN.RIGHT if j == 4 else None)

used_relations = {r["relation"] for r in audit["diagram_relations"]}
assert len(used_relations) == 7
rel_by_name = {r["name"]: r for r in ontology_relationships}


def direction(r):
    return f"{JP_ENTITY[r['from']]} → {JP_ENTITY[r['to']]}"


s = new_slide(
    "39種類の関係と読み方",
    "関係には向きがあり、「実体 → 実体」の形で名前を付けます。\n"
    "表の外部キー列で表すものが36種類、リンク表（※）で表すものが3種類です。",
    "構成の解説 6  |  Ontologyの関係",
    """【目安】60秒。
【話す】「関係には向きがあり、『実体から実体へ』の形で名前を付けています。EventHasEffectなら、イベントから効果へ、という読み方です。」
【数】39種類で、点と点を結ぶ線は合計65,684本です。36種類は表の外部キーの列で、3種類はリンク表で表します。
【リンク表の3本】EventImpactsActivity（イベント→活動）、EffectRealizedInMetric（効果→日次指標）、MetricRollsUpToGroupKPI（日次指標→全社KPI）です。IDの対応を明示した表で結び、日付だけの結合が起こす曖昧さを避けます。
【グループ】原因をたどる4本、取引をたどる5本、金額を集約する2本、会社に結ぶ4本（本体15本）に、業種拡張の24本（8実体×3本）を加えて39種類です。
【●】Step03の図で使っている7本です（本体6本と、食品日本の業種拡張1本）。残りの関係は、図に出ていないだけで、グラフには入っています。
【避ける】本数を取引の件数と読みません。たとえばActivityUnderContractの8,760は、活動の数と同じです（1つの活動が1つの契約に属します）。
【つなぎ】「最後に、なぜこの構成にしたのかを3つの判断で説明します。」""",
    ["ontology_tmdl/entityRelationships.tmdl", "reports/graph_edge_counts.json",
     "reports/step3_live_audit.json:diagram_relations"],
    "●はStep03の図で使う7本。※のリンク表は、IDの対応で結び、日付だけの結合を避ける")


def rel_group(card_x, card_w, y0, title, names):
    text(s, title, card_x + .2, y0, card_w - .4, .3, 13, TEAL, True)
    for k, nm in enumerate(names):
        r = rel_by_name[nm]
        y = y0 + .34 + k * .275
        if nm in used_relations:
            text(s, "●", card_x + .2, y, .25, .26, 11, TEAL, True)
        text(s, nm + (" ※" if r["link_table"] else ""), card_x + .45, y, 2.65, .26, 11, INK, True)
        text(s, direction(r), card_x + 3.1, y, 2.25, .26, 11, MUTED)
        text(s, f"{edge_counts[nm]:,}", card_x + card_w - .85, y, .65, .26, 11, INK, False, PP_ALIGN.RIGHT)
    return y0 + .34 + len(names) * .275 + .12


groups_left = [
    ("原因をたどる", ["EvidenceSupportsEvent", "EventHasEffect", "EffectExplainedByEvidence", "EventImpactsActivity"]),
    ("取引をたどる", ["EffectOnActivity", "EffectOnContract", "ActivityUnderContract",
                  "ShipmentForActivity", "ShipmentUnderContract"]),
]
groups_right = [
    ("金額を集約する", ["EffectRealizedInMetric", "MetricRollsUpToGroupKPI"]),
    ("会社に結ぶ", ["ContractOwnedByCompany", "ActivityOwnedByCompany", "EffectOwnedByCompany", "MetricForCompany"]),
]
assert sum(len(g[1]) for g in groups_left + groups_right) == 15
box(s, .62, 2.62, 6.0, 3.76, WHITE, rounded=True)
box(s, 6.85, 2.62, 5.85, 3.76, WHITE, rounded=True)
y = 2.74
for title, names in groups_left:
    y = rel_group(.62, 6.0, y, title, names)
y = 2.74
for title, names in groups_right:
    y = rel_group(6.85, 5.85, y, title, names)
text(s, "業種拡張（24種類）", 7.05, y, 5.4, .3, 13, TEAL, True)
text(s, "●", 7.05, y + .34, .25, .26, 11, TEAL, True)
text(s, "FoodsJPExtensionForActivity ほか", 7.3, y + .34, 3.6, .26, 11, INK, True)
text(s, "各365", 11.8, y + .34, .65, .26, 11, INK, False, PP_ALIGN.RIGHT)
text(s, "8実体が、活動、契約、会社に1本ずつ結ぶ", 7.3, y + .62, 5.2, .26, 11, MUTED)

s = new_slide(
    "この構成にした3つの理由",
    "Ontologyの実体は、1つの主なデータ元（表）の行から作ります。8社に分かれた原本のままでは、\n"
    "同じ種類の実体（契約など）を、会社ごとに別々に定義することになります。",
    "",
    """【目安】90秒。技術に詳しい方への補足です。
【話す】「Ontologyの実体は、主なデータ元の1つの表の行から作ります。ほかの表は、共通の列で結んで列を足す使い方です。8社の同じ形の表を、縦に積む用途ではありません。8社の原本をそのまま使うと、同じ種類の実体を、会社ごとに別々に定義することになります。そこで3つの判断をしました。」
【根拠】Microsoft Learn（Bind data、2026-10-21更新）は、実体の2つ目以降のデータ元を「副のデータ元」として追加し、主のデータ元との共通の列で関係付けて、列（プロパティ）を足すと説明しています。この構成は、2026年10月の実機確認で、実体ごとの静的なデータ元を1つにする必要があった結果です。Previewなので、将来変わる可能性があります。
【判断1：8社分を親の投影表に統合】親のLakehouseのsemantic_*の表に8社分を統合し、source_lakehouse列で原本の場所を残しました。原本は子会社のLakehouseに残り、投影は分析用のコピーです。困りごとは、同じ形の表が8社に別々にあり、そのまま結ぶと「契約」の定義が8個になることでした。得られるのは、実体、関係、質問を全社で1回ずつ書けば済むことです。代償は、コピーなので原本が変わったら作り直すことです。
【判断2：業種拡張は別の実体にして、ショートカットで参照】金属の純度、エネルギーの貯蔵率、食品の温度、金融のスプレッドなど、項目が業種ごとに違います。1つの表にまとめると、空欄だらけになります。そこで、日米×4業種の8実体に分け、親のショートカットで参照します。データはコピーしません。8つの子会社Lakehouseを別々のソースとして直接結ぶ方法は、最初のテストでグラフを照会できず、ショートカット方式に切り替えました。
【判断3：多対多の関係はリンク表で明示】1つのイベントが多くの活動に影響し、1つの活動が複数のイベントの影響を受けます。イベントと活動、効果と日次指標、日次指標と全社KPIの3本は、IDの対応を持つ表で結びます。日付だけで結ぶと、別の行どうしが結び付いて、金額が二重に数えられます。対応表の整備が要り、イベント台帳と同じ課題になります。
【運用上の注意】公開APIで定義を取り込むと、表との結び付き（バインディング）に使う接続情報が外れます。定義を差し替えたら、バインディングを再適用し、グラフを更新してから照会します。取り込みが成功しただけでは、グラフで照会できる保証になりません。この内容は、2026年10月のPreviewでの確認結果です。
【避ける】これらを製品の恒久的な仕様として説明しません。実演の前に、現在の製品の挙動を確認します。
【つなぎ】「最後に、実装の技術質問が出たときの説明を用意しています。」""",
    ["docs/architecture.md", "config/ontology_connection_bindings.json", "config/sector_shortcuts.json",
     "Microsoft Learn：fabric/iq/ontology/how-to-bind-data（2026-10-21更新）"],
    "定義を差し替えたら、バインディングの再適用とグラフの更新が要る。取り込みが成功しただけでは照会できない")
reason_rows = [
    ["1　8社分を\n1つの表にまとめた",
     "同じ形の表が8社のLakehouseに別々にあり、そのまま結ぶと「契約」の定義が8個になる。",
     "親のLakehouseに、8社分を1つにまとめた分析用の表を作り、「契約」はその表に結ぶ。会社の区別は列（source_lakehouse）に残す。",
     "実体、関係、質問を、全社で1回ずつ書けば済む。コピーなので、原本が変わったら作り直す。"],
    ["2　業種固有の項目は\n別の実体に分けた",
     "項目が業種で違う（金属の純度、食品の温度、金融のスプレッドなど）。1表にまとめると空欄だらけになる。",
     "日米×4業種の8つの拡張実体に分け、親のショートカット（ゼロコピー）で各社の拡張表を参照する。",
     "共通の質問と業種固有の質問を、同じOntologyで扱える。子会社のLakehouseを直接結ぶ方式は、最初のテストで照会できなかった。"],
    ["3　多対多の関係は\n対応表で結んだ",
     "1つのイベントが多くの活動に影響する。日付だけで結ぶと、別の行どうしが結び付いて金額が二重に数えられる。",
     "IDの対応を持つ表（semantic_event_impact_bridgeなど3表）を作り、IDが一致する行だけを結ぶ。",
     "84効果が、84件として数えられる。対応表の整備が要り、イベント台帳と同じ課題になる。"],
]
grid_table(s, .62, 2.6, [2.15, 3.05, 3.45, 3.43], ["理由", "困りごと", "やったこと", "得られること、代償"],
           reason_rows, [1.12, 1.15, 1.15], size=12.5,
           colors=[TEAL, INK, INK, INK], bolds=[True, False, False, False])


# New-experience features and the minimal-start guide. Counts come from the repo's ontology
# definition files; the live Fabric state was also checked in the browser on 2026-10-26.
runtime_model = json.loads((PROJECT / "reports/ontology_runtime_current.json").read_text())["model"]
runtime_status = {
    "namespaces": len(runtime_model["namespaces"]),
    "rules": len(runtime_model["rules"]),
    "metrics": len(runtime_model["metrics"]),
    "shared_properties": len(runtime_model["reusableProperties"]),
    "entity_descriptions": sum(1 for e in runtime_model["entities"] if e.get("description")),
    "property_descriptions": sum(1 for e in runtime_model["entities"]
                                 for p in e["properties"] if p.get("description")),
    "synonyms": sum(1 for e in runtime_model["entities"] if e.get("synonyms")),
    "resource_links": sum(1 for e in runtime_model["entities"] if e.get("resourceLinks")),
    "properties": sum(len(e["properties"]) for e in runtime_model["entities"]),
}
assert runtime_status == {"namespaces": 1, "rules": 0, "metrics": 0, "shared_properties": 0,
                          "entity_descriptions": 1, "property_descriptions": 0, "synonyms": 0,
                          "resource_links": 0, "properties": 172}

property_names = {}
for f in sorted((ONT / "entities").glob("*.tmdl")):
    t = f.read_text()
    property_names[re.search(r"^entity (\w+)", t, re.M).group(1)] = re.findall(r"^\tproperty (\w+)", t, re.M)
property_users = defaultdict(list)
for entity_name, props in property_names.items():
    for p in props:
        property_users[p].append(entity_name)
SHARED_CANDIDATES = ["entity_id", "contract_id", "activity_id", "observation_date",
                     "source_lakehouse", "currency_code", "event_id"]
assert [len(property_users[p]) for p in SHARED_CANDIDATES] == [15, 12, 11, 8, 7, 6, 3]
assert sum(len(v) for v in property_names.values()) == runtime_status["properties"]
EXT_ORDER = [f"{sec}{cc}Extension" for sec in ("Metals", "Energy", "Foods", "Finance") for cc in ("JP", "US")]
assert sorted(EXT_ORDER) == sorted(n for n in property_names if n.endswith("Extension"))
common_extension_props = [p for p in property_names[EXT_ORDER[0]]
                          if all(p in property_names[e] for e in EXT_ORDER)]
assert sorted(common_extension_props) == ["activity_id", "contract_id", "entity_id",
                                          "extension_id", "observation_date"]
extension_specific = {e: [p for p in property_names[e] if p not in common_extension_props]
                      for e in EXT_ORDER}
assert sorted(len(v) for v in extension_specific.values()) == [2, 2, 2, 2, 2, 2, 3, 3]
CURRENCY_USERS = property_users["currency_code"]
assert sorted(CURRENCY_USERS) == ["Activity", "Company", "Contract", "FinancialEffect",
                                  "FinancialMetric", "GroupKPI"]
SECTOR_JP = {"Metals": "金属", "Energy": "エネルギー", "Foods": "食品", "Finance": "金融"}
ENTITY_RELATIONS_USED_BY_Q4 = ["EventHasEffect", "EffectOnContract", "ContractOwnedByCompany"]
assert [(rel_by_name[n]["from"], rel_by_name[n]["to"]) for n in ENTITY_RELATIONS_USED_BY_Q4] == [
    ("RiskEvent", "FinancialEffect"), ("FinancialEffect", "Contract"), ("Contract", "Company")]
NEW_FEATURE_SOURCES = "Microsoft Learn：fabric/iq/ontology（2026-10-26に確認、Preview）"


def chip(slide, label, x, y, w, color=TEAL, fill=PALE, size=10.5):
    shp = box(slide, x, y, w, .27, fill, rounded=True, flat=True)
    shp.adjustments[0] = .5
    text(slide, label, x, y + .04, w, .2, size, color, True, PP_ALIGN.CENTER)


# New 1
s = new_slide(
    "新しいOntologyで増えた機能",
    "新しい体験では、定義を「整理する」「意味を足す」「運用する」ための機能が増えました。\n"
    "このデモで使っているのは、実体、関係、データの結び付け、グラフ、DA連携が中心です。",
    "新機能 1  |  全体像",
    """【目安】60秒。「新しい機能は何が増えましたか」と聞かれたときの入口です。
【話す】「Ontologyは新しい体験になり、定義を整理する機能、意味を足す機能、運用する機能が増えました。このデモは実体、関係、データの結び付けが中心です。新機能は、これから足していける状態です。」
【指す】左から「整理する」「意味を足す」「運用する」の順です。各行の右の印は、このデモのOntologyの現状です（2026-10-26に実機で確認）。印のない行は、今回確認していません。
【現状の数】名前空間は「既定値」の1つだけで、17実体と39関係が所属します。ルールは0件、共有プロパティは0件、メトリックは0件です。メタデータは、会社（Company）の説明が1件あるだけです。同義語、追加の属性、リソースリンクは、17実体のすべてで0件でした（17実体の詳細画面を確認）。継承はなく、バージョン履歴も0件です。172個のプロパティの説明は、リポジトリのOntology定義（reports/ontology_runtime_current.json）で0件です。
【各機能の一言】名前空間：同じ名前の概念を領域で区別する。継承：共通の親から実体を派生する。共有プロパティ：1つの定義を複数の実体が参照する。ルール：守る条件を自然言語で書き、実体やプロパティに結ぶ。メタデータ：説明、同義語（実体型だけに付けられる）、単位などの追加の属性。メトリック：Power BIセマンティックモデルのDAXメジャーを取り込む（セマンティックモデルから生成したOntologyで使う）。リソースリンク：関連するレポートやセマンティックモデルを実体に紐付ける。バージョン履歴：名前を付けた版を保存し、復元できる。インポート／エクスポート：.ttl、.rdf、.owlのファイルを空のOntologyに取り込み、Turtle（.ttl）かRDF/XML（.rdf）で書き出せる（対応する範囲のみで、OWLの制約などは失われ得る）。Ontology agent：定義やバインディングの案を作り、確認してから反映する。MCPサーバー：外部のAIや開発ツールからOntologyを参照する。
【表示の切替】キャンバスの右上に「完全なオントロジー」「近隣」「継承」の切替があります。「近隣」は、選んだ実体と、直接つながる実体だけを表示します。Activityを選ぶと、Activityと13実体が表示されることを確認しました。17実体が絡んで見えるときの説明に使えます。
【旧体験】旧体験は2027-01-31に廃止されます。このデモのOntologyは新体験で、画面にも「これは新しいオントロジ エクスペリエンスです」と表示されます。お客様が旧体験で作っている場合は、画面のバナーから、新体験にコピーを作成して移行します。下流のエージェントやダッシュボードの再接続と、ルールの再作成が必要です。
【避ける】すべてPreviewの機能です。製品の恒久的な仕様として説明せず、実演の前に現在の画面で確認します。「このデモは新機能を使っている」とは言いません。
【つなぎ】「気になる3つ、名前空間、ルール、共有プロパティを、順に説明します。」""",
    [NEW_FEATURE_SOURCES + " overview", "reports/ontology_runtime_current.json（namespaces、rules、metrics、reusableProperties、説明、同義語、リソースリンクの件数）",
     "実機確認 2026-10-26：名前空間、ルール、共有プロパティ、バージョン履歴、近隣ビュー、17実体の詳細（説明、同義語、追加のメタデータ、継承、リソース）"],
    "旧体験は2027-01-31に廃止。印はこのデモの現状（2026-10-26に実機確認）、印のない行は未確認")
n1_cols = [
    ("整理する", "同じ名前や共通項目を管理する", TEAL, [
        ("名前空間", "同じ名前の概念を領域で区別", "既定値のみ"),
        ("継承", "共通の親から実体を派生", "なし"),
        ("共有プロパティ", "1つの定義を複数の実体で参照", "0件"),
        ("表示の切替", "全体・近隣・継承の3つのビュー", "使える"),
    ]),
    ("意味を足す", "言葉の決まりや単位を持たせる", PURPLE, [
        ("ビジネスルール", "守る条件を文章で定義", "0件"),
        ("メタデータ", "説明・同義語・追加の属性", "説明1件のみ"),
        ("メトリック", "DAXメジャーを取り込む", "0件"),
        ("リソースリンク", "関連レポートを実体に紐付け", "0件"),
    ]),
    ("運用する", "版の管理、移行、AIとの連携", CORAL, [
        ("バージョン履歴", "名前付きの版を保存・復元", "0版"),
        ("インポート／エクスポート", "TTL・RDF・OWLを取り込み、TTL・RDFで出力", None),
        ("Ontology agent", "定義やバインディングの案を作成", None),
        ("MCPサーバー", "外部のAIからOntologyを参照", None),
    ]),
]
for c, (head, sub, accent, items) in enumerate(n1_cols):
    x = .62 + c * 4.09
    box(s, x, 2.62, 3.89, 3.8, WHITE, rounded=True)
    text(s, head, x + .25, 2.76, 3.4, .36, 20, accent, True)
    text(s, sub, x + .25, 3.17, 3.5, .26, 11.5, MUTED)
    for r, (name, desc, status) in enumerate(items):
        y = 3.58 + r * .7
        text(s, name, x + .25, y, 2.3 if status else 3.4, .28, 14, NAVY, True)
        text(s, desc, x + .25, y + .3, 3.4, .26, 11, INK)
        if status:
            chip(s, status, x + 3.89 - .2 - 1.15, y + .02, 1.15, accent, PALE, 10)

# New 2
s = new_slide(
    "名前空間：同じ名前を領域ごとに区別する",
    "同じ名前で意味が違う概念を、1つのOntologyの中で区別して持てます。\n"
    "概念は必ず1つの名前空間に属し、指定しなければ「既定値」に入ります。",
    "新機能 2  |  名前空間（Namespace）",
    """【目安】90秒。
【話す】「同じ『顧客』という言葉でも、営業と経理では意味が違います。名前空間は、概念にどの領域のものかという所属を与えて、同じ名前の概念を1つのOntologyの中で区別する仕組みです。」
【指す】左の図で、営業と経理に同じCustomerがあることを示します。次に4つの要点、右の「このデモに当てはめると」の順です。左の図は、公式ドキュメントの例を借りた一般的な例で、このデモの実体ではありません。
【仕組み】すべてのOntologyに「既定値（Default）」の名前空間があり、削除できません。実体型と関係型は、必ず1つの名前空間に属します。作成時に「追加の構成」（公式ドキュメントの表記）で指定するか、作成後に付け替えます。付け替えで同名の概念が重複する場合は、ブロックされます。名前空間の作成と同時には、概念を割り当てられません。作成してから割り当てます。名前、説明、接頭辞は、参照を壊さずに変更でき、名前を変えると、概念への参照は自動で追従します。参照が残っている名前空間は、削除できません。エクスプローラーでは、名前空間でグループ化やフィルターができます。RDFの名前空間の考え方に沿っているので、業界の標準語彙を取り込んで対応づけるときにも使えます。
【画面（日本語UI）】リボンの「名前空間」で作成、「名前空間の管理」で一覧と編集です。このデモは、「名前空間の管理」に「既定値」の1行だけがあり、「17件のエンティティ型; 39件のリレーションシップ」と表示されます（2026-10-26に確認）。
【このデモの案】本体9実体をcore、業種拡張8実体を4業種（metals、energy、foods、finance）に分ける案です。実装していません。名前は説明者の設計案です。
【注意】公式ドキュメントの制限事項として、名前空間を使う実体と関係は、グラフ上のノードとエッジのラベルに、名前空間が接頭辞として付きます。確認済みのDAの質問やGQLの問い合わせが変わり得るため、導入後はQ1～Q4を再実行します。このデモのグラフでの挙動は検証していません。
【使いどころ】複数の部門やチームがそれぞれ概念を持つ大きなOntology、外部の標準語彙の取り込み、所有者の境界を分けたい場合に効きます。最初の検証では、既定値のままで困りません（説明者の提案）。
【避ける】「名前空間を入れると回答が良くなる」とは言いません。整理と衝突回避のための機能です。
【つなぎ】「次は、業務の決まりを書くルールです。」""",
    [NEW_FEATURE_SOURCES + " how-to-use-namespaces（2026-09-10更新）", NEW_FEATURE_SOURCES + " how-to-use-ontology-graph（グラフ化の制限事項）",
     "ontology_tmdl/namespaces/default.tmdl", "実機確認 2026-10-26：名前空間の管理（既定値のみ）"],
    "規模が大きくなってから効く機能。最初の検証は既定値のままでよい（説明者の提案）")
box(s, .62, 2.62, 6.2, 3.82, WHITE, rounded=True)
text(s, "名前空間：概念に「所属」を与える", .85, 2.76, 5.7, .34, 17, TEAL, True)
dashed_frame(s, .85, 3.2, 5.74, 1.55, TEAL)
text(s, "1つのOntology", 1.0, 3.27, 3.0, .24, 11, MUTED)
for k, (ns, nm_, meaning) in enumerate([("営業 名前空間", "Customer", "売上先を表す"),
                                         ("経理 名前空間", "Customer", "請求先を表す")]):
    bx = 1.05 + k * 2.8
    box(s, bx, 3.62, 2.5, 1.0, PALE, rounded=True, flat=True)
    text(s, ns, bx + .17, 3.7, 2.2, .24, 11.5, TEAL, True)
    text(s, nm_, bx + .17, 3.97, 2.2, .3, 17, NAVY, True)
    text(s, meaning, bx + .17, 4.32, 2.2, .24, 11, MUTED)
text(s, "≠", 3.55, 3.9, .3, .4, 20, CORAL, True, PP_ALIGN.CENTER)
text(s, "・全てのOntologyに「既定値」があり、削除できない\n"
        "・概念は必ず1つに所属。作成後も付け替えられる\n"
        "・名前を変えても、参照は自動で追従する\n"
        "・外部の語彙（RDFなど）と対応づけられる",
     .9, 4.95, 5.6, 1.4, 12, INK)
box(s, 7.0, 2.62, 5.7, 3.82, WHITE, rounded=True)
text(s, "このデモに当てはめると（案）", 7.25, 2.76, 5.2, .34, 17, TEAL, True)
text(s, "現状：「既定値」1つに、17実体と39関係が所属（実機確認）", 7.25, 3.2, 5.3, .28, 12, INK)
text(s, "案：本体と4つの業種に分ける", 7.25, 3.58, 5.2, .28, 12.5, NAVY, True)
box(s, 7.25, 3.95, 5.2, .5, PALE, rounded=True, flat=True)
text(s, "core：本体9実体（会社、イベント、契約など）", 7.42, 4.08, 4.9, .26, 11.5, INK)
for k, (ns, sec_jp) in enumerate([("metals", "金属"), ("energy", "エネルギー"),
                                   ("foods", "食品"), ("finance", "金融")]):
    bx = 7.25 + k * 1.3
    box(s, bx, 4.55, 1.2, .78, PALE, rounded=True, flat=True)
    text(s, ns, bx + .1, 4.62, 1.0, .26, 12, PURPLE, True)
    text(s, f"{sec_jp}\n日米の2実体", bx + .1, 4.9, 1.05, .4, 9.5, INK)
text(s, "注意：既定値以外に入れると、グラフ上のラベルに名前空間が接頭辞で付く（公式の制限事項）。"
        "DAの問い合わせが変わり得るため、導入後はQ1〜Q4を再実行する。",
     7.25, 5.5, 5.2, .85, 10.5, CORAL)

# New 3
s = new_slide(
    "ビジネスルール：業務の決まりを文章で定義する",
    "守る条件を自然言語で書き、実体、プロパティ、関係に結び付けます。\n"
    "Ontologyは実行せず、エージェントなどが読んで回答の根拠にします。",
    "新機能 3  |  ルール（Business rules）",
    """【目安】90秒。
【話す】「いまは、数字の扱いの決まりをDAの指示文に書いています。たとえば、actualとbaselineを足さない、51%の子会社も全額連結にする、などです。ルール機能を使うと、こうした決まりをOntologyに書いて、関係する実体やプロパティに結び付けられます。」
【指す】左の4行（DAの指示文にある決まりと、結び付ける先の案）、右の「仕組みと限界」、その下の囲みの順です。
【仕組み】ビジネスルールは、自然言語の文章です。「名前」と「定義」が必須で、リンクする実体、プロパティ、関係、説明、追加の属性（category、owner、policyFamilyなど）は任意です。定義は、必ず守る、してはいけない、通常は守る、のどれかの形で書くと伝わります。リンクした実体やプロパティの名前を変えると、ルールの参照も更新され、リンク先を削除すると、ルールから外れます。ルールを削除すると、すべてのリンク先から外れ、元に戻せません。
【画面（日本語UI）】エクスプローラーの「概要」の「ルール」を開き、「ルールの作成」を押します。「名前」「定義」「リンクされたオントロジー要素」「メタデータ」（説明、追加のメタデータ）を入力して「作成する」です。このデモのルールのページは「最初のルールの作成を開始する」と表示され、0件でした（2026-10-26に確認。作成画面を開いたあと、保存せずにキャンセルしました）。
【読み取り】認可されたエージェントやアプリは、ontology MCPのlist_ontology_rulesで、ルールの名前、定義、リンク先、説明、追加の属性を取得できます。読み取り専用で、ルールの解釈と運用は利用側の役割です。
【限界】ルールは文章で、Ontologyはデータに対して実行せず、アクションも起こしません。旧体験にあったActivatorとのルール連携は、新体験にはありません。Activatorは、Ontologyとは別に設定します。
【このデモの案】4件は、prompts/da_instructions_ja.txtの19、23～24、26、38、45行目（2026-10-26の更新後の行番号）にある決まりを言い換えたものです。実装していません。
【未検証】Fabric Data Agentが、Ontologyのルールを自動で読んで回答に使うかは、今回確認していません。公式ドキュメントは、エージェントが根拠づけに使えると説明しますが、利用するエージェントごとに対応が異なり得ます。当面は、DAの指示文を残して併用し、検証してから移します。
【避ける】「ルールを書けばAIが必ず守る」「ルールで自動チェックする」とは言いません。
【つなぎ】「次は、同じ項目を何度も定義しなくて済む共有プロパティです。」""",
    [NEW_FEATURE_SOURCES + " how-to-use-rules（2026-09-15更新）", "prompts/da_instructions_ja.txt（19、23～24、26、38、45行目）",
     "実機確認 2026-10-26：ルールのページ（0件）と作成画面の項目名"],
    "決まりをDAの指示文から切り出して1か所で管理できる。エージェントが読む前提で、実行はしない")
text(s, "ルール定義（案）", .85, 2.64, 4.6, .3, 13, TEAL, True)
text(s, "結び付ける先", 5.6, 2.64, 2.3, .3, 13, TEAL, True)
for i, (rule_text, links) in enumerate([
    ("actualとbaselineは合算しない。\nactualにdeltaを再加算しない。",
     ["FinancialMetric.scenario_tag", "GroupKPI.scenario_tag"]),
    ("支配する子会社は、持分51%でも全額連結する。\n全額連結の値に持分比率を掛けない。",
     ["Company.ownership_ratio"]),
    ("横断イベントのprimary_sectorは'all'。\n4セクターに絞って落とさない。",
     ["RiskEvent.primary_sector"]),
    ("効果と根拠はevidence_idで結ぶ。\nevent_idだけで結ぶと件数が増える。",
     ["EffectExplainedByEvidence（関係）"]),
]):
    y = 3.0 + i * .85
    box(s, .62, y, 7.4, .77, WHITE, rounded=True)
    text(s, rule_text, .8, y + .1, 4.65, .6, 11.5, INK)
    for k, link in enumerate(links):
        chip(s, link, 5.6, y + (.08 if len(links) == 2 else .25) + k * .34, 2.3, TEAL, PALE, 9.5)
box(s, 8.2, 2.62, 4.5, 3.82, WHITE, rounded=True)
text(s, "仕組みと限界", 8.45, 2.76, 4.0, .34, 17, TEAL, True)
text(s, "・書く：名前と定義は必須。リンク先などは任意\n"
        "・場所：エクスプローラーの「概要」→「ルール」\n"
        "・読む：ontology MCPのlist_ontology_rulesで取得（読み取り専用）\n"
        "・限界：ルールは実行されない\n"
        "・旧体験のActivator連携は新体験にない\n"
        "・現状：このデモは0件（実機確認）",
     8.45, 3.2, 4.0, 2.0, 11.5, INK)
box(s, 8.4, 5.3, 4.1, .98, PALE, rounded=True, flat=True)
text(s, "DAがルールを自動で読むかは未検証。DAの指示文を残して併用し、検証してから移す。",
     8.58, 5.42, 3.75, .8, 11, NAVY, True)

# New 4
s = new_slide(
    "共有プロパティ：同じ定義を1回だけ書く",
    "currency_codeのように、複数の実体で同じ意味の項目を、1回だけ定義して参照できます。\n"
    "共有するのは定義だけで、データの結び付け（バインディング）は実体ごとです。",
    "新機能 4  |  共有プロパティ（Shared properties）",
    """【目安】90秒。
【話す】「currency_codeのような項目は、複数の実体に同じ意味で出てきます。共有プロパティを使うと、型や説明を1回だけ定義して、各実体が参照できます。コピーではなく参照なので、定義を直すと、参照している実体に反映されます。」
【指す】左の図で、上の共有プロパティから、下の6実体へ矢印が伸びている点を示します。次に右の棒グラフで、このデモの共通項目の数を示します。
【仕組み】共有プロパティは、Ontologyの最上位にある定義で、データ型、説明などを持ちます。名前はOntologyの中で一意です。実体の「構成」ページで、プロパティのバインディングの管理から既存のプロパティを追加する操作を選び、共有プロパティを参照します（英語の公式ドキュメントの操作名を訳したもので、日本語UIの表記は未確認です）。データのバインディングは、その実体の中だけの設定です。共有プロパティ自体はバインドされないままで、別の実体が別々にバインドします。説明などのメタデータは、実体ごとに上書きできます（上書きには印が付き、1操作で元に戻せます）。共有側の値が変わっても、上書きした項目は維持されます。
【参照を外す】その実体にローカルのバインディング、上書き、関係の端点がある場合だけ、選択肢が出ます。「ソフトデタッチ」は、プロパティをローカルのプロパティとして残し、バインディング、上書き、関係の端点を保持します。「完全削除」は、プロパティとローカルの設定を消します。ローカルの設定がなければ、選択なしでそのまま外れます。共有プロパティ自体を削除するときは、影響分析（バインディング、関係、上書き）が表示され、確認後は、影響を受ける実体にローカルのコピーが残ります。
【出所】実体の構成ページで、プロパティの出所を「ローカル」「継承」「共有」で絞り込めます。
【このデモの実測】17実体のプロパティ名を、ontology_tmdlから集計しました。entity_idは15実体、contract_idは12、activity_idは11、observation_dateは8（業種拡張の8実体）、source_lakehouseは7、currency_codeは6、event_idは3実体に出てきます。
【向き不向き】意味が同じ項目が向きます。currency_codeとsource_lakehouseが候補です。entity_idは、会社では主キー、他の実体では会社への参照なので、意味を整理してから共有します。共有プロパティを実体のキーにできるかは、今回確認していません。
【画面（日本語UI）】エクスプローラーの「概要」の「共有プロパティ」で一覧を開き、「新しい共有プロパティ」で作成します。このデモは「共有プロパティはまだありません」と表示されます（2026-10-26に確認）。
【避ける】共有プロパティは、データを共有する機能ではありません。「1つの列を複数の実体で使い回す」とは言いません。
【つなぎ】「似た機能の継承も見ます。業種拡張の8実体が題材です。」""",
    [NEW_FEATURE_SOURCES + " how-to-reuse-properties（2026-10-20更新）", "ontology_tmdl/entities/（プロパティ名の出現数を集計）",
     "実機確認 2026-10-26：共有プロパティのページ（0件）"],
    "共有するのは定義（型・説明）だけ。バインディングは実体ごと。出所はローカル／継承／共有で絞り込める")
box(s, .62, 2.62, 6.4, 3.82, WHITE, rounded=True)
box(s, 1.0, 2.82, 5.64, .8, TEAL, rounded=True, flat=True)
text(s, "共有プロパティ　currency_code", 1.22, 2.9, 5.2, .3, 17, WHITE, True)
text(s, "型：string　説明：ISO 4217の通貨コード（1回だけ定義）", 1.22, 3.27, 5.2, .26, 11.5, WHITE)
arrow(s, 3.82, 3.64, 3.82, 3.98, TEAL)
text(s, "参照する（コピーではない）", 3.98, 3.7, 2.6, .24, 10.5, MUTED)
box(s, 1.0, 4.02, 5.64, 1.58, PALE, rounded=True, flat=True)
text(s, "参照する実体（このデモでは6つ）", 1.18, 4.1, 5.2, .26, 12, TEAL, True)
for k, name in enumerate(["Company", "Contract", "Activity", "FinancialEffect", "FinancialMetric", "GroupKPI"]):
    assert name in CURRENCY_USERS
    bx, by = 1.18 + (k % 3) * 1.82, 4.45 + (k // 3) * .56
    box(s, bx, by, 1.7, .46, WHITE, rounded=True, flat=True)
    text(s, name, bx, by + .12, 1.7, .24, 11, NAVY, True, PP_ALIGN.CENTER)
text(s, "共有するのは定義だけ。バインディングは実体ごとに別で、説明は上書きできる。\n"
        "参照を外すとき、ローカルの設定があれば、残すか削除するかを選べる。",
     1.0, 5.7, 5.64, .72, 10.5, INK)
box(s, 7.2, 2.62, 5.5, 3.82, WHITE, rounded=True)
text(s, "このデモの共通項目（17実体で集計）", 7.45, 2.76, 5.0, .34, 15, TEAL, True)
for k, prop in enumerate(SHARED_CANDIDATES):
    n_users = len(property_users[prop])
    y = 3.24 + k * .33
    text(s, prop, 7.45, y, 1.9, .26, 11.5, INK)
    box(s, 9.4, y + .03, 2.6 * n_users / 15, .2, TEAL, flat=True)
    text(s, f"{n_users}", 9.4 + 2.6 * n_users / 15 + .1, y, .5, .26, 11.5, NAVY, True)
box(s, 7.4, 5.65, 5.1, .68, PALE, rounded=True, flat=True)
text(s, "向くのは意味が同じ項目（currency_code、source_lakehouse）。entity_idは、会社では主キー、他では参照なので、意味を整理してから共有する。",
     7.55, 5.7, 4.85, .6, 10, NAVY)

# New 5
s = new_slide(
    "継承：共通の親から実体を派生させる",
    "共通の項目を親の実体に持たせ、子の実体が引き継ぎます。親は1つだけです（単一継承）。\n"
    "このデモの業種拡張8実体は、5項目が共通で、当てはめやすい題材です。",
    "新機能 5  |  継承（Inheritance）",
    """【目安】90秒。
【話す】「業種拡張の8実体は、5項目が共通で、残りの2〜3項目だけが業種ごとに違います。継承は、共通の項目を親の実体に持たせて、子の実体がそれを引き継ぐ仕組みです。」
【指す】左の図で、上の親、下の8つの子の順に示します。右の「引き継ぐ」「引き継がない」「決まり」の順に読みます。
【このデモの実測】8つの業種拡張は、extension_id、activity_id、contract_id、entity_id、observation_dateの5項目が共通です。固有の項目は、金属が純度と製錬所の稼働率、エネルギーが貯蔵率と配分量、食品が温度と廃棄率（各2項目）、金融が想定元本、信用スプレッド、ヘッジ比率（3項目）です。日本と米国で項目は同じです。
【仕組み】親は1つだけです（単一継承）。子は、親のプロパティ（項目ごとの説明などを含む）と、実体の説明、同義語、追加の属性を引き継ぎ、親子であることを示すis-aの関係を持ちます。引き継がないものは、データバインディングと、通常の関係です。子ごとに表を結び、関係を定義します。親の項目は、子で削除できず、名前とデータ型も変えられません。説明などは、子で上書きでき、上書きを戻すと親の値に戻ります。親の項目を直すと、上書きしていない子へ即時に反映されます。循環する継承は拒否されます。
【画面】実体を追加するダイアログの「追加の構成」（公式ドキュメントの表記）で、継承元の実体を選びます。キャンバスの右上の「継承」ビューで、親子の関係を表示します。このデモは継承がなく（会社の詳細は「直接の親要素：なし」、子孫0）、Activityを選んで「継承」を押すと、1実体だけが表示されます（2026-10-26に確認）。
【減らせるもの】共通5項目の定義です。関係24本（8実体×3本）と、表の結び付け8本は、子ごとに必要です。
【既知の制限】公式のトラブルシューティングに、「親の実体を照会して、継承する実体のインスタンスを取得できない」が既知の問題として載っています。照会は子の実体に対して行います。グラフ化では、子のノードに、親の項目と子の項目が含まれます。RDFやOWLのインポートでは、複数の親を持つ定義は、1つの親だけが保持されます。
【未確認】既存の実体を、後から派生型に付け替えられるかは、今回確認していません。実体を追加するときに継承元を選ぶ手順だけを、公式ドキュメントで確認しました。
【避ける】「継承すると関係も引き継ぐ」とは言いません。
【つなぎ】「5つの機能を、入れる順番と注意点にまとめます。」""",
    [NEW_FEATURE_SOURCES + " how-to-use-inheritance（2026-10-20更新）", NEW_FEATURE_SOURCES + " resources-troubleshooting、how-to-use-ontology-graph、how-to-import-export",
     "ontology_tmdl/entities/*Extension.tmdl（共通5項目と固有項目）", "実機確認 2026-10-26：継承ビューと会社の詳細（親なし）"],
    "減らせるのは共通5項目の定義。関係24本と表の結び付け8本は、子ごとに必要")
box(s, .62, 2.62, 7.3, 3.82, WHITE, rounded=True)
box(s, 1.0, 2.8, 6.54, .92, TEAL, rounded=True, flat=True)
text(s, "SectorExtension（案：親の実体）", 1.2, 2.87, 6.1, .32, 17, WHITE, True)
text(s, "共通5項目：" + " ／ ".join(common_extension_props), 1.2, 3.31, 6.2, .26, 10.5, WHITE)
arrow(s, 4.27, 3.74, 4.27, 4.08, TEAL)
text(s, "is-a（派生）", 4.42, 3.78, 1.6, .24, 10.5, MUTED)
box(s, 1.0, 4.12, 6.54, 1.82, PALE, rounded=True, flat=True)
text(s, "子の実体（8つ）：共通5項目を引き継ぎ、固有項目だけを持つ", 1.18, 4.19, 6.2, .26, 12, TEAL, True)
for k, name in enumerate(EXT_ORDER):
    sec, cc = re.match(r"(Metals|Energy|Foods|Finance)(JP|US)", name).groups()
    bx, by = 1.18 + (k % 4) * 1.6, 4.5 + (k // 4) * .72
    box(s, bx, by, 1.5, .66, WHITE, rounded=True, flat=True)
    text(s, name, bx + .06, by + .06, 1.4, .2, 9, TEAL, True, PP_ALIGN.CENTER)
    text(s, f"{SECTOR_JP[sec]}・{'日本' if cc == 'JP' else '米国'}", bx + .06, by + .26, 1.4, .2, 9, INK, False, PP_ALIGN.CENTER)
    text(s, f"＋固有{len(extension_specific[name])}項目", bx + .06, by + .45, 1.4, .2, 9, MUTED, False, PP_ALIGN.CENTER)
text(s, "既知の制限：親を照会しても、子のインスタンスは返らない（公式の既知の問題）。",
     1.0, 6.02, 6.54, .36, 10.5, CORAL)
box(s, 8.1, 2.62, 4.6, 3.82, WHITE, rounded=True)
text(s, "引き継ぐ", 8.35, 2.74, 4.1, .28, 13, TEAL, True)
text(s, "・プロパティと、その説明など\n・実体の説明、同義語（子で上書き可）\n・親子（is-a）の記録", 8.35, 3.06, 4.1, .8, 11, INK)
text(s, "引き継がない", 8.35, 3.98, 4.1, .28, 13, CORAL, True)
text(s, "・データバインディング（子ごとに表を結ぶ）\n・通常の関係（子ごとに定義する）", 8.35, 4.3, 4.1, .6, 11, INK)
text(s, "決まり", 8.35, 5.0, 4.1, .28, 13, NAVY, True)
text(s, "・親の項目は、子で削除・改名・型変更できない\n・親を直すと、子へ即時に反映される\n・現状：このデモは継承なし（実機確認）", 8.35, 5.32, 4.1, 1.0, 11, INK)

# New 6
s = new_slide(
    "新機能を入れる順番と注意点",
    "効果が大きく、確認済みのデモに影響しにくいものから入れます。\n"
    "以下は公式ドキュメントとこのデモの構成から整理した、説明者の提案です（実機では未検証）。",
    "新機能 6  |  導入の優先順位",
    """【目安】60秒。
【話す】「新機能は、効果が大きく、いまの確認済みデモに影響しにくいものから入れるのが安全です。メタデータとルール、共有プロパティ、継承、名前空間の順を提案します。」
【指す】上から順に、効果、このデモへの影響、入れ方の案を読みます。「影響」の印は、小が緑、中が紫、大が赤です。
【前提】この順番は、公式ドキュメントとこのデモの構成から、説明者が整理した提案です。実機では入れて確認していません。
【1 メタデータ】説明、同義語、追加の属性は、エージェントの回答の正確さを上げるとされています。特に、単位や用語の定義が必要な質問に効きます。同義語は、実体型にだけ付けられます（プロパティや関係には付けられません）。このデモでは、金額の列の説明に単位（円、百万円、億円）を書く、「全額連結」「持分管理」「baseline」「actual」の意味は、プロパティの説明やルールで定義する、実体名に日本語の呼び名を同義語として付ける（例：GroupKPIに「全社KPI」）、が候補です。現状は、会社の説明が1件あるだけです。
【2 ルール】ビジネスルールのページ（第7章）のとおりです。DAが読むかを確認するまでは、DAの指示文と併用します。
【3 共有プロパティ】currency_codeとsource_lakehouseが候補です。
【4 継承】業種拡張の8実体が候補です。既存の実体を付け替えられるかは未確認で、バインディングと関係の作り直しが必要になる可能性があります。親の実体の照会は、既知の問題です。
【5 名前空間】グラフ上のラベルに接頭辞が付くため、DAの問い合わせに影響する可能性があります。実体が増えて、複数チームで持つ段階で導入します。
【どの機能でも】Previewです。Ontologyを変更したら、公開済みDAでQ1～Q4を再実行し、結果を確かめます。
【避ける】「入れると必ず精度が上がる」とは言いません。
【つなぎ】「進め方の全体は、第5章の「小さく始める」のページにまとめています。」""",
    [NEW_FEATURE_SOURCES + " overview、how-to-add-metadata、how-to-use-rules、how-to-use-inheritance、how-to-use-namespaces",
     NEW_FEATURE_SOURCES + " how-to-use-ontology-graph、resources-troubleshooting", "提案：説明者の整理（実機では未検証）"],
    "どれもPreview。入れたら、DAの確認済み質問（Q1〜Q4）を必ず再実行して結果を確かめる")
for label, x, w in [("順", .82, .4), ("機能", 1.25, 2.2), ("効果", 3.55, 3.1),
                    ("このデモへの影響", 6.85, 2.9), ("入れ方（案）", 10.0, 2.6)]:
    text(s, label, x, 2.64, w, .28, 13, TEAL, True)
for i, (feat, effect, level, impact, how) in enumerate([
    ("メタデータ", "単位や用語を、回答に反映しやすくする", "小",
     "説明を足すだけ。構造は変わらない", "金額列の説明に単位\n実体名に日本語の同義語"),
    ("ビジネスルール", "決まりを1か所で管理し、エージェントと共有", "小",
     "DAが読むかは未検証。併用が前提", "ビジネスルールのページの4件から"),
    ("共有プロパティ", "共通項目の定義の重複を減らす", "小",
     "バインディングは実体ごとのまま", "currency_code、source_lakehouse"),
    ("継承", "業種拡張8実体の共通5項目を、親に集約する", "中",
     "付け替え可否は未確認\n親の照会は既知の問題", "業種拡張だけ。照会は子に対して行う"),
    ("名前空間", "実体が増えたときの整理と、同名の衝突防止", "大",
     "グラフのラベルに接頭辞が付く\nDAの問い合わせを再確認", "複数チームで持つ段階で"),
]):
    y = 3.0 + i * .68
    box(s, .62, y, 12.08, .6, WHITE, rounded=True)
    text(s, str(i + 1), .82, y + .12, .4, .36, 20, TEAL, True)
    text(s, feat, 1.25, y + .17, 2.25, .3, 14, NAVY, True)
    # centre each text block vertically in the .6in row (explicit line breaks are the only wraps here)
    def mid(value, size):
        return (.6 - (value.count("\n") + 1) * size * 1.25 / 72) / 2
    text(s, effect, 3.55, y + mid(effect, 11), 3.15, .5, 11, INK)
    chip(s, level, 6.85, y + .17, .38, {"小": TEAL, "中": PURPLE, "大": CORAL}[level], PALE, 10.5)
    text(s, impact, 7.33, y + mid(impact, 10.5), 2.55, .5, 10.5, INK)
    text(s, how, 10.0, y + mid(how, 10.5), 2.6, .5, 10.5, INK)

# New 7
s = new_slide(
    "小さく始める：最小のOntologyから広げる",
    "最初は、答えたい質問1つに必要な、数個の実体と関係だけを作ります。\n"
    "このデモなら、17実体・39関係のうち、4実体・3関係から始められます。",
    "導入の進め方  |  小さく始める",
    """【目安】90秒。「作るのが大変そう」という質問に答えるページです。
【話す】「最初から17実体、39関係を作る必要はありません。答えたい質問を1つ決めて、その質問に出てくる名詞を実体に、名詞をつなぐ動詞を関係にします。」
【指す】左の5ステップを上から。次に右の「このデモに当てはめると」を、段階1、2、3の順に示します。
【このデモの段階】段階1は、RiskEvent、FinancialEffect、Contract、Companyの4実体と、EventHasEffect、EffectOnContract、ContractOwnedByCompanyの3関係です。Q4（契約ごとの影響額、子会社、商品、影響したイベント）に必要な範囲です。段階2は、Step03の図が使う8実体と7関係で、Q3（根拠までの経路）に必要な範囲です。段階3が、現在の17実体と39関係です。
【段階の数の根拠】段階1は、ontology_tmdlの関係定義から、段階2は、reports/step3_live_audit.jsonのdiagram_relationsから数えました。Q4の表の列（契約ID、子会社、商品、イベント、影響額）が、どの実体に当たるかを確かめた結果です。DAがQ4をOntologyだけで答える、という意味ではありません。DAの定量の集計は、LH_Group_SAPの表で行っています。
【近道】既存のPower BIセマンティックモデルがあれば、「Generate Ontology」で、テーブルを実体、列をプロパティ、リレーションを関係、DAXメジャーをメトリックとして自動生成できます（Preview）。生成後に、キーの確認、関係のバインディング、名前の整理が必要です。ただし、公式ドキュメントでは、セマンティックモデルにバインドしたOntologyは、Data Agentが動作しない既知の問題があり、グラフ化の対象外とされています。DAやグラフを使うなら、このデモのように、Lakehouseの表に結びます。
【実際の手間】定義は少なく済みます。時間がかかるのは、キーをそろえる、8社分を統合した表を作る、リンク表を用意する、といったデータ側の準備です。
【避ける】「数日でできる」と約束しません。データの状態で大きく変わります。
【つなぎ】「まず質問を1つ決めるところから、始められます。」""",
    [NEW_FEATURE_SOURCES + " how-to-generate-from-semantic-models、how-to-create-data-agent、how-to-use-ontology-graph",
     "ontology_tmdl/entityRelationships.tmdl", "reports/step3_live_audit.json:diagram_relations"],
    "最小のOntologyで価値を確かめてから広げる。17実体・39関係は、広げた後の姿")
box(s, .62, 2.62, 6.35, 3.82, WHITE, rounded=True)
text(s, "進め方（5ステップ）", .85, 2.76, 5.8, .34, 17, TEAL, True)
for i, (head, detail) in enumerate([
    ("質問を1つ決める", "例：どの契約が、どのイベントで減ったか"),
    ("名詞を実体にする", "質問に出てくる、イベント、効果、契約、会社"),
    ("動詞を関係にする", "イベント → 効果 → 契約 → 会社 のつながり"),
    ("表を1つずつ結び、質問で確かめる", "実体ごとに、Lakehouseの表とキー列を指定"),
    ("足すのは、その後", "実体・関係を増やし、ルールなどの新機能を足す"),
]):
    y = 3.22 + i * .62
    text(s, str(i + 1), .85, y + .03, .45, .4, 22, TEAL, True)
    text(s, head, 1.4, y, 5.4, .3, 13.5, NAVY, True)
    text(s, detail, 1.4, y + .3, 5.4, .26, 11, MUTED)
box(s, 7.15, 2.62, 5.55, 3.82, WHITE, rounded=True)
text(s, "このデモに当てはめると", 7.4, 2.76, 5.0, .34, 17, TEAL, True)
for i, (big, accent, head, detail) in enumerate([
    ("4 / 3", TEAL, "Q4：契約ごとの影響と対応", "RiskEvent → FinancialEffect → Contract → Company"),
    ("8 / 7", PURPLE, "Q3：根拠までの経路（Step03の図）", "図の8ノード（食品日本の拡張を含む）と7本の関係"),
    ("17 / 39", CORAL, "現在の全体", "業種拡張、出荷、会社への関係まで含む"),
]):
    y = 3.22 + i * .7
    box(s, 7.4, y, 5.05, .62, PALE, rounded=True, flat=True)
    text(s, big, 7.52, y + .13, 1.3, .36, 19, accent, True)
    text(s, head, 8.9, y + .06, 3.5, .26, 11.5, NAVY, True)
    text(s, detail, 8.9, y + .33, 3.5, .26, 9, INK)
text(s, "数字は「実体の数／関係の数」",
     7.4, 5.3, 5.0, .24, 10, MUTED)
box(s, 7.4, 5.58, 5.05, .75, PALE, rounded=True, flat=True)
text(s, "近道：既存のPower BIセマンティックモデルから自動生成できる（Preview）。ただしDAとの連携は既知の問題で、グラフ化も対象外。DAやグラフを使うならLakehouseへ結ぶ。",
     7.55, 5.62, 4.78, .7, 9.5, NAVY)

# Casual questions (placed after the FAQ appendix)
casual_before = json.loads((PROJECT / "reports/da_casual_baseline_2026-10-26.json").read_text())
casual_after = json.loads((PROJECT / "reports/da_casual_after_2026-10-26.json").read_text())
assert [q["verdict"][:2] for q in casual_before["casual_questions"]] == [
    "OK", "OK", "OK", "OK", "NG", "OK", "OK", "NG", "一部", "一部"]
assert "10 of 10" not in casual_after["summary"] and "All ten" in casual_after["summary"]
CASUAL_MARK = {"OK": ("○", TEAL), "一部": ("△", PURPLE), "NG": ("×", CORAL)}
CASUAL_MEMO = {"A2": "縦軸が「3B」表記", "D": "FOD_NEGに結び付かず", "G": "「いま」が直近30日で0件",
               "H": "他のイベントが混在", "L": "契約×イベントで順位"}
s = new_slide(
    "カジュアルな質問に答えさせる",
    "指示文に言い回しの既定値を足したところ、短い質問10問が通るようになりました。\n"
    "足す前は6問OK、足したあとは10問すべてOKでした（確認した範囲の結果です）。",
    "",
    """【目安】60秒。「もっとカジュアルにDAへ質問できないか」と聞かれたときに使います。
【話す】「入力①〜④は、結果がぶれにくいように、条件を細かく書いた質問です。普段使いでは、もっと短く聞きたいはずです。そこで、DAの指示文に、言い回しの読み替えを足しました。たとえば、『今期』は期間全体、『予算』はbaseline、『食品の損失』は食品損失イベントです。」
【指す】左の表の「前」と「後」の列を指します。前は、10問のうち6問がOK、2問が一部、2問がNGでした。後は、10問すべてOKです。右のカードは、足した指示の要点です。
【前の問題】「食品の損失は、どの契約に響いたの？」は、食品損失イベントに結び付かず、食品全体の実績と基準の差を、現地通貨のまま契約別に並べました。「いま手を打つべき契約ってどれ？」は、「いま」が指示文の「直近30日」と読まれ、0件になりました。食品損失の根拠は、他のイベントが混ざりました。契約のトップ5は、契約とイベントの組で順位が付きました。推移のグラフは、縦軸が「3B」の表記でした。
【足した指示】期間（今期・いま＝期間全体、最近・直近＝直近30日）、言葉（予算＝baseline、稼ぎ・利益・儲け＝EBITDA）、イベントの呼び名（食品損失＝FOD_NEGなど）、業種と国、集計の単位（契約別は全イベントの合算、「どの契約」は全件）、通貨と単位（JPYで億円、グラフの軸も億円、太字記号なし）、期間の値（全日分を合計し、最終日の1行を期間の値にしない）、前の質問の引き継ぎ（指定がなければ全社・全イベント・今期全体）です。指示文は、約2,900字から約5,600字になりました（上限は15,000字）。
【確認の範囲】足したあとの結果は、公開前のドラフトで10問すべてと、追加の質問（最近どう？、持分で見ると、エネルギーの会社は、与信費用が増えたのは、など）、入力①〜④で確認しました。公開後は、公開済み版で、今期のEBITDA（3回）、金融の子会社（3回）、食品の損失から、いま手を打つべき契約への連続、入力①〜④を確認しました。
【経緯】最初に公開したあとの確認で、「今期のEBITDA」が1日分の値（1.27億円）で答えられる、金額の桁を誤る、日次の明細しか取得できず答えが出ない、といった失敗が見つかり、指示文を足して計5回公開し直しました。最後の版では、上の確認を通っています。
【変えていないもの】データ、Ontology、Lakehouse、データソースの設定は変えていません。変えたのはエージェントの指示だけです。
【避ける】「どんな質問でも答える」「常に同じ答えが出る」とは言いません。LLMのため、文面や、まれに取得の経路が実行ごとに変わります。""",
    ["reports/da_casual_baseline_2026-10-26.json", "reports/da_casual_after_2026-10-26.json",
     "prompts/da_instructions_ja.txt", "実機確認 2026-10-26：公開済みDA（5回目の公開後）"],
    "変えたのは指示文だけ。データ、Ontology、Lakehouseは変更していない")
text(s, "短い質問（入力した文）", .8, 2.66, 4.3, .28, 12.5, TEAL, True)
text(s, "前", 5.2, 2.66, .5, .28, 12.5, TEAL, True, PP_ALIGN.CENTER)
text(s, "後", 5.75, 2.66, .5, .28, 12.5, TEAL, True, PP_ALIGN.CENTER)
text(s, "前の問題", 6.4, 2.66, 2.0, .28, 12.5, TEAL, True)
for i, q in enumerate(casual_before["casual_questions"]):
    y = 3.0 + i * .335
    if i % 2 == 0:
        box(s, .65, y - .04, 7.85, .32, WHITE, flat=True)
    mark, color = CASUAL_MARK[q["verdict"][:2]]
    text(s, q["question"], .8, y, 4.35, .26, 11.5, INK)
    text(s, mark, 5.2, y - .02, .5, .28, 14, color, True, PP_ALIGN.CENTER)
    text(s, "○", 5.75, y - .02, .5, .28, 14, TEAL, True, PP_ALIGN.CENTER)
    text(s, CASUAL_MEMO.get(q["id"], ""), 6.4, y + .01, 2.1, .26, 10, MUTED)
box(s, 8.7, 2.62, 4.0, 3.82, WHITE, rounded=True)
text(s, "足した指示（要点）", 8.92, 2.74, 3.6, .3, 15, TEAL, True)
for k, (a_, b_) in enumerate([
    ("期間", "今期・いま＝期間全体、最近＝直近30日"),
    ("言葉", "予算＝baseline、稼ぎ・利益＝EBITDA"),
    ("イベント名", "食品損失＝FOD_NEG など10件"),
    ("業種と国", "金融、米国などをコードに対応づけ"),
    ("集計", "契約別は全イベントを合算して順位付け"),
    ("表記", "億円で表示。グラフの軸も億円"),
    ("期間の値", "全日分を合計（最終日だけにしない）"),
    ("引き継ぎ", "話題が変われば前の質問を持ち越さない"),
]):
    y = 3.1 + k * .37
    text(s, a_, 8.92, y, 1.05, .26, 11, PURPLE, True)
    text(s, b_, 9.95, y, 2.65, .3, 10.5, INK)
text(s, "指示文：約2,900字 → 約5,600字（上限15,000字）", 8.92, 6.12, 3.6, .25, 10, MUTED)

s = new_slide(
    "カジュアルに聞くときの使い方と限界",
    "短い質問でも、読み替えた前提を答えに1行で書くよう、指示文で指定しています。\n"
    "毎回同じ答えとは限らないので、使い分けと取り直しの方法を決めておきます。",
    "",
    """【目安】60秒。
【話す】「短い質問でも、読み替えた前提を答えに1行で書くよう、指示文で指定しています。たとえば『今期は2025-10から2026-09』のように書かれます。前提が違っていれば、その場で言い直せます。ただし、LLMなので毎回まったく同じ答えにはなりません。」
【指す】左の4行は、質問、DAの読み替え、答えの順です。右は、限界と使い分けです。
【読み替えの例】①今期のEBITDAは予算に対して：447.57億円、予算448.60億円、差は−1.03億円。②食品の損失は、どの契約に：食品損失イベント（FOD_NEG）の6契約で、各−1,650万〜−2,056万円（合計−1.14億円）。③いま手を打つべき契約は：今期全体の全ネガティブイベントを契約ごとに合算し、金融の4契約が上位（−4,283万〜−4,535万円）。④最近どう：直近30日（2026-09-01〜09-30）で、実績40.19億円、予算34.35億円、差は+5.84億円。
【限界】確認の途中では、期間や会社単位の質問で、5〜10回に1回ほどの割合で、日次の明細のまま答えが出ない、金額の桁を誤る、といった失敗が見られました。前の質問の対象を引き継ぐ失敗も、まれにありました。指示文を足して減らし、最後の版は確認を通りましたが、確認した回数は少なく、ゼロとは言えません。
【使い分け】実演の本編は、条件を細かく書いた入力①〜④の長い質問を使います。結果のぶれは小さくなりますが、毎回同じとは限りません。短い質問は、「もっと気軽に聞けますか」と聞かれたときの追加の見せ場に使います。話題を変えるときは、チャットをクリアしてから聞きます。答えが変なときは、同じ質問をもう一度入力します。金額は、答えの下の表（result）で確認します。
【実演のおすすめ】チャットをクリアしてから、次の3つを順に聞きます。「今期のEBITDAって、予算に対してどうだった？」「一番足を引っ張ったイベントは？」「食品の損失は、どの契約に響いたの？」。
【この先】読み替えのルールや同義語は、現在はDAの指示文に書いています。Ontologyのルール、同義語、説明に移せば、DA以外のエージェントとも共有できる可能性があります（第7章のビジネスルールのページ。DAがOntologyのルールを読むかは未検証です）。
【避ける】「質問の言い方は自由」「毎回同じ答え」とは言いません。取引先名など、データにない項目を聞かれたときの扱いは、DAの指示文では定めていません。Q4の質問文では、取引先名を作らず契約IDで示すよう指定しています。""",
    ["reports/da_casual_after_2026-10-26.json:known_limits", "prompts/da_instructions_ja.txt",
     "実機確認 2026-10-26：公開済みDA（5回目の公開後）"],
    "気軽に聞く場面は短い質問、結果のぶれを減らしたい場面は入力①〜④。金額は答えの下の表で確認する")
text(s, "読み替えの例", .85, 2.64, 4.0, .3, 15, TEAL, True)
for i, (q_, rd, ans) in enumerate([
    ("「今期のEBITDAは、予算に対してどうだった？」", "今期＝2025-10〜2026-09、予算＝baseline",
     "447.57億円（予算448.60億円、−1.03億円）"),
    ("「食品の損失は、どの契約に響いたの？」", "食品損失＝FOD_NEG、「どの契約」＝該当を全件",
     "6契約、−1,650万〜−2,056万円（計−1.14億円）"),
    ("「いま手を打つべき契約ってどれ？」", "いま＝今期全体、全ネガティブイベントを契約ごとに合算",
     "金融の4契約が上位（−4,283万〜−4,535万円）"),
    ("「最近どう？」", "最近＝直近30日（2026-09-01〜09-30）",
     "40.19億円（予算より＋5.84億円）"),
]):
    y = 3.02 + i * .85
    box(s, .62, y, 7.3, .78, WHITE, rounded=True)
    text(s, q_, .8, y + .06, 6.95, .24, 12, NAVY, True)
    text(s, "読み替え：" + rd, .8, y + .33, 6.95, .22, 10.5, MUTED)
    text(s, "答え：" + ans, .8, y + .54, 6.95, .22, 11, TEAL, True)
box(s, 8.1, 2.62, 4.6, 3.82, WHITE, rounded=True)
text(s, "限界と使い分け", 8.35, 2.74, 4.1, .3, 15, TEAL, True)
for lead, body, y, h in [
    ("限界", "毎回同じ文面にはならない。確認の途中では、5〜10回に1回ほど、日次の明細のまま答えが出ない、金額の桁を誤る、などの失敗があった。ゼロとは言えない", 3.12, .95),
    ("取り直し", "答えが変なときは、同じ質問をもう一度入力する。話題を変えるときは、チャットをクリアする", 4.12, .6),
    ("使い分け", "実演の本編は、条件を細かく書いた入力①〜④の長い質問を使う。短い質問は、追加の見せ場に使う", 4.98, .6),
    ("確認", "金額は、答えの下の表（result）で確認する", 5.82, .45),
]:
    text(s, lead, 8.35, y, 4.1, .24, 11.5, PURPLE, True)
    text(s, body, 8.35, y + .26, 4.15, h - .26, 11.5, INK)

# Chapter dividers are generated last, so that their page lists use the final page numbers.
LIGHT_ON_NAVY = "A9DCD9"
BODY_ON_NAVY = "E6EEF2"


def chapter_divider(ch):
    title = f"第{ch['n']}章 {ch['name']}"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    n = len(prs.slides)
    text(slide, "資料の章立て", .8, .55, 5, .3, 12, LIGHT_ON_NAVY, True)
    text(slide, f"第{ch['n']}章", .8, 1.2, 5, .75, 34, LIGHT_ON_NAVY, True)
    text(slide, ch["name"], .8, 2.0, 4.9, 1.6, 34 if len(ch["name"]) <= 9 else (30 if len(ch["name"]) <= 11 else 26), WHITE, True)
    text(slide, ch["summary"], .8, 3.95, 4.75, 1.5, 16, BODY_ON_NAVY)
    text(slide, "使う場面", .8, 5.6, 2, .3, 12, LIGHT_ON_NAVY, True)
    text(slide, ch.get("when_long", ch["when"]), .8, 5.92, 4.75, .8, 14, WHITE)
    count = len(ch["pages"])
    box(slide, 6.05, .55, 6.7, min(6.3, 1.2 + count * min(.52, 5.3 / count) + .3), WHITE, rounded=True)
    text(slide, f"この章のページ（{count}ページ）", 6.35, .78, 6.0, .35, 15, TEAL, True)
    row_h = min(.52, 5.3 / count)
    size = 14 if count <= 8 else 12
    for i, (page_title, _) in enumerate(ch["pages"]):
        y = 1.32 + i * row_h
        text(slide, f"p.{PAGE_OF[page_title]}", 6.35, y, .8, .3, size, TEAL, True)
        text(slide, page_title, 7.2, y, 5.4, .3, size, INK)
    text(slide, f"{n:02}", 12.17, 7.06, .55, .3, 13, LIGHT_ON_NAVY, True)
    first, last = PAGE_OF[ch["pages"][0][0]], PAGE_OF[ch["pages"][-1][0]]
    slide.notes_slide.notes_text_frame.text = (
        f"【この章】{ch['summary'].replace(chr(10), '')}\n【使う場面】{ch.get('when_long', ch['when'])}\n"
        f"【ページ】p.{first}〜{last}（{count}ページ）\n" + ch["note"] +
        "\n【読み方】各ページの上部の小さな見出しに、章と、章の中での位置が書いてあります。章を飛ばして読む場合は、"
        "各ページの【つなぎ】の一文で、前後の話がつながるようにしてあります。"
        "\n\n【出典と照合先】\n章扉（この資料の構成）。ページ番号はPowerPointのスライド番号です。")
    slide_metadata.append({"slide": n, "title": title, "sources": ["章扉（この資料の構成）"]})


for _ch in CHAPTERS:
    chapter_divider(_ch)


def reorder(final_titles):
    ids = list(prs.slides._sldIdLst)
    assert len(ids) == len(slide_metadata)
    by_title = {}
    for sid, meta in zip(ids, slide_metadata):
        assert meta["title"] not in by_title, f"duplicate title: {meta['title']}"
        by_title[meta["title"]] = (sid, meta)
    assert set(by_title) == set(final_titles), sorted(set(by_title) ^ set(final_titles))
    for sid in ids:
        prs.slides._sldIdLst.remove(sid)
    slide_metadata[:] = []
    for title in final_titles:
        sid, meta = by_title[title]
        prs.slides._sldIdLst.append(sid)
        slide_metadata.append(meta)


reorder(FINAL_TITLES)
for i, (slide, meta) in enumerate(zip(prs.slides, slide_metadata), 1):
    meta["slide"] = i
    assert PAGE_OF[meta["title"]] == i, (meta["title"], i)
    for shape in slide.shapes:
        if shape.has_text_frame and shape.left == Inches(12.17) and shape.top == Inches(7.06):
            shape.text_frame.paragraphs[0].runs[0].text = f"{i:02}"

assert len(prs.slides) == len(FINAL_TITLES)

prs.save(OUT)
(HERE / "slide_sources.json").write_text(
    json.dumps(slide_metadata, ensure_ascii=False, indent=2))
prompt_catalog = [
    ("Q1 年間と月次の稼ぎ", Q1),
    ("Q2 増減の理由を比較", Q2),
    ("Q3 食品損失の根拠", Q3),
    ("Q4 影響点の特定と対応案", Q_ACT),
    ("追加A ENE_POSの好影響", Q4),
    ("追加B 横断2件のセクター比較", Q5),
    ("Q1が崩れた場合の追い指示",
     "期間合計をグラフから外し、12か月だけ、縦軸を億円にして作り直してください。"),
]
(PROJECT / "prompts" / "presenter_prompts_ja.txt").write_text(
    "公開済みFabric Data Agentへ入力する質問の全文です。説明者ガイド（decks/）の入力①～④と同じ内容で、"
    "presentation/build_presenter_deck.py から生成されます。\n\n" +
    "\n\n".join(f"【{name}】\n{body}" for name, body in prompt_catalog) + "\n")
(QA / "slide_text_and_notes.txt").write_text(
    "\n\n".join(
        f"=== SLIDE {i} ===\n" +
        "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) +
        "\n[NOTES]\n" + slide.notes_slide.notes_text_frame.text
        for i, slide in enumerate(prs.slides, 1)))
print(f"Saved {OUT} ({len(prs.slides)} slides)")
